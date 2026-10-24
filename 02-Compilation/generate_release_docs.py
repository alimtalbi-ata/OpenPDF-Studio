"""Produce version-specific DOCX/PDF documentation before cumulative packaging.
Author: Abdelhalim Talbi. No fabricated screenshots or test evidence.
"""
from pathlib import Path
from datetime import datetime, timezone
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from xml.sax.saxutils import escape

ROOT=Path(__file__).resolve().parents[1]
version=(ROOT/'VERSION.txt').read_text(encoding='utf-8').strip()
if not version or any(c not in '0123456789.' for c in version):
    raise ValueError('Invalid version')
docs=ROOT/'05-Documentation'
docs.mkdir(exist_ok=True)
release=docs/f'RELEASE-V{version}-QUALITE-ET-ROADMAP.md'
if not release.is_file():
    raise FileNotFoundError(f'Missing release notes: {release}')
notes=release.read_text(encoding='utf-8')
sections=[
('Présentation',f'OpenPDF Studio V{version} est une application Windows pour traiter les PDF localement. Auteur : Abdelhalim Talbi.'),
('Fonctions principales','Fusion, extraction, découpage, réorganisation, compression, OCR local et conversion textuelle vers DOCX, XLSX et PPTX.'),
('Architecture','Interface Tkinter -> contrôleur -> dispatcher -> services PDF PyMuPDF / conversion Office / OCR Tesseract optionnel -> publication exclusive -> reçus SHA-256.'),
('Installation Windows','Utiliser le programme Setup présent dans 01-Application/binary. La version portable est également fournie dans le même répertoire. Ne pas lancer les deux exécutables simultanément pour une même destination.'),
('Utilisation','1. Choisir une opération. 2. Ajouter les fichiers PDF sources. 3. Vérifier le chemin de destination et les options. 4. Exécuter. 5. Examiner le résultat et les éventuels avertissements.'),
('Compilation','Depuis une invite Windows : 02-Compilation\\BUILD-WINDOWS-INSTALLER.cmd. La chaîne utilise Python 3.12, PyInstaller et Inno Setup; voir les scripts de compilation pour les prérequis exacts.'),
('Qualification','La CI Windows exécute des tests automatiques et compile un installateur et une version portable. Ces contrôles ne constituent pas une recette ergonomique manuelle. Consulter le run GitHub correspondant pour le résultat exact.'),
('Limites et transparence','Aucune capture réelle de l’interface n’est générée par cette étape. Aucune maquette n’est présentée comme une capture. OCR sur pages scannées non démontré sans Tesseract et données linguistiques. Fidélité visuelle Office non garantie.'),
('Notes de version',notes),
]
title=f'OpenPDF Studio V{version} - Guide technique et utilisateur'
doc=Document()
doc.add_heading(title,0)
doc.add_paragraph('Auteur : Abdelhalim Talbi')
doc.add_paragraph('Documentation spécifique à la version. Générée en CI.')
for heading,body in sections:
    doc.add_heading(heading,level=1)
    for line in body.splitlines():
        if line.strip():doc.add_paragraph(line.strip())
doc.save(docs/f'OpenPDF-Studio-V{version}-Guide.docx')
styles=getSampleStyleSheet()
story=[Paragraph(escape(title),styles['Title']),Paragraph('Auteur : Abdelhalim Talbi',styles['Normal']),Spacer(1,16)]
for heading,body in sections:
    story.append(Paragraph(escape(heading),styles['Heading2']))
    for line in body.splitlines():
        if line.strip():story.append(Paragraph(escape(line.strip()),styles['BodyText']))
    story.append(Spacer(1,10))
SimpleDocTemplate(str(docs/f'OpenPDF-Studio-V{version}-Guide.pdf'),pagesize=A4).build(story)
print('Generated version-specific Word/PDF documentation:',version)
