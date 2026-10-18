import unittest
import tempfile
from pathlib import Path
from pdf_table_audit import match_tables, inspect_pdf_tables, compare_pdf_office_tables

class TestPDFTables(unittest.TestCase):
    def test_exact_match(self):
        a={'id':'a','cells':[['Name','Cost'],['Disk','100']],'header':['Name','Cost'],'columns':2}
        b={**a,'id':'b'}
        result=match_tables([a],[b])
        self.assertEqual(len(result['matches']),1)
        self.assertEqual(result['matches'][0]['missing_cells'],[])
    def test_missing_cell(self):
        a={'id':'a','cells':[['Name','Cost'],['Disk','100']],'header':['Name','Cost'],'columns':2}
        b={'id':'b','cells':[['Name','Cost'],['Disk','200']],'header':['Name','Cost'],'columns':2}
        result=match_tables([a],[b],threshold=.4)
        self.assertIn('100',result['matches'][0]['missing_cells'])
    def test_unmatched(self):
        a={'id':'a','cells':[['red']],'header':['red'],'columns':1}
        b={'id':'b','cells':[['blue']],'header':['blue'],'columns':1}
        self.assertEqual(match_tables([a],[b])['unmatched_source'],['a'])
    def test_pdf_extract_and_match(self):
        import fitz
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            pdf=Path(tmp)/'table.pdf'; docx=Path(tmp)/'table.docx'
            d=fitz.open(); p=d.new_page()
            for y,row in enumerate([['Name','Cost'],['Disk','100']]):
                for x,value in enumerate(row):p.insert_text((50+x*150,70+y*25),value)
            d.save(str(pdf)); d.close()
            doc=Document(); t=doc.add_table(rows=2,cols=2)
            for i,row in enumerate([['Name','Cost'],['Disk','100']]):
                for j,value in enumerate(row):t.cell(i,j).text=value
            doc.save(docx)
            extracted=inspect_pdf_tables(pdf)
            self.assertIn('tables',extracted)
            result=compare_pdf_office_tables(pdf,docx)
            self.assertEqual(result['office_table_count'],1)
            self.assertIn('matches',result)
    def test_threshold(self):
        with self.assertRaises(ValueError):match_tables([],[],threshold=2)

if __name__=='__main__':unittest.main()
