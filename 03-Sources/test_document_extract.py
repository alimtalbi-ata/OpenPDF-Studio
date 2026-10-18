import tempfile,unittest
from pathlib import Path
from document_extract import extract_document,compare_documents,ExtractionError

class DocumentExtractTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder=Path(self.tmp.name)
        import fitz
        pdf=fitz.open(); page=pdf.new_page(width=500,height=700);page.insert_text((50,70),'Hello PDF test');pdf.save(self.folder/'source.pdf');pdf.close()
        from docx import Document
        d=Document();d.add_paragraph('Hello DOCX test');d.add_table(rows=1,cols=2).cell(0,0).text='Cell A';d.save(self.folder/'target.docx')
    def test_pdf_real_geometry(self):
        r=extract_document(self.folder/'source.pdf');self.assertEqual(r['geometry'],'native_pdf');self.assertEqual(len(r['pages']),1);self.assertTrue(any('Hello PDF' in b['text'] for b in r['pages'][0]['blocks']))
    def test_docx_synthetic_geometry(self):
        r=extract_document(self.folder/'target.docx');self.assertEqual(r['geometry'],'synthetic_logical_order');self.assertTrue(r['warnings']);self.assertTrue(any('Cell A' in b['text'] for b in r['pages'][0]['blocks']))
    def test_pipeline_pdf_to_pdf(self):
        r=compare_documents(self.folder/'source.pdf',self.folder/'source.pdf');self.assertIn('report',r['diagnostic'])
    def test_pipeline_pdf_to_docx(self):
        r=compare_documents(self.folder/'source.pdf',self.folder/'target.docx');self.assertTrue(r['warnings']);self.assertIn('report',r['diagnostic'])
    def test_unsupported_format(self):
        p=self.folder/'other.txt';p.write_text('test')
        with self.assertRaises(ExtractionError):extract_document(p)
    def test_missing_file(self):
        with self.assertRaises(ExtractionError):extract_document(self.folder/'missing.pdf')
if __name__=='__main__':unittest.main(verbosity=2)
