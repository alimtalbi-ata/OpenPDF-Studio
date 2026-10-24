"""Independent annotated-ground-truth quality gate for PDF -> editable XLSX tables.

Expected cells are supplied by an independent annotator, NEVER by PyMuPDF.
The report contains only counts, ratios and SHA-256 hashes, not document text.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re
import unicodedata

import fitz
from openpyxl import load_workbook

SCHEMA = 'openpdf.xlsx.ground_truth.v1'
MAX_ANNOTATION_BYTES = 2 * 1024 * 1024
MAX_TABLES = 256
MAX_GRID_CELLS = 200_000
MAX_ROWS = 4096
MAX_COLS = 128
SHEET_PATTERN = re.compile(r'^P([0-9]{4})-T([0-9]{3})$')


def _sha256(path: Path, max_bytes: int) -> str:
    if path.is_symlink() or not path.is_file() or path.stat().st_size > max_bytes:
        raise ValueError('Fichier absent, symbolique ou hors limite')
    before = path.stat()
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    after = path.stat()
    if (before.st_size, before.st_mtime_ns, before.st_ino) != (after.st_size, after.st_mtime_ns, after.st_ino):
        raise ValueError('Fichier modifie pendant le controle')
    return h.hexdigest()


def _unique_pairs(pairs):
    obj = {}
    for key, value in pairs:
        if key in obj:
            raise ValueError('Cle JSON dupliquee')
        obj[key] = value
    return obj


def _positive_int(value, upper):
    return type(value) is int and 1 <= value <= upper


def _load_ground_truth(path: Path, source_hash: str, pages: int):
    if path.is_symlink():
        raise ValueError('Annotation symbolique interdite')
    if not path.is_file() or path.stat().st_size > MAX_ANNOTATION_BYTES:
        raise ValueError('Annotation absente ou trop volumineuse')
    raw = path.read_bytes()
    if len(raw) > MAX_ANNOTATION_BYTES:
        raise ValueError('Annotation trop volumineuse')
    try:
        data = json.loads(raw.decode('utf-8'), object_pairs_hook=_unique_pairs,
                          parse_constant=lambda _: (_ for _ in ()).throw(ValueError('Nombre non fini')))
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise ValueError('Annotation JSON invalide') from exc
    if not isinstance(data, dict) or set(data) != {'schema', 'source_sha256', 'origin', 'tables'}:
        raise ValueError('Schema annotation invalide')
    if data['schema'] != SCHEMA or data['source_sha256'] != source_hash:
        raise ValueError('Annotation non liee au PDF source')
    if data['origin'] not in ('human_reviewed', 'synthetic_known_at_generation'):
        raise ValueError('Origine de verite terrain non reconnue')
    entries = data['tables']
    if not isinstance(entries, list) or len(entries) > MAX_TABLES:
        raise ValueError('Nombre de tableaux annotes hors limite')
    expected = {}
    total_cells = 0
    for entry in entries:
        if not isinstance(entry, dict) or set(entry) != {'page', 'index', 'cells', 'merges'}:
            raise ValueError('Tableau annote invalide')
        page, index = entry['page'], entry['index']
        if not _positive_int(page, min(pages, 9999)) or not _positive_int(index, 999):
            raise ValueError('Provenance annotee invalide')
        key = (page, index)
        if key in expected:
            raise ValueError('Tableau annote en doublon')
        rows = entry['cells']
        if not isinstance(rows, list) or not 1 <= len(rows) <= MAX_ROWS:
            raise ValueError('Lignes annotees invalides')
        width = len(rows[0]) if isinstance(rows[0], list) else 0
        if not 1 <= width <= MAX_COLS:
            raise ValueError('Colonnes annotees invalides')
        for row in rows:
            if not isinstance(row, list) or len(row) != width:
                raise ValueError('Grille annotee non rectangulaire')
            for cell in row:
                if cell is not None and (not isinstance(cell, str) or len(cell.encode('utf-16-le')) > 65534):
                    raise ValueError('Valeur de cellule annotee invalide')
        total_cells += len(rows) * width
        if total_cells > MAX_GRID_CELLS:
            raise ValueError('Trop de cellules annotees')
        merges = entry['merges']
        if not isinstance(merges, list) or len(merges) > total_cells:
            raise ValueError('Fusions annotees invalides')
        covered = set()
        merge_set = set()
        for rect in merges:
            if (not isinstance(rect, list) or len(rect) != 4 or
                    not all(type(v) is int for v in rect)):
                raise ValueError('Fusion annotee invalide')
            r0, c0, r1, c1 = rect
            if not (1 <= r0 <= r1 <= len(rows) and 1 <= c0 <= c1 <= width and (r0 != r1 or c0 != c1)):
                raise ValueError('Fusion annotee hors grille')
            span = tuple(rect)
            if span in merge_set:
                raise ValueError('Fusion annotee en doublon')
            merge_set.add(span)
            for r in range(r0, r1 + 1):
                for c in range(c0, c1 + 1):
                    if (r, c) in covered:
                        raise ValueError('Fusions annotees chevauchantes')
                    covered.add((r, c))
                    if (r, c) != (r0, c0) and rows[r - 1][c - 1] is not None:
                        raise ValueError('Cellule couverte par fusion doit etre nulle')
        expected[key] = (rows, merge_set)
    return expected, data['origin'], hashlib.sha256(raw).hexdigest()


def _normalized(value):
    if value is None:
        return None
    return ' '.join(unicodedata.normalize('NFKC', value).split())


def _ratio(numerator, denominator):
    return round(numerator / denominator, 6) if denominator else 1.0


def evaluate_ground_truth(pdf, xlsx, annotation) -> dict:
    """Compare real editable XLSX to independent annotations; never echo cells."""
    pdf, xlsx, annotation = Path(pdf), Path(xlsx), Path(annotation)
    if len({p.resolve() for p in (pdf, xlsx, annotation)}) != 3:
        raise ValueError('Fichiers source, resultat et annotation doivent etre distincts')
    source_hash = _sha256(pdf, 256 * 1024 * 1024)
    workbook_hash = _sha256(xlsx, 128 * 1024 * 1024)
    with fitz.open(pdf) as document:
        if document.needs_pass or document.page_count < 1:
            raise ValueError('PDF protege ou vide')
        pages = document.page_count
    expected, origin, annotation_hash = _load_ground_truth(annotation, source_hash, pages)
    # Reuse the production fail-closed OOXML gate, including relationship checks.
    from openpdf_xlsx_fidelity_v515 import _xlsx_preflight
    _xlsx_preflight(xlsx)
    workbook = load_workbook(xlsx, read_only=False, data_only=False, keep_links=False)
    try:
        if not workbook.worksheets or workbook.worksheets[0].title != 'Texte PDF':
            raise ValueError('Feuille Texte PDF absente ou deplacee')
        if sum(len(sheet._cells) for sheet in workbook.worksheets) > 300_000:
            raise ValueError('Classeur Excel trop volumineux pour audit')
        for sheet in workbook.worksheets:
            if sheet.sheet_state != 'visible':
                raise ValueError('Feuille Excel masquee : preuve non visible')
            for cell in sheet._cells.values():
                if cell.data_type == 'f' or cell.hyperlink is not None:
                    raise ValueError('Formule ou hyperlien actif dans Excel')
        actual = {}
        for sheet in workbook.worksheets[1:]:
            match = SHEET_PATTERN.fullmatch(sheet.title)
            if match is None:
                raise ValueError('Feuille de tableau sans provenance')
            key = (int(match.group(1)), int(match.group(2)))
            if key in actual or (sheet['A1'].value, sheet['B1'].value, sheet['C1'].value, sheet['D1'].value) != ('Page PDF', key[0], 'Tableau', key[1]):
                raise ValueError('Provenance Excel incoherente')
            actual[key] = sheet
        matching = set(expected) & set(actual)
        annotated_cells = exact = normalized = nonempty = exact_nonempty = extra_nonempty = 0
        expected_merges = observed_merges = matched_merges = 0
        for key, (rows, merges) in expected.items():
            height, width = len(rows), len(rows[0])
            annotated_cells += height * width
            nonempty += sum(cell not in (None, '') for row in rows for cell in row)
            expected_merges += len(merges)
            if key not in actual:
                continue
            sheet = actual[key]
            got_merges = {(m.min_row - 2, m.min_col, m.max_row - 2, m.max_col)
                          for m in sheet.merged_cells.ranges}
            observed_merges += len(got_merges)
            matched_merges += len(merges & got_merges)
            for r, row in enumerate(rows, 3):
                for c, value in enumerate(row, 1):
                    cell = sheet.cell(r, c)
                    if cell.data_type == 'f' or cell.hyperlink is not None:
                        raise ValueError('Formule ou hyperlien actif dans tableau Excel')
                    observed = cell.value
                    observed = None if observed == '' else observed
                    reference = None if value == '' else value
                    # A spurious value INSIDE an annotated blank cell is also
                    # an extra extraction. Previously only cells outside the
                    # expected rectangle were counted, allowing a false PASS
                    # at the corpus gate despite invented data in blank cells.
                    if reference is None and observed is not None:
                        extra_nonempty += 1
                    if observed == reference:
                        exact += 1
                        if value not in (None, ''):
                            exact_nonempty += 1
                    if isinstance(observed, (str, type(None))) and _normalized(observed) == _normalized(reference):
                        normalized += 1
            # Every nonempty cell outside the annotated grid is an extraction
            # anomaly, including row 2 (reserved blank spacer), additional
            # provenance-row columns, and out-of-grid cells. Do not count the
            # four trusted provenance fields A1:D1 as content.
            for (r, c), cell in sheet._cells.items():
                outside_grid = not (3 <= r <= height + 2 and 1 <= c <= width)
                provenance = r == 1 and 1 <= c <= 4
                if outside_grid and not provenance and cell.value not in (None, ''):
                    extra_nonempty += 1
        for key, sheet in actual.items():
            if key not in expected:
                observed_merges += len(sheet.merged_cells.ranges)
                extra_nonempty += sum(1 for (r, _), cell in sheet._cells.items()
                                      if r >= 3 and cell.value not in (None, ''))
        table_precision = _ratio(len(matching), len(actual)) if actual else (1.0 if not expected else 0.0)
        table_recall = _ratio(len(matching), len(expected)) if expected else 1.0
        merge_precision = _ratio(matched_merges, observed_merges) if observed_merges else (1.0 if not expected_merges else 0.0)
        merge_recall = _ratio(matched_merges, expected_merges) if expected_merges else 1.0
        perfect = (len(matching) == len(expected) == len(actual)
                   and exact == annotated_cells and extra_nonempty == 0
                   and expected_merges == observed_merges == matched_merges)
        result = {
            'schema': 'openpdf.xlsx.ground_truth.report.v1',
            'status': 'PASS' if perfect else 'DEVIATION',
            'scope': 'independent_annotated_cell_and_merge_comparison',
            'annotation_origin': origin,
            'source_sha256': source_hash,
            'workbook_sha256': workbook_hash,
            'annotation_sha256': annotation_hash,
            'expected_tables': len(expected), 'observed_tables': len(actual),
            'matched_tables': len(matching),
            'table_precision': table_precision, 'table_recall': table_recall,
            'annotated_cells': annotated_cells, 'exact_cells': exact,
            'cell_exact_accuracy': _ratio(exact, annotated_cells),
            'cell_normalized_accuracy': _ratio(normalized, annotated_cells),
            'annotated_nonempty_cells': nonempty, 'exact_nonempty_cells': exact_nonempty,
            'nonempty_cell_accuracy': _ratio(exact_nonempty, nonempty),
            'extra_nonempty_cells': extra_nonempty,
            'expected_merges': expected_merges, 'observed_merges': observed_merges,
            'matched_merges': matched_merges,
            'merge_precision': merge_precision, 'merge_recall': merge_recall,
            'semantic_fidelity': 'annotated_cell_content_only',
            'visual_fidelity': 'not_assessed',
        }
    finally:
        workbook.close()
    if _sha256(pdf, 256 * 1024 * 1024) != source_hash or _sha256(xlsx, 128 * 1024 * 1024) != workbook_hash:
        raise ValueError('PDF ou XLSX modifie pendant le controle')
    if _sha256(annotation, MAX_ANNOTATION_BYTES) != annotation_hash:
        raise ValueError('Annotation modifiee pendant le controle')
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description='Audit de fidelite PDF vers XLSX sur verite terrain annotee')
    parser.add_argument('--pdf', required=True)
    parser.add_argument('--xlsx', required=True)
    parser.add_argument('--annotations', required=True)
    parser.add_argument('--report')
    args = parser.parse_args(argv)
    try:
        result = evaluate_ground_truth(args.pdf, args.xlsx, args.annotations)
        if args.report:
            from openpdf_xlsx_fidelity_v515 import write_report_new
            write_report_new(args.report, result)
    except (ValueError, OSError, RuntimeError, FileExistsError) as exc:
        print(json.dumps({'status': 'ERROR', 'reason': str(exc)}, ensure_ascii=False))
        return 2
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0 if result['status'] == 'PASS' else 3


if __name__ == '__main__':
    raise SystemExit(main())
