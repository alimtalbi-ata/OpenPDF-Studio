"""V46.0 exact DOCX text/page-boundary gate for the native PDF exporter.

Scope: extracted PyMuPDF text and exporter-created page breaks, not visual
fidelity or OCR. Reject unexpected OOXML content rather than claiming that
unsupported constructs were validated.
Author: Abdelhalim Talbi.
"""
from pathlib import Path

import fitz
from docx import Document
from docx.oxml.ns import qn


class DocxRoundtripError(ValueError):
    """Native DOCX output differs from extracted PDF text or page boundaries."""


def _expected_paragraphs(source):
    """Reproduce the native exporter contract, including blank PDF pages."""
    with fitz.open(source) as pdf:
        if pdf.needs_pass:
            raise DocxRoundtripError('PDF protégé : vérification DOCX impossible')
        for page_index, page in enumerate(pdf):
            if page_index:
                yield ('page_break', '')
            for line in page.get_text(sort=True).splitlines():
                yield ('text', line)


def _actual_paragraphs(document):
    body = document.element.body
    for node in body:
        if node.tag not in (qn('w:p'), qn('w:sectPr')):
            raise DocxRoundtripError('Élément DOCX inattendu (tableau/objet)')
    if len(document.sections) != 1 or document.tables or document.inline_shapes:
        raise DocxRoundtripError('Structure DOCX inattendue')
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header,
                     section.even_page_footer):
            if any(p.text or p._p.xpath('.//w:drawing | .//w:fldChar | .//w:instrText')
                   for p in part.paragraphs) or part.tables:
                raise DocxRoundtripError('Contenu hors corps DOCX inattendu')
    for paragraph in document.paragraphs:
        runs = paragraph._p.xpath('./w:r')
        if any(node.tag not in (qn('w:pPr'), qn('w:r')) for node in paragraph._p):
            raise DocxRoundtripError('Élément de paragraphe DOCX inattendu')
        breaks = paragraph._p.xpath('.//w:br')
        if breaks:
            if (len(breaks) != 1 or len(runs) != 1
                    or breaks[0].get(qn('w:type')) != 'page'
                    or paragraph.text.strip()):
                raise DocxRoundtripError('Saut de page DOCX invalide')
            for child in runs[0]:
                if child.tag not in (qn('w:rPr'), qn('w:br')):
                    raise DocxRoundtripError('Saut de page DOCX contenant du texte')
            yield ('page_break', '')
            continue
        for run in runs:
            if any(child.tag not in (qn('w:rPr'), qn('w:t'), qn('w:tab')) for child in run):
                raise DocxRoundtripError('Contenu DOCX non textuel inattendu')
        yield ('text', paragraph.text)


def _reject_hidden_text(document):
    """Fail closed if any native text or style is marked invisible in OOXML."""
    flags = './/w:vanish | .//w:webHidden | .//w:specVanish'
    if document.element.body.xpath(flags) or document.styles.element.xpath(flags):
        raise DocxRoundtripError('Texte DOCX masqué par un attribut OOXML : sortie rejetée')


def verify_docx_text(source, destination):
    """Require an exact ordered match; never silently accept missing lines."""
    source, destination = Path(source), Path(destination)
    document = Document(destination)
    _reject_hidden_text(document)
    # python-docx does not interpret arbitrary extra OOXML parts; refuse any
    # external relationship that could make a document depend on a remote URI.
    if any(rel.is_external for part in document.part.package.parts
           for rel in part.rels.values()):
        raise DocxRoundtripError('Référence DOCX externe interdite')
    expected = iter(_expected_paragraphs(source))
    actual = iter(_actual_paragraphs(document))
    count = 0
    while True:
        marker = object()
        left = next(expected, marker)
        right = next(actual, marker)
        if left is marker and right is marker:
            break
        if left != right:
            raise DocxRoundtripError(
                f'Texte ou pagination DOCX modifié/manquant au paragraphe {count+1}')
        count += 1
    return count
