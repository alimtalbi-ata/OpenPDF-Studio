"""V531: real OCR scanned PDF to editable Excel tables, with safe failures.
Synthetic fixtures only; not a business-document or human GUI qualification.
Author: Abdelhalim Talbi.
"""
import hashlib
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

import fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
from openpdf_ocr_office_pipeline_v511 import convert_scanned_office, main
from openpdf_ocr_language_preflight_v486 import preflight_ocr_languages

TABLE = (('ID', 'STOCK', 'PRICE'), ('A01', '102', '55'),
         ('A02', '103', '80'), ('A03', '104', '75'))


def ocr_available():
    try:
        preflight_ocr_languages('eng')
        return True
    except (OSError, RuntimeError, ValueError):
        return False


def fixture(path, *, scanned=True, mixed=False, prose=False):
    with fitz.open() as doc:
        if mixed:
            doc.new_page(width=600, height=400).insert_text(
                (50, 70), 'Native page retained', fontsize=18)
        page = doc.new_page(width=600, height=400)
        if prose:
            page.insert_text((40, 80), 'INVOICE TOTAL 12345', fontsize=16)
        else:
            for row, values in enumerate(TABLE):
                for col, value in enumerate(values):
                    page.insert_text((40 + col * 140, 80 + row * 45),
                                     value, fontsize=16)
        if not scanned:
            doc.save(path)
            return
        png = page.get_pixmap(matrix=fitz.Matrix(2, 2),
                              alpha=False).tobytes('png')
    with fitz.open() as doc:
        if mixed:
            doc.new_page(width=600, height=400).insert_text(
                (50, 70), 'Native page retained', fontsize=18)
        doc.new_page(width=600, height=400).insert_image(
            fitz.Rect(0, 0, 600, 400), stream=png)
        doc.save(path)


class OCRExcelTablesV531(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.pdf = self.root / 'scan.pdf'
        self.xlsx = self.root / 'output.xlsx'

    @unittest.skipUnless(ocr_available(), 'English Tesseract model unavailable')
    def test_scanned_table_real_ocr_and_twelve_editable_cells(self):
        fixture(self.pdf)
        source_sha = hashlib.sha256(self.pdf.read_bytes()).hexdigest()
        receipt = convert_scanned_office(self.pdf, self.xlsx, language='eng',
                                         table_mode=True)
        self.assertEqual((receipt.ocr_pages, receipt.table_worksheets), (1, 1))
        self.assertTrue(receipt.table_mode)
        self.assertEqual(receipt.sha256, hashlib.sha256(self.xlsx.read_bytes()).hexdigest())
        self.assertEqual(source_sha, hashlib.sha256(self.pdf.read_bytes()).hexdigest())
        self.assertIn('human verification', ' '.join(receipt.warnings))
        book = load_workbook(self.xlsx)
        try:
            self.assertEqual(book.sheetnames, ['Texte PDF', 'P0001-T001'])
            observed = tuple(tuple(book.worksheets[1].cell(r, c).value
                                   for c in range(1, 4)) for r in range(3, 7))
            self.assertEqual(observed, TABLE)
            self.assertFalse(any(book.worksheets[1].cell(r, c).data_type == 'f'
                                 for r in range(3, 7) for c in range(1, 4)))
        finally:
            book.close()

    @unittest.skipUnless(ocr_available(), 'English Tesseract model unavailable')
    def test_mixed_document_keeps_page_provenance(self):
        fixture(self.pdf, mixed=True)
        receipt = convert_scanned_office(self.pdf, self.xlsx, language='eng',
                                         table_mode=True)
        self.assertEqual((receipt.pages, receipt.scanned_pages), (2, (2,)))
        book = load_workbook(self.xlsx)
        try:
            self.assertEqual(book.sheetnames, ['Texte PDF', 'P0002-T001'])
            self.assertEqual(book.worksheets[1]['A3'].value, 'ID')
        finally:
            book.close()

    @unittest.skipUnless(ocr_available(), 'English Tesseract model unavailable')
    def test_no_table_fails_closed(self):
        fixture(self.pdf, prose=True)
        with self.assertRaisesRegex(ValueError, 'Aucun tableau'):
            convert_scanned_office(self.pdf, self.xlsx, language='eng',
                                   table_mode=True)
        self.assertFalse(self.xlsx.exists())
        self.assertFalse(list(self.root.glob('.openpdf-*')))

    def test_invalid_format_and_type_rejected(self):
        fixture(self.pdf, scanned=False)
        with self.assertRaisesRegex(ValueError, 'reserve au format XLSX'):
            convert_scanned_office(self.pdf, self.root / 'out.docx', table_mode=True)
        for invalid in ('yes', 1, None):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                convert_scanned_office(self.pdf, self.xlsx, table_mode=invalid)
        self.assertFalse(self.xlsx.exists())

    def test_native_table_does_not_invoke_ocr(self):
        fixture(self.pdf, scanned=False)
        with patch('openpdf_ocr_office_pipeline_v511.searchable_pdf',
                   side_effect=AssertionError('Unexpected OCR')):
            receipt = convert_scanned_office(self.pdf, self.xlsx, table_mode=True)
        self.assertEqual((receipt.ocr_pages, receipt.table_worksheets), (0, 1))

    @unittest.skipUnless(ocr_available(), 'English Tesseract model unavailable')
    def test_legacy_default_remains_single_sheet(self):
        fixture(self.pdf)
        receipt = convert_scanned_office(self.pdf, self.xlsx, language='eng')
        self.assertFalse(receipt.table_mode)
        book = load_workbook(self.xlsx)
        try:
            self.assertEqual(book.sheetnames, ['Texte PDF'])
        finally:
            book.close()

    @unittest.skipUnless(ocr_available(), 'English Tesseract model unavailable')
    def test_cli_explicit_opt_in(self):
        fixture(self.pdf)
        self.assertEqual(main([str(self.pdf), str(self.xlsx),
                               '--language', 'eng', '--editable-tables']), 0)
        book = load_workbook(self.xlsx)
        try:
            self.assertEqual(book.sheetnames, ['Texte PDF', 'P0001-T001'])
        finally:
            book.close()

    def test_gui_dispatch_combined_options(self):
        import openpdf_operations as operations
        fixture(self.pdf, scanned=False)
        seen = []
        fake = types.ModuleType('openpdf_ocr_office_pipeline_v511')
        def convert(*args, **kwargs):
            seen.append(kwargs)
            return 'ROUTED'
        fake.convert_scanned_office = convert
        with patch.dict(sys.modules, {'openpdf_ocr_office_pipeline_v511': fake}), \
             patch.object(operations, 'write_event'), \
             patch.object(operations, '_legacy_run_operation',
                          side_effect=AssertionError('Legacy converter called')):
            result = operations.run_operation('xlsx', (str(self.pdf),),
                                              str(self.xlsx), table_mode=True,
                                              ocr_before_office=True, language='eng')
        self.assertEqual(result, 'ROUTED')
        self.assertEqual(seen, [{'language': 'eng', 'table_mode': True}])
        gui = (ROOT / '03-Sources/openpdf_desktop.py').read_text(encoding='utf-8')
        self.assertNotIn('Tableaux Excel et OCR Office ne sont pas combinables', gui)


if __name__ == '__main__':
    unittest.main()
