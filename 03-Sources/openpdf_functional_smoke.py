"""Isolated PDF round-trip smoke; no external files. Author: Abdelhalim Talbi."""
import json
import sys
import tempfile
from pathlib import Path

def run():
    try:
        import fitz
        with tempfile.TemporaryDirectory(prefix="openpdf-smoke-") as tmp:
            root = Path(tmp)
            first = root / "first.pdf"
            second = root / "second.pdf"
            merged = root / "merged.pdf"
            split = root / "split.pdf"
            with fitz.open() as d:
                page = d.new_page()
                page.insert_text((72, 72), "OpenPDF functional smoke page A")
                d.save(first)
            with fitz.open() as d:
                page = d.new_page()
                page.insert_text((72, 72), "OpenPDF functional smoke page B")
                d.save(second)
            with fitz.open() as result:
                for path in (first, second):
                    with fitz.open(path) as source:
                        result.insert_pdf(source)
                result.save(merged)
            with fitz.open(merged) as d:
                assert len(d) == 2, "merge page count"
                assert "page A" in d[0].get_text(), "merge first text"
                assert "page B" in d[1].get_text(), "merge second text"
                with fitz.open() as output:
                    output.insert_pdf(d, from_page=1, to_page=1)
                    output.save(split)
            with fitz.open(split) as d:
                assert len(d) == 1 and "page B" in d[0].get_text(), "split validation"
        # Real end-to-end conversions executed from the packaged Windows binary.
        from openpdf_core_v260 import pdf_to_docx, pdf_to_xlsx, pdf_to_pptx
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        with tempfile.TemporaryDirectory(prefix="openpdf-conversion-smoke-") as tmp:
            root=Path(tmp)
            source=root/"source.pdf"
            with fitz.open() as d:
                page=d.new_page()
                page.insert_text((72,100),"FACTURE TEST CONVERSION",fontsize=16)
                page.insert_text((72,140),"Ligne Excel controle")
                page.insert_text((72,180),'=HYPERLINK("https://example.invalid","ouvrir")')
                d.save(source)
            docx=root/"output.docx"; xlsx=root/"output.xlsx"; pptx=root/"output.pptx"
            assert pdf_to_docx(source,docx)==[docx]
            assert pdf_to_xlsx(source,xlsx)==[xlsx]
            assert pdf_to_pptx(source,pptx)==[pptx]
            word=Document(docx)
            assert "FACTURE TEST CONVERSION" in "\\n".join(p.text for p in word.paragraphs)
            excel=load_workbook(xlsx,read_only=True)
            try:
                assert any("FACTURE TEST CONVERSION" in str(row[2]) for row in excel.active.iter_rows(min_row=2,values_only=True))
                values=[row[2] for row in excel.active.iter_rows(min_row=2,values_only=True)]
                assert "'=HYPERLINK(\"https://example.invalid\",\"ouvrir\")" in values,repr(values)
            finally:
                excel.close()
            slides=Presentation(pptx)
            assert any("FACTURE TEST CONVERSION" in shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame)
        # Verify that image-bearing PDF pages retain embedded image content in DOCX.
        with tempfile.TemporaryDirectory(prefix="openpdf-image-smoke-") as tmp:
            root=Path(tmp); source=root/"image-source.pdf"; target=root/"image-output.docx"
            with fitz.open() as graphic:
                p=graphic.new_page(width=220,height=100)
                p.draw_rect(fitz.Rect(10,10,210,90),color=(0,0,0),fill=(0.3,0.6,0.9))
                png=p.get_pixmap(dpi=100).tobytes("png")
            with fitz.open() as d:
                p=d.new_page();p.insert_text((72,72),"ILLUSTRATION TEST")
                p.insert_image(fitz.Rect(72,100,292,200),stream=png)
                d.save(source)
            pdf_to_docx(source,target)
            from zipfile import ZipFile
            with ZipFile(target) as archive:
                media=[name for name in archive.namelist() if name.startswith("word/media/")]
                assert media,"DOCX lost the PDF embedded image"
        # Real bordered table: ensure PDF cell geometry becomes editable Excel cells.
        with tempfile.TemporaryDirectory(prefix="openpdf-table-smoke-") as tmp:
            root=Path(tmp); source=root/"table-source.pdf"; target=root/"table-output.xlsx"
            with fitz.open() as d:
                p=d.new_page(width=595,height=842)
                xs=(60,250,440); ys=(100,145,190,235)
                for x in xs: p.draw_line((x,ys[0]),(x,ys[-1]),color=(0,0,0),width=1)
                for y in ys: p.draw_line((xs[0],y),(xs[-1],y),color=(0,0,0),width=1)
                for i,row in enumerate((("Produit","Montant"),("Stockage","125"),("Sauvegarde","75"))):
                    for j,value in enumerate(row): p.insert_text((xs[j]+10,ys[i]+28),value,fontsize=12)
                d.save(source)
            pdf_to_xlsx(source,target)
            workbook=load_workbook(target,read_only=True,data_only=True)
            try:
                tables=[workbook[name] for name in workbook.sheetnames if name.startswith("Tableau ")]
                assert tables,"No structured Excel table detected from bordered PDF"
                values=[[str(cell or "").strip() for cell in row] for sheet in tables for row in sheet.iter_rows(values_only=True)]
                assert any(row[:2]==["Produit","Montant"] for row in values),repr(values)
                assert any(row[:2]==["Stockage","125"] for row in values),repr(values)
                assert any(row[:2]==["Sauvegarde","75"] for row in values),repr(values)
            finally:
                workbook.close()
        # Validate rotation parameters at the desktop UI boundary, not only in the PDF engine.
        from openpdf_desktop import validate_page_options
        for valid_angle in ('90','180','270'):
            assert validate_page_options('rotation',order=valid_angle) is True
        for invalid_angle in ('0','45','360','abc',''):
            try:
                validate_page_options('rotation',order=invalid_angle)
            except ValueError:
                pass
            else:
                raise AssertionError(f"UI accepted invalid rotation angle: {invalid_angle!r}")
        # Exercise the same public dispatcher used by the GUI, for all rotation angles.
        from openpdf_operations import run_operation
        with tempfile.TemporaryDirectory(prefix="openpdf-rotation-smoke-") as tmp:
            root=Path(tmp); source=root/"rotation-source.pdf"
            with fitz.open() as d:
                p=d.new_page(width=320,height=480)
                p.insert_text((50,80),"ROTATION TEST PAGE")
                d.save(source)
            original=source.read_bytes()
            for angle in (90,180,270):
                target=root/("rotation-"+str(angle)+".pdf")
                run_operation("rotation",[source],target,order=str(angle))
                with fitz.open(target) as d:
                    assert d.page_count==1 and d[0].rotation==angle,(angle,d[0].rotation)
                    assert "ROTATION TEST PAGE" in d[0].get_text(),angle
                assert source.read_bytes()==original,"Rotation modified source PDF"
        print("OPENPDF_FUNCTIONAL_TEST_OK")
        return 0
    except Exception as exc:
        print(json.dumps({"status": "FAILED", "error": str(exc)}), file=sys.stderr)
        return 1

if __name__ == "__main__":
    sys.exit(run())
