"""Generate V50.8-specific Word/PDF from the actual source contract.

Optional real UI screenshot is used only when a capture exists and is explicitly
identified as Linux Xvfb, never as a Windows capture.
Author: Abdelhalim Talbi.
"""
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / '05-Documentation'
PROOF = ROOT / '04-Tests-et-Qualification' / 'Preuves-Reelles-V50.8'
TITLE = 'OpenPDF Studio V50.8 - Architecture et qualification'
AUTHOR = 'Abdelhalim Talbi'
SECTIONS = [
    ('Probleme de fidelite corrige',
     'Le rendu RGB sur fond blanc peut masquer une perte de transparence. Un PDF vide transparent et un PDF rempli de blanc ont le meme rendu RGB sur fond blanc, mais des comportements differents sur un fond colore. V50.8 ajoute une verification independante de la couche alpha avant publication.'),
    ('Architecture et pipeline',
     'Source PDF -> candidats prives -> moteur de compression -> verification du texte, des liens, des formulaires et de la structure -> audit RGB a echelle bornee -> couverture pages/pixels V50.7 -> verification alpha V50.8 -> controle SHA-256 -> publication exclusive du plus petit candidat conforme. Chaque echec rejette le candidat.'),
    ('Algorithme et budgets',
     'Le nouveau module openpdf_compression_alpha_v508.py controle les pages, les geometries, les rotations et les budgets de pixels avant allocation. Il rend chaque page source/candidat en RGBA avec annotations a la meme echelle que la politique RGB, compare tous les octets alpha, puis verifie a nouveau les SHA-256. Les pixels transparents, translucides et opaques sont comptes dans le recu.'),
    ('Preuves liees aux fichiers',
     'Le recu selected_quality.alpha contient les SHA-256 source et candidat, les pages et pixels controles, le nombre de divergences alpha (zero pour une sortie acceptee), les classes de transparence et l echelle. Le rapport ajoute alpha_verification. Le schema v50.6 reste conserve pour compatibilite.'),
    ('Qualification fonctionnelle locale',
     'Quatorze tests V50.8 ont ete executes avec PyMuPDF sur de vrais PDF synthetiques : transparence partielle, page tournee, fond blanc et transparence visuellement equivalentes en RGB mais distinctes en alpha, recus falsifies, fichiers chiffres, budgets, collision de destination et refus de publication. Les tests V50.5 a V50.7 sont rejoues en regression. Les resultats CI et Windows doivent etre lus dans leurs journaux, et ne sont pas supposes par le document.'),
    ('Limites et risques',
     'La comparaison alpha couvre tous les pixels uniquement a l echelle de rendu configuree (0.5 par defaut). Elle augmente le temps de calcul et ne garantit ni fidelite a haute resolution, ni PDF/A, OCR, impression, accessibilite, ni conversions Office editables. Un corpus de documents metier reste necessaire.'),
    ('Prochaines priorites',
     'Etendre les corpus PDF metier aux scans, transparences complexes, graphiques, formulaires et liens ; mesurer les temps et la memoire ; qualifier visuellement l interface Windows avec captures reelles ; approfondir la conversion Word/Excel/PowerPoint editable et la fidelite des tableaux.'),
]


def generate():
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from docx.shared import Inches, Pt
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PDFImage, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors

    DOCS.mkdir(exist_ok=True)
    diagram = DOCS / 'Architecture-Fidelite-Alpha-V50.8.png'
    im = Image.new('RGB', (1320, 280), 'white')
    d = ImageDraw.Draw(im)
    font = ImageFont.load_default(size=24)
    boxes = [('PDF source', 15), ('Compression\nprivee', 230), ('Audit RGB\nV50.6', 445),
             ('Couverture\nV50.7', 660), ('Alpha RGBA\nV50.8', 875), ('Publication\nexclusive', 1090)]
    for label, x in boxes:
        d.rounded_rectangle((x, 65, x+190, 205), radius=16, outline='#1a3850', width=4, fill='#eef5fa')
        d.multiline_text((x+15, 111), label, fill='#102a43', spacing=9, font=font)
        if x < 1090:
            d.line((x+191, 135, x+212, 135), fill='#1a3850', width=4)
    im.save(diagram)

    doc = Document()
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = 'OPENPDF STUDIO  /  V50.8'
    doc.add_heading(TITLE, 0)
    doc.add_paragraph('Auteur : ' + AUTHOR)
    doc.add_paragraph('Documentation specifique V50.8 - 23 octobre 2026')
    doc.add_picture(str(diagram), width=Inches(6.3))
    for heading, body in SECTIONS:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)
    doc.add_heading('Statuts de qualification (ne pas confondre)', level=1)
    rows = [('Controle', 'Statut a documenter'),
            ('Tests fonctionnels Linux', '14 nouveaux tests PASS (execution locale)'),
            ('CI Ubuntu/Windows', 'A confirmer dans les journaux GitHub Actions'),
            ('Portable Windows V50.8', 'Recu de smoke natif requis'),
            ('Installateur Windows V50.8', 'Recu installation/desinstallation requis'),
            ('Validation humaine GUI', 'Non realisee')]
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Light Shading Accent 1'
    for a, b in rows:
        c = t.add_row().cells
        c[0].text, c[1].text = a, b
    screenshot = PROOF / 'UI-Linux-Xvfb-V50.8.png'
    if screenshot.is_file():
        doc.add_heading('Capture reelle de l interface (Linux Xvfb)', level=1)
        doc.add_paragraph('Capture de l application lancee sous Linux/Xvfb. Ce n est pas une preuve de qualification graphique Windows.')
        doc.add_picture(str(screenshot), width=Inches(6.2))
    doc.core_properties.author = AUTHOR
    doc.save(DOCS / 'OpenPDF-Studio-V50.8-Guide.docx')

    styles = getSampleStyleSheet()
    story = [Paragraph(TITLE, styles['Title']), Spacer(1, 12),
             Paragraph('Auteur : ' + AUTHOR + ' | 23 octobre 2026', styles['Normal']),
             Spacer(1, 16), PDFImage(str(diagram), width=495, height=105), Spacer(1, 16)]
    for heading, body in SECTIONS:
        story += [Paragraph(heading, styles['Heading2']), Paragraph(body, styles['BodyText']), Spacer(1, 10)]
    story.append(Paragraph('Statuts de qualification (preuves distinctes)', styles['Heading2']))
    table = Table(rows, colWidths=[165, 310])
    table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
        ('GRID', (0,0), (-1,-1), 0.4, colors.grey), ('VALIGN',(0,0),(-1,-1),'TOP'),
        ('LEFTPADDING',(0,0),(-1,-1),8), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story += [table, Spacer(1, 15)]
    if screenshot.is_file():
        story += [Paragraph('Capture reelle de l interface - Linux Xvfb (pas Windows)', styles['Heading2']),
                  PDFImage(str(screenshot), width=475, height=309)]
    SimpleDocTemplate(str(DOCS / 'OpenPDF-Studio-V50.8-Guide.pdf'), pagesize=A4,
                      leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48).build(story)
    return diagram


if __name__ == '__main__':
    print(generate())
