"""Real PDF and Office lexical coverage regression tests.
Author: Abdelhalim Talbi.
"""
import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'03-Sources'))
import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from openpdf_fidelity_v352 import measure_text_coverage,_words

class CoverageTests(unittest.TestCase):
    def test_unicode_nfkc_and_casefold(self):
        self.assertEqual(_words('ＣＡＦＥ́ Société'),['café','société'])
    def test_docx_duplicate_tokens_are_counted(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);src=p/'in.pdf';dst=p/'out.docx'
            doc=fitz.open();page=doc.new_page();page.insert_text((40,70),'alpha alpha beta');doc.save(src);doc.close()
            word=Document();word.add_paragraph('alpha beta');word.save(dst)
            score=measure_text_coverage(src,dst)
            self.assertEqual(score.source_words,3)
            self.assertEqual(score.matched_words,2)
            self.assertAlmostEqual(score.coverage,0.6667,places=4)
    def test_xlsx_real_text(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);src=p/'in.pdf';dst=p/'out.xlsx'
            doc=fitz.open();doc.new_page().insert_text((40,70),'revenue forecast');doc.save(src);doc.close()
            wb=Workbook();wb.active['A1']='revenue forecast';wb.save(dst)
            self.assertEqual(measure_text_coverage(src,dst).coverage,1.0)
    def test_pptx_real_text(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d);src=p/'in.pdf';dst=p/'out.pptx'
            doc=fitz.open();doc.new_page().insert_text((40,70),'storage performance');doc.save(src);doc.close()
            prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='storage performance';prs.save(dst)
            self.assertEqual(measure_text_coverage(src,dst).coverage,1.0)
if __name__=='__main__':unittest.main()
