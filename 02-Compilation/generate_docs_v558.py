"""Generate version-specific V558 Word/PDF documentation and architecture diagrams.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
AUTHOR = "Abdelhalim Talbi"
TITLE = "OpenPDF Studio V51.4 / V558 - Bounded strict JSON release parsing"
DOCX_NAME = "OpenPDF-Studio-V51.4-V558-Bounded-JSON-Guide.docx"
PDF_NAME = "OpenPDF-Studio-V51.4-V558-Bounded-JSON-Guide.pdf"
PNG_NAME = "Architecture-V558-Bounded-JSON.png"
SVG_NAME = "Architecture-V558-Bounded-JSON.svg"


def build_context(*, commit_sha=None, pr_number=None, workflow_run_id=None):
    sha = commit_sha
    if sha is None:
        sha = os.environ.get("OPENPDF_HEAD_SHA") or os.environ.get("GITHUB_SHA", "")
    if not isinstance(sha, str) or not re.fullmatch(r"[0-9a-f]{40}", sha):
        sha = "NOT-VERIFIED"
    pr = str(pr_number if pr_number is not None else os.environ.get("OPENPDF_PULL_REQUEST_NUMBER", "")).strip()
    run_id = str(workflow_run_id if workflow_run_id is not None else os.environ.get("OPENPDF_RUN_ID", "")).strip()
    return {"commit_sha": sha, "pr_number": pr or "not supplied", "workflow_run_id": run_id or "not supplied"}


def build_sections(context):
    pr_url = ("https://github.com/alimtalbi-ata/OpenPDF-Studio/pull/" + context["pr_number"]
              if context["pr_number"].isdigit() else "not available")
    run_url = ("https://github.com/alimtalbi-ata/OpenPDF-Studio/actions/runs/" + context["workflow_run_id"]
               if context["workflow_run_id"].isdigit() else "not available")
    sha_verified_format = context["commit_sha"] != "NOT-VERIFIED"
    return [
        ("1. Risque de ressources et d'interopérabilité",
         "V557 a supprimé les clés JSON répétées et les nombres non finis, mais un JSON très volumineux, profondément imbriqué, portant un entier arbitrairement grand ou un substitut Unicode isolé peut encore coûter des ressources importantes ou être interprété différemment selon l'outil."),
        ("2. Implémentation",
         "V558 complète V557 par des plafonds vérifiés avant analyse : 32 MiB par document JSON, 128 niveaux d'objets/tableaux, et des entiers dans l'intervalle signé 64 bits. Après parsing, les clés et valeurs sont parcourues pour refuser les substituts Unicode isolés (surrogates) ; les paires de caractères supplémentaires valides restent acceptées. Le contrôle couvre toujours l'inventaire, le manifeste binaire et les reçus Windows."),
        ("3. Défense en profondeur",
         "Les contrôles V553 de chemins, V554 d'amplification, V555 des enregistrements ZIP, V556 du descripteur et V557 des clés JSON restent actifs. Les budgets V558 sont appliqués avant la création de l'arbre JSON et ne remplacent pas la validation de schéma, les tailles déclarées, les empreintes ni l'attestation des binaires Windows."),
        ("4. Tests fonctionnels",
         "Les tests V558 vérifient la limite de taille avant tout appel au parseur, la limite de nesting avant parsing récursif, les bornes signées 64 bits, les substituts Unicode isolés dans les clés et les valeurs, les caractères supplémentaires valides, ainsi que la conservation des régressions V557 sur les inventaires ZIP réels."),
        ("5. Limites et qualification",
         "Les plafonds réduisent les risques de déni de service par JSON et les divergences de chaînes Unicode, sans prouver l'absence de toutes les vulnérabilités de parsing. Les tests fonctionnels restent distincts des builds natifs Windows, du cycle installation/désinstallation et de la vérification CRC/inventaire/SHA-256 du ZIP sur le SHA exact. Aucun SHA-256 ne signe l'identité de l'éditeur."),
        ("6. Prochaine priorité produit",
         "Consolider la publication V554-V558, puis qualifier la fidélité PDF vers DOCX/XLSX/PPTX à partir d'un corpus métier annoté indépendamment : précision, rappel, ordre des objets, géométrie et fidélité visuelle mesurés séparément."),
        ("Traçabilité",
         "Commit: " + context["commit_sha"] + "\nPR: " + pr_url + "\nWorkflow: " + run_url + "\n" +
         ("Le SHA a le format attendu; les reçus CI doivent établir la qualification de ce commit."
          if sha_verified_format else
          "Le SHA source est absent ou invalide; ce document est explicitement non qualifié.")),
    ]


def svg_diagram():
    return """<svg xmlns="http://www.w3.org/2000/svg" width="1120" height="360" viewBox="0 0 1120 360" role="img" aria-labelledby="title desc">
<title id="title">OpenPDF Studio V558 strict JSON release integrity</title>
<desc id="desc">Generated architecture diagram, not an application screenshot.</desc>
<defs><marker id="arrow" markerWidth="9" markerHeight="9" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#24566b"/></marker></defs>
<rect x="48" y="92" width="270" height="148" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="183" y="124" text-anchor="middle" font-family="Segoe UI,Arial" font-size="16" fill="#153747">JSON inputs</text>
<text x="183" y="154" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#153747">UTF-8 / optional BOM</text>
<text x="183" y="180" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#153747">32 MiB / 128 nesting levels</text>
<path d="M318 165 L402 165" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="410" y="72" width="300" height="188" rx="12" fill="#fff3df" stroke="#a75c19" stroke-width="2"/>
<text x="560" y="108" text-anchor="middle" font-family="Segoe UI,Arial" font-size="16" fill="#75410d">Strict JSON parser</text>
<text x="560" y="140" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#75410d">unique keys + bounded depth</text>
<text x="560" y="166" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#75410d">non-finite, huge integers, surrogates rejected</text>
<text x="560" y="192" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#75410d">schema validation follows</text>
<path d="M710 165 L794 165" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="802" y="92" width="270" height="148" rx="12" fill="#e7f5e9" stroke="#2b7545" stroke-width="2"/>
<text x="937" y="124" text-anchor="middle" font-family="Segoe UI,Arial" font-size="16" fill="#1d5833">Fail-closed receipt</text>
<text x="937" y="154" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#1d5833">inventory + SHA-256</text>
<text x="937" y="180" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#1d5833">binary / smoke / install</text>
<rect x="48" y="288" width="1024" height="48" rx="8" fill="#f4f0ff" stroke="#6750a4" stroke-width="1.5"/>
<text x="560" y="308" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#3e2d73">V553 paths | V554 ZIP budgets | V555 ZIP records | V556 one descriptor | V557 strict JSON | V558 JSON size/depth/scalar bounds</text>
<text x="560" y="326" text-anchor="middle" font-family="Segoe UI,Arial" font-size="11" fill="#3e2d73">Architecture diagram only — not a screenshot or publisher signature.</text>
</svg>"""


def draw_diagram_png(path):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1120, 360), "white")
    draw = ImageDraw.Draw(image)
    font_path = next((candidate for candidate in (
        "C:/Windows/Fonts/segoeui.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ) if Path(candidate).is_file()), None)
    title_font = ImageFont.truetype(font_path, 16) if font_path else ImageFont.load_default()
    body_font = ImageFont.truetype(font_path, 12) if font_path else ImageFont.load_default()
    boxes = [
        (48, 92, 318, 240, "JSON inputs", "UTF-8 / optional BOM", "32 MiB / 128 nesting levels", "#e8f3f5", "#24566b"),
        (410, 72, 710, 260, "Strict JSON parser", "unique keys + bounded depth", "non-finite, huge integers, surrogates rejected", "#fff3df", "#a75c19"),
        (802, 92, 1072, 240, "Fail-closed receipt", "inventory + SHA-256", "binary / smoke / install", "#e7f5e9", "#2b7545"),
    ]
    for x0, y0, x1, y1, title, line1, line2, fill, outline in boxes:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=12, fill=fill, outline=outline, width=2)
        draw.text((x0 + 14, y0 + 28), title, font=title_font, fill="#153747")
        draw.text((x0 + 14, y0 + 68), line1, font=body_font, fill="#153747")
        draw.text((x0 + 14, y0 + 94), line2, font=body_font, fill="#153747")
    for start, end in ((318, 410), (710, 802)):
        y = 165
        draw.line((start + 3, y, end - 9, y), fill="#24566b", width=3)
        draw.polygon([(end - 8, y), (end - 19, y - 6), (end - 19, y + 6)], fill="#24566b")
    draw.rounded_rectangle((48, 288, 1072, 336), radius=8, fill="#f4f0ff", outline="#6750a4", width=2)
    draw.text((64, 298), "V553 paths | V554 ZIP budgets | V555 ZIP records | V556 one descriptor | V557 strict JSON | V558 JSON size/depth/scalar bounds",
              font=body_font, fill="#3e2d73")
    draw.text((64, 316), "Architecture diagram only — not a screenshot or publisher signature.",
              font=body_font, fill="#3e2d73")
    image.save(path, format="PNG", optimize=True)


def generate(root=ROOT, *, commit_sha=None, pr_number=None, workflow_run_id=None):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PdfImage

    root = Path(root)
    docs = root / "05-Documentation"
    docs.mkdir(parents=True, exist_ok=True)
    context = build_context(commit_sha=commit_sha, pr_number=pr_number, workflow_run_id=workflow_run_id)
    sections = build_sections(context)
    svg_path, png_path = docs / SVG_NAME, docs / PNG_NAME
    docx_path, pdf_path = docs / DOCX_NAME, docs / PDF_NAME
    svg_path.write_text(svg_diagram(), encoding="utf-8")
    draw_diagram_png(png_path)

    document = Document()
    document.core_properties.author = AUTHOR
    document.core_properties.title = TITLE
    document.core_properties.subject = "Bounded strict JSON parsing for release integrity evidence"
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading(TITLE, level=0)
    intro = document.add_paragraph()
    intro.add_run("Auteur : ").bold = True
    intro.add_run(AUTHOR + "  ")
    intro.add_run("Version applicative : ").bold = True
    intro.add_run("51.4  ")
    intro.add_run("Jalon : ").bold = True
    intro.add_run("V558")
    document.add_paragraph("Commit source : " + context["commit_sha"])
    document.add_paragraph("PR : " + context["pr_number"] + " | Workflow : " + context["workflow_run_id"])
    if context["commit_sha"] == "NOT-VERIFIED":
        note = document.add_paragraph("SHA absent ou invalide : guide non qualifié.")
        for run in note.runs:
            run.bold = True
    document.add_picture(str(png_path), width=Inches(6.25))
    caption = document.add_paragraph("Figure 1. Schéma généré, pas une capture réelle de l'interface.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, body in sections:
        document.add_heading(heading, level=1)
        for line in body.split("\n"):
            if line.strip():
                document.add_paragraph(line)
    document.save(docx_path)

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="V558Body", parent=styles["BodyText"],
                              fontName="Helvetica", fontSize=9, leading=12, spaceAfter=6))
    story = [
        Paragraph(escape(TITLE), styles["Title"]),
        Paragraph(escape("Auteur : " + AUTHOR + " | Commit : " + context["commit_sha"]), styles["V558Body"]),
        Paragraph(escape("PR : " + context["pr_number"] + " | Workflow : " + context["workflow_run_id"]), styles["V558Body"]),
        Spacer(1, 4 * mm),
        PdfImage(str(png_path), width=180 * mm, height=180 * mm * 360 / 1120),
        Spacer(1, 4 * mm),
    ]
    for heading, body in sections:
        story.append(Paragraph(escape(heading), styles["Heading2"]))
        story.append(Paragraph(escape(body).replace("\n", "<br/>"), styles["V558Body"]))
    pdf = SimpleDocTemplate(str(pdf_path), pagesize=A4, rightMargin=14 * mm,
                            leftMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=TITLE, author=AUTHOR)
    pdf.build(story)
    return {"docx": docx_path, "pdf": pdf_path, "png": png_path, "svg": svg_path}


if __name__ == "__main__":
    result = generate()
    for key, path in result.items():
        print(key + ": " + str(path))
