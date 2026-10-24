"""V51.5: real PDF geometric merged cells preserved as editable XLSX ranges.

The fixtures are synthetic PDF documents generated during the tests, not
screenshots or a qualification of the Windows user interface.
Author: Abdelhalim Talbi.
"""
from pathlib import Path
from types import SimpleNamespace
import sys
import tempfile
import unittest
from unittest.mock import patch
import hashlib

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))

import fitz
from openpyxl import load_workbook
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_tables_v514 import _merged_regions, verify_xlsx_tables


def fixture(path, *, orientation='horizontal', formula=False):
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    x, y, w, h = 50, 80, 400, 120
    sh = page.new_shape()
    if orientation == 'horizontal':
        for j in range(4):
            sh.draw_line((x, y + j * h / 3), (x + w, y + j * h / 3))
        for start, end in (((x, y), (x, y+h)), ((x+w, y), (x+w, y+h)),
                           ((x+w/2, y+h/3), (x+w/2, y+h))):
            sh.draw_line(start, end)
        labels = [(x+12, y+25, '=2+2' if formula else 'Global storage'),
                  (x+12, y+65, 'SAN-A'), (x+w/2+12, y+65, '20 To'),
                  (x+12, y+105, 'SAN-B'), (x+w/2+12, y+105, '50 To')]
    elif orientation == 'vertical':
        for j in (0, 1, 3):
            sh.draw_line((x, y + j * h / 3), (x + w, y + j * h / 3))
        sh.draw_line((x+w/2, y+2*h/3), (x+w, y+2*h/3))
        for j in (0, 1, 2):
            sh.draw_line((x + j*w/2, y), (x + j*w/2, y+h))
        labels = [(x+12, y+25, 'Volume'), (x+w/2+12, y+25, 'Capacite'),
                  (x+12, y+65, 'SAN-A'), (x+w/2+12, y+65, '20 To'),
                  (x+w/2+12, y+105, '50 To')]
    else:
        raise ValueError(orientation)
    sh.finish(color=(0, 0, 0), width=1)
    sh.commit()
    for px, py, label in labels:
        page.insert_text((px, py), label, fontsize=10)
    doc.save(path)
    doc.close()


class TestMergedCellsV515(unittest.TestCase):
    def _convert(self, orientation, *, formula=False):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        src, dst = Path(temp.name)/'source.pdf', Path(temp.name)/'result.xlsx'
        fixture(src, orientation=orientation, formula=formula)
        convert_office(src, dst, table_mode=True)
        return src, dst

    def test_horizontal_header_is_real_merged_cell(self):
        src, dst = self._convert('horizontal')
        wb = load_workbook(dst)
        try:
            ws = wb.worksheets[1]
            self.assertEqual(str(next(iter(ws.merged_cells.ranges))), 'A3:B3')
            self.assertEqual(ws['A3'].value, 'Global storage')
            self.assertEqual(ws['B3'].value, None)
            self.assertEqual(ws['B4'].value, '20 To')
        finally:
            wb.close()
        self.assertEqual(verify_xlsx_tables(src, dst).tables, 1)

    def test_vertical_span_is_real_merged_cell(self):
        src, dst = self._convert('vertical')
        wb = load_workbook(dst)
        try:
            ws = wb.worksheets[1]
            self.assertEqual(str(next(iter(ws.merged_cells.ranges))), 'A4:A5')
            self.assertEqual(ws['A4'].value, 'SAN-A')
            self.assertIsNone(ws['A5'].value)
            self.assertEqual(ws['B5'].value, '50 To')
        finally:
            wb.close()
        self.assertEqual(verify_xlsx_tables(src, dst).tables, 1)

    def test_merged_formula_stays_text(self):
        _, dst = self._convert('horizontal', formula=True)
        wb = load_workbook(dst)
        try:
            self.assertEqual(wb.worksheets[1]['A3'].value, '=2+2')
            self.assertEqual(wb.worksheets[1]['A3'].data_type, 's')
        finally:
            wb.close()

    def test_tampered_merge_is_detected(self):
        src, dst = self._convert('horizontal')
        wb = load_workbook(dst)
        wb.worksheets[1].unmerge_cells('A3:B3')
        wb.save(dst)
        wb.close()
        with self.assertRaisesRegex(ValueError, 'Fusions'):
            verify_xlsx_tables(src, dst)

    def test_extra_merge_is_detected(self):
        src, dst = self._convert('horizontal')
        wb = load_workbook(dst)
        wb.worksheets[1].merge_cells('A4:B4')
        wb.save(dst)
        wb.close()
        with self.assertRaisesRegex(ValueError, 'Fusions'):
            verify_xlsx_tables(src, dst)

    def test_ambiguous_pdf_grid_is_rejected(self):
        table = SimpleNamespace(
            bbox=(0, 0, 200, 100),
            rows=[SimpleNamespace(cells=[(0, 0, 130, 50), None]),
                  SimpleNamespace(cells=[(0, 50, 100, 100), (100, 50, 200, 100)])],
        )
        with self.assertRaisesRegex(ValueError, 'ambigues'):
            _merged_regions(table, [['a', None], ['b', 'c']], 2)

    def test_no_output_after_ambiguous_geometry(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d)/'source.pdf', Path(d)/'result.xlsx'
            fixture(src)
            original = hashlib.sha256(src.read_bytes()).hexdigest()
            with patch('openpdf_xlsx_tables_v514._merged_regions',
                       side_effect=ValueError('Frontieres ambigues')):
                with self.assertRaisesRegex(ValueError, 'ambigues'):
                    convert_office(src, dst, table_mode=True)
            self.assertFalse(dst.exists())
            self.assertEqual(hashlib.sha256(src.read_bytes()).hexdigest(), original)


if __name__ == '__main__':
    unittest.main()
