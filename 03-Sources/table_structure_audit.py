"""Office table structure checks; not a PDF table reconstruction. Author: Abdelhalim Talbi."""
from pathlib import Path
from collections import Counter


def inspect_office_tables(path):
    ext = Path(path).suffix.lower()
    tables = []
    if ext == '.docx':
        from docx import Document
        doc = Document(str(path))
        for index, table in enumerate(doc.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            widths = [len(row) for row in rows]
            # python-docx expands merged cells into repeated references; detect adjacent duplicates
            merged = sum(sum(row.cells[i]._tc is row.cells[i-1]._tc for i in range(1, len(row.cells))) for row in table.rows)
            tables.append({'id': f'table-{index+1}', 'rows': len(rows), 'columns': max(widths, default=0),
                           'header': rows[0] if rows else [], 'cells': rows,
                           'merged_cell_candidates': merged, 'source': 'docx_native'})
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=False, data_only=False)
        try:
            for ws in wb.worksheets:
                for table_name, table in ws.tables.items():
                    ref = table if isinstance(table, str) else table.ref
                    region = ws[ref]
                    rows = [[str(c.value) if c.value is not None else '' for c in row] for row in region]
                    tables.append({'id': f'{ws.title}:{table_name}', 'rows': len(rows),
                                   'columns': len(rows[0]) if rows else 0, 'header': rows[0] if rows else [],
                                   'cells': rows, 'merged_cell_candidates': 0,
                                   'source': 'xlsx_defined_table'})
        finally:
            wb.close()
    return tables


def compare_table_structures(source_path, target_path):
    """Compare native Office table structures by position; report uncertainties explicitly."""
    source = inspect_office_tables(source_path)
    target = inspect_office_tables(target_path)
    findings = []
    for index in range(max(len(source), len(target))):
        if index >= len(source):
            findings.append({'kind': 'added_table', 'index': index, 'severity': 'review'})
            continue
        if index >= len(target):
            findings.append({'kind': 'missing_table', 'index': index, 'severity': 'review'})
            continue
        a, b = source[index], target[index]
        for field in ('rows', 'columns'):
            if a[field] != b[field]:
                findings.append({'kind': f'{field}_changed', 'index': index, 'source': a[field], 'target': b[field], 'severity': 'review'})
        if a['header'] != b['header']:
            if Counter(a['header']) == Counter(b['header']):
                kind = 'header_columns_reordered'
            else:
                kind = 'header_content_changed'
            findings.append({'kind': kind, 'index': index, 'source': a['header'], 'target': b['header'], 'severity': 'review'})
        if a['cells'] != b['cells']:
            left = Counter(v for row in a['cells'] for v in row if v)
            right = Counter(v for row in b['cells'] for v in row if v)
            missing = list((left - right).elements())
            added = list((right - left).elements())
            if missing or added:
                findings.append({'kind': 'cell_values_changed', 'index': index, 'missing': missing[:100], 'added': added[:100], 'severity': 'review'})
        if a['merged_cell_candidates'] != b['merged_cell_candidates']:
            findings.append({'kind': 'merged_cells_changed', 'index': index, 'source': a['merged_cell_candidates'], 'target': b['merged_cell_candidates'], 'severity': 'review'})
    return {'source_table_count': len(source), 'target_table_count': len(target), 'findings': findings,
            'limitations': 'Office native tables only; positional matching, no PDF table reconstruction. XLSX only named tables.'}
