"""V521: phantom XLSX table/merge structures must fail the real PDF corpus gate.

Creates actual PDF -> editable XLSX conversions, then injects only blank
structural artifacts. Ground truth is synthetic, known at generation, never
represented as human-reviewed. Author: Abdelhalim Talbi.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
from openpyxl import load_workbook
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_corpus_gate_v519 import SCHEMA as CORPUS_SCHEMA, evaluate_corpus
from openpdf_xlsx_ground_truth_v517 import SCHEMA as TRUTH_SCHEMA, evaluate_ground_truth
from test_xlsx_tables_v514 import fixture


class PhantomStructureTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.cases = []
        for n in (1, 2):
            pdf = self.root / f'doc{n}.pdf'
            xlsx = self.root / f'doc{n}.xlsx'
            annotation = self.root / f'doc{n}.json'
            fixture(pdf, injection=(n == 2))
            convert_office(pdf, xlsx, table_mode=True)
            annotation.write_text(json.dumps({
                'schema': TRUTH_SCHEMA,
                'source_sha256': hashlib.sha256(pdf.read_bytes()).hexdigest(),
                'origin': 'synthetic_known_at_generation',
                'tables': [{'page': 1, 'index': 1,
                            'cells': [['Volume', 'Capacite'], ['SAN-A', '20 To'],
                                      ['SAN-B', '=1+1' if n == 2 else '50 To']], 'merges': []}],
            }), encoding='utf-8')
            self.cases.append((pdf, xlsx, annotation))
        self.manifest = self.root / 'corpus.json'
        self.manifest.write_text(json.dumps({
            'schema': CORPUS_SCHEMA,
            'cases': [{'id': f'doc-{n}', 'pdf': pdf.name, 'xlsx': xlsx.name,
                       'annotations': annotation.name}
                      for n, (pdf, xlsx, annotation) in enumerate(self.cases, 1)],
            'thresholds': {'min_nonempty_cell_accuracy': 1.0,
                           'min_table_recall': 1.0, 'min_merge_recall': 1.0,
                           'max_extra_nonempty_cells': 0, 'min_human_reviewed': 0},
        }), encoding='utf-8')

    def _edit(self, case, *, phantom_table=False, phantom_merge=False):
        xlsx = self.cases[case][1]
        workbook = load_workbook(xlsx)
        try:
            if phantom_table:
                sheet = workbook.create_sheet('P0001-T002')
                sheet['A1'], sheet['B1'] = 'Page PDF', 1
                sheet['C1'], sheet['D1'] = 'Tableau', 2
            if phantom_merge:
                workbook.worksheets[1].merge_cells('C7:D7')
            workbook.save(xlsx)
        finally:
            workbook.close()

    def test_clean_real_conversions_pass(self):
        report = evaluate_corpus(self.manifest)
        self.assertEqual(report['status'], 'PASS')
        self.assertEqual(report['failed_thresholds'], [])
        self.assertEqual(report['human_reviewed_cases'], 0)
        self.assertEqual(report['nonempty_cell_accuracy'], 1.0)

    def test_phantom_empty_table_fails_even_with_perfect_recall(self):
        self._edit(0, phantom_table=True)
        individual = evaluate_ground_truth(*self.cases[0])
        self.assertEqual(individual['extra_nonempty_cells'], 0)
        self.assertEqual(individual['table_recall'], 1.0)
        self.assertEqual(individual['table_precision'], 0.5)
        report = evaluate_corpus(self.manifest)
        self.assertEqual(report['status'], 'FAIL')
        self.assertEqual(report['nonempty_cell_accuracy'], 1.0)
        self.assertIn('unexpected_tables', report['failed_thresholds'])
        self.assertNotIn('max_extra_nonempty_cells', report['failed_thresholds'])

    def test_phantom_empty_merge_fails_even_with_perfect_recall(self):
        self._edit(1, phantom_merge=True)
        individual = evaluate_ground_truth(*self.cases[1])
        self.assertEqual(individual['extra_nonempty_cells'], 0)
        self.assertEqual(individual['merge_recall'], 1.0)
        self.assertEqual(individual['merge_precision'], 0.0)
        report = evaluate_corpus(self.manifest)
        self.assertEqual(report['status'], 'FAIL')
        self.assertIn('unexpected_merges', report['failed_thresholds'])

    def test_both_phantoms_fail_without_leaking_source_cells(self):
        self._edit(0, phantom_table=True)
        self._edit(1, phantom_merge=True)
        report = evaluate_corpus(self.manifest)
        self.assertEqual(report['status'], 'FAIL')
        self.assertEqual(set(report['failed_thresholds']),
                         {'unexpected_tables', 'unexpected_merges'})
        self.assertEqual(report['extra_nonempty_cells'], 0)
        self.assertNotIn('SAN-A', json.dumps(report))
        self.assertNotIn('50 To', json.dumps(report))

    def test_existing_corpus_schema_and_thresholds_unchanged(self):
        manifest = json.loads(self.manifest.read_text(encoding='utf-8'))
        self.assertEqual(manifest['schema'], 'openpdf.xlsx.corpus.v1')
        self.assertEqual(set(manifest['thresholds']), {
            'min_nonempty_cell_accuracy', 'min_table_recall',
            'min_merge_recall', 'max_extra_nonempty_cells',
            'min_human_reviewed',
        })


if __name__ == '__main__':
    unittest.main()
