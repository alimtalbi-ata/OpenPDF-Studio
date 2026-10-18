import unittest,tempfile
from pathlib import Path
import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from openpdf_fidelity_v352 import measure_text_coverage
from openpdf_conversion_audit_v346 import audit_conversion
class FidelityTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.root=Path(self.temp.name)
        self.source=self.root/'in.pdf'
        with fitz.open() as d:
            p=d.new_page();p.insert_text((50,60),'alpha beta gamma');d.save(self.source)
    def tearDown(self):self.temp.cleanup()
    def test_docx_full_coverage(self):
        p=self.root/'a.docx';d=Document();d.add_paragraph('alpha beta gamma');d.save(p)
        self.assertEqual(measure_text_coverage(self.source,p).coverage,1.0)
        self.assertTrue(any('Couverture lexicale' in x for x in audit_conversion(self.source,p).warnings))
    def test_docx_partial_coverage(self):
        p=self.root/'a.docx';d=Document();d.add_paragraph('alpha');d.save(p)
        self.assertLess(measure_text_coverage(self.source,p).coverage,0.8)
        self.assertTrue(any('faible' in x for x in audit_conversion(self.source,p).warnings))
    def test_xlsx(self):
        p=self.root/'a.xlsx';w=Workbook();w.active['A1']='alpha beta gamma';w.save(p)
        self.assertEqual(measure_text_coverage(self.source,p).coverage,1.0)
    def test_pptx(self):
        p=self.root/'a.pptx';prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6]);box=slide.shapes.add_textbox(0,0,1000000,1000000);box.text='alpha beta gamma';prs.save(p)
        self.assertEqual(measure_text_coverage(self.source,p).coverage,1.0)
    def test_bad_format(self):
        with self.assertRaises(ValueError):measure_text_coverage(self.source,self.root/'x.txt')
if __name__=='__main__':unittest.main()
