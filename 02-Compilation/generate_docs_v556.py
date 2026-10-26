"""Generate version-specific V556 verification-stability documentation. Author: Abdelhalim Talbi."""
from __future__ import annotations
import os
import re
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
AUTHOR = "Abdelhalim Talbi"
TITLE = "OpenPDF Studio V51.4 / V556 - Stable cumulative archive verification"
DOCX_NAME = "OpenPDF-Studio-V51.4-V556-Archive-Stability-Guide.docx"
PDF_NAME = "OpenPDF-Studio-V51.4-V556-Archive-Stability-Guide.pdf"
PNG_NAME = "Architecture-V556-Archive-Stability.png"
SVG_NAME = "Architecture-V556-Archive-Stability.svg"


def build_context(*, commit_sha=None, pr_number=None, workflow_run_id=None):
    sha = commit_sha
    if sha is None:
        sha = os.environ.get("OPENPDF_HEAD_SHA") or os.environ.get("GITHUB_SHA", "")
    if not isinstance(sha, str) or not re.fullmatch(r"[0-9a-f]{40}", sha):
        sha = "NOT-VERIFIED"
    pr = str(pr_number if pr_number is not None else os.environ.get("OPENPDF_PULL_REQUEST_NUMBER", "")).strip()
    run_id = str(workflow_run_id if workflow_run_id is not None else os.environ.get("OPENPDF_RUN_ID", "")).strip()
    return {"commit_sha": sha, "pr_number": pr or "not supplied", "workflow_run_id": run_id or "not supplied"}


def build_sections(context):
    pr_url = ("https://github.com/alimtalbi-ata/OpenPDF-Studio/pull/" + context["pr_number"]
              if context["pr_number"].isdigit() else "not available")
    run_url = ("https://github.com/alimtalbi-ata/OpenPDF-Studio/actions/runs/" + context["workflow_run_id"]
               if context["workflow_run_id"].isdigit() else "not available")
    sha_verified = context["commit_sha"] != "NOT-VERIFIED"
    return [
        ("1. Problème corrigé",
         "Le vérificateur cumulatif vérifiait les membres ZIP, l'inventaire et les SHA-256, puis calculait l'empreinte du chemin en rouvrant le fichier après la fermeture du lecteur ZIP. Un remplacement du chemin pendant la vérification pouvait donc dissocier le contenu réellement contrôlé de l'empreinte inscrite au reçu."),
        ("2. Nouvelle chaîne de vérification",
         "V556 ouvre le ZIP une fois et remet ce même descripteur au lecteur ZIP. Il calcule l'empreinte initiale et finale sur ce descripteur, puis compare les métadonnées de fichier ouvert avec le chemin publié. Les liens symboliques, dossiers et chemins devenus différents sont refusés. Une mutation de contenu à longueur identique et un remplacement du chemin sont reproduits dans des tests dédiés et doivent échouer en mode fail-closed."),
        ("3. Intégrité du reçu",
         "Le reçu retourne l'empreinte du même flux qui a été vérifié, et non une nouvelle ouverture du chemin. La taille reportée provient du descripteur ouvert. Le contrôle structurel V555, les seuils d'amplification V554, les contrôles de chemin NFKC/casefold, l'inventaire et les sommes SHA-256 restent en place."),
        ("4. Tests et qualification",
         "Les tests V556 couvrent une archive synthétique valide, la mutation en place pendant le précontrôle, le remplacement du chemin et le refus des entrées non régulières. Ce sont des tests fonctionnels du vérificateur. Ils ne constituent pas une qualification d'exécutable Windows. La qualification native doit rester séparée : portable, installateur, démarrage, tests fonctionnels, OCR runtime, installation/désinstallation et génération/validation du ZIP cumulatif sur le SHA exact."),
        ("5. Limites explicites",
         "Les contrôles de stabilité détectent les changements observés pendant le cycle de vérification; ils ne remplacent pas un verrouillage du système de fichiers contre un acteur privilégié contrôlant les écritures concurrentes. Le SHA-256 confirme l'intégrité par rapport à la valeur calculée mais n'authentifie pas l'éditeur. Aucune signature Authenticode, revue graphique humaine, certification OCR métier ou fidélité Office certifiée n'est revendiquée."),
        ("6. Prochaine priorité",
         "Après validation CI sur le SHA exact, poursuivre la consolidation cumulative V554-V556. Ensuite, établir un corpus métier indépendant par format et par page pour mesurer rappel/précision, ordre des objets, géométrie et fidélité visuelle PDF vers DOCX/XLSX/PPTX."),
        ("Traçabilité",
         "Commit: " + context["commit_sha"] + "\nPR: " + pr_url + "\nWorkflow: " + run_url + "\n" +
         ("Le SHA a le format attendu; les reçus CI doivent établir la qualification de ce commit." if sha_verified
          else "Le SHA source est absent/invalide; ce document est explicitement non qualifié.")),
    ]


def svg_diagram():
    return """<svg xmlns="http://www.w3.org/2000/svg" width="1120" height="350" viewBox="0 0 1120 350" role="img" aria-labelledby="title desc">
<title id="title">OpenPDF Studio V556 stable archive verification</title>
<desc id="desc">Generated architecture diagram, not an application screenshot.</desc>
<defs><marker id="arrow" markerWidth="9" markerHeight="9" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#24566b"/></marker></defs>
<rect x="18" y="74" width="190" height="112" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="113" y="106" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#153747">Archive path</text>
<text x="113" y="130" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">lstat / regular file</text>
<text x="113" y="152" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">no symlink</text>
<path d="M208 130 L239 130" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="246" y="58" width="208" height="144" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="350" y="89" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#153747">One open descriptor</text>
<text x="350" y="114" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">identity + initial SHA-256</text>
<text x="350" y="137" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">ZIP parser reads same fd</text>
<text x="350" y="160" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">inventory + member hashes</text>
<path d="M454 130 L486 130" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="493" y="74" width="185" height="112" rx="12" fill="#fff3df" stroke="#a75c19" stroke-width="2"/>
<text x="585" y="105" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#75410d">Recheck binding</text>
<text x="585" y="130" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#75410d">path + descriptor identity</text>
<text x="585" y="152" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#75410d">final SHA-256</text>
<path d="M678 130 L710 130" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="717" y="74" width="174" height="112" rx="12" fill="#e7f5e9" stroke="#2b7545" stroke-width="2"/>
<text x="804" y="105" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#1d5833">Stable?</text>
<text x="804" y="130" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#1d5833">same identity + digest</text>
<text x="804" y="152" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#1d5833">receipt allowed</text>
<path d="M891 130 L922 130" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="929" y="74" width="173" height="112" rx="12" fill="#fbe7e7" stroke="#a33333" stroke-width="2"/>
<text x="1015" y="105" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#8b2525">Mismatch?</text>
<text x="1015" y="130" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#8b2525">fail closed</text>
<text x="1015" y="152" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#8b2525">no trusted receipt</text>
<rect x="246" y="239" width="856" height="72" rx="8" fill="#f4f0ff" stroke="#6750a4" stroke-width="1.5"/>
<text x="674" y="266" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#3e2d73">V555 structural gate + V554 expansion limits + SHA-256 inventory</text>
<text x="674" y="289" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#3e2d73">Architecture diagram only — not a real application screenshot or publisher signature.</text>
</svg>"""


def draw_diagram_png(path):
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1120, 350), "white")
    draw = ImageDraw.Draw(image)
    font_path = next((p for p in (
        "C:/Windows/Fonts/segoeui.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ) if Path(p).is_file()), None)
    font = ImageFont.truetype(font_path, 15) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(font_path, 12) if font_path else ImageFont.load_default()
    boxes = [
        (18, 74, 208, 186, "Archive path", "lstat / regular file", "no symlink", "#e8f3f5", "#24566b"),
        (246, 58, 454, 202, "One open descriptor", "identity + initial SHA-256", "ZIP parser / member hashes", "#e8f3f5", "#24566b"),
        (493, 74, 678, 186, "Recheck binding", "path + descriptor identity", "final SHA-256", "#fff3df", "#a75c19"),
        (717, 74, 891, 186, "Stable?", "same identity + digest", "receipt allowed", "#e7f5e9", "#2b7545"),
        (929, 74, 1102, 186, "Mismatch?", "fail closed", "no trusted receipt", "#fbe7e7", "#a33333"),
    ]
    for box in boxes:
        x0, y0, x1, y1, title, sub1, sub2, fill, outline = box
        draw.rounded_rectangle((x0, y0, x1, y1), radius=12, fill=fill, outline=outline, width=2)
        draw.text((x0 + 8, y0 + 25), title, font=small, fill="#153747")
        draw.text((x0 + 8, y0 + 51), sub1, font=small, fill="#153747")
        draw.text((x0 + 8, y0 + 74), sub2, font=small, fill="#153747")
    for x0, x1 in ((208, 246), (454, 493), (678, 717), (891, 929)):
        y = 130
        draw.line((x0 + 2, y, x1 - 6, y), fill="#24566b", width=3)
        draw.polygon([(x1 - 6, y), (x1 - 16, y - 6), (x1 - 16, y + 6)], fill="#24566b")
    draw.rounded_rectangle((246, 239, 1102, 311), radius=8, fill="#f4f0ff", outline="#6750a4", width=2)
    draw.text((263, 257), "V555 structural gate + V554 expansion limits + SHA-256 inventory", font=small, fill="#3e2d73")
    draw.text((263, 281), "Architecture diagram only — not an application screenshot or publisher signature", font=small, fill="#3e2d73")
    image.save(path, format="PNG", optimize=True)


def generate(root=ROOT, *, commit_sha=None, pr_number=None, workflow_run_id=None):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PdfImage, Table, TableStyle, KeepTogether

    root = Path(root)
    docs = root / "05-Documentation"
    docs.mkdir(parents=True, exist_ok=True)
    context = build_context(commit_sha=commit_sha, pr_number=pr_number, workflow_run_id=workflow_run_id)
    sections = build_sections(context)
    svg_path, png_path = docs / SVG_NAME, docs / PNG_NAME
    docx_path, pdf_path = docs / DOCX_NAME, docs / PDF_NAME
    svg_path.write_text(svg_diagram(), encoding="utf-8")
    draw_diagram_png(png_path)

    document = Document()
    document.core_properties.author = AUTHOR
    document.core_properties.title = TITLE
    document.core_properties.subject = "V556 stable archive verification and evidence levels"
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading(TITLE, level=0)
    p = document.add_paragraph()
    p.add_run("Auteur : ").bold = True
    p.add_run(AUTHOR + "   ")
    p.add_run("Version applicative : ").bold = True
    p.add_run("51.4   ")
    p.add_run("Jalon : ").bold = True
    p.add_run("V556")
    document.add_paragraph("Commit source : " + context["commit_sha"])
    document.add_paragraph("PR : " + context["pr_number"] + " | Workflow : " + context["workflow_run_id"])
    if context["commit_sha"] == "NOT-VERIFIED":
        note = document.add_paragraph("SHA absent ou invalide : guide non qualifié.")
        for run in note.runs:
            run.bold = True
    document.add_picture(str(png_path), width=Inches(6.35))
    caption = document.add_paragraph("Figure 1. Schéma explicatif généré, pas une capture réelle de l'interface.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, body in sections:
        document.add_heading(heading, level=1)
        for line in body.split("\n"):
            if line.strip():
                document.add_paragraph(line)
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("Niveau", "Contrôle", "Portée")):
        cell.text = text
    for row in (
        ("Preflight", "Lstat + descripteur ouvert", "Lier le chemin au fichier vérifié"),
        ("Empreintes", "SHA-256 avant/après", "Même flux de lecture"),
        ("Tests", "Mutation + remplacement", "Fonctionnel, fixtures synthétiques"),
        ("Authenticité", "Non revendiquée", "SHA-256 n'authentifie pas l'éditeur"),
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    document.save(docx_path)

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="V556Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12, spaceAfter=6))
    styles.add(ParagraphStyle(name="V556Small", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, leading=9, spaceAfter=5))
    story = [
        Paragraph(escape(TITLE), styles["Title"]),
        Spacer(1, 3 * mm),
        Paragraph(escape("Auteur : " + AUTHOR + " | Version applicative : 51.4 | Jalon : V556"), styles["V556Body"]),
        Paragraph(escape("Commit : " + context["commit_sha"] + " | PR : " + context["pr_number"] + " | Workflow : " + context["workflow_run_id"]), styles["V556Small"]),
        PdfImage(str(png_path), width=174 * mm, height=174 * mm * 350 / 1120),
        Paragraph("Schéma d'architecture généré — pas une capture de l'application.", styles["Italic"]),
        Spacer(1, 3 * mm),
    ]
    for heading, body in sections:
        block = [Paragraph(escape(heading), styles["Heading2"])]
        block.extend(Paragraph(escape(line), styles["V556Body"]) for line in body.split("\n") if line.strip())
        if heading == "Traçabilité":
            story.append(KeepTogether(block))
        else:
            story.extend(block)
    rows = [
        ["Niveau", "Contrôle", "Portée"],
        ["Preflight", "Lstat + descripteur", "Chemin lié au fichier vérifié"],
        ["Empreintes", "SHA-256 avant/après", "Même flux de lecture"],
        ["Tests", "Mutation + remplacement", "Fixtures fonctionnelles"],
        ["Authenticité", "Non revendiquée", "SHA-256 n'authentifie pas l'éditeur"],
    ]
    table = Table(rows, colWidths=[35 * mm, 53 * mm, 86 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24566b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9db3bd")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([Spacer(1, 3 * mm), KeepTogether([Paragraph("Niveaux de preuve", styles["Heading2"]), table])])
    SimpleDocTemplate(str(pdf_path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title=TITLE, author=AUTHOR).build(story)

    for path in (docx_path, pdf_path, png_path, svg_path):
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError("Generated file is missing/empty: " + str(path))
    return {"context": context, "docx": str(docx_path), "pdf": str(pdf_path), "png": str(png_path), "svg": str(svg_path)}


if __name__ == "__main__":
    for key, value in generate().items():
        print(str(key) + ": " + str(value))
