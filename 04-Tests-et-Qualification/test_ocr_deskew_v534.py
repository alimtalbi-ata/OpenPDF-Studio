"""V534: real tilted scanned PDF to editable XLSX, explicit raster deskew.
Synthetic regression only; not a claim of recognition accuracy.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations
import hashlib
import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import fitz
from PIL import Image
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
from openpdf_ocr_office_pipeline_v511 import convert_scanned_office
from openpdf_ocr_language_preflight_v486 import preflight_ocr_languages
from openpdf_ocr_service_v347 import searchable_pdf
import openpdf_operations as operations

TABLE = (('ID','STOCK','PRICE'),('A01','102','55'),
         ('A02','103','80'),('A03','104','74'))


def english_available():
    try:
        preflight_ocr_languages('eng')
        return True
    except (RuntimeError, ValueError, OSError):
        return False


def make_tilted_scan(path, *, native_first=False):
    with fitz.open() as doc:
        page = doc.new_page(width=600, height=400)
        for row, values in enumerate(TABLE):
            for col, value in enumerate(values):
                page.insert_text((55 + col * 160, 100 + row * 50), value, fontsize=17)
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
        tilted = image.rotate(8, resample=Image.Resampling.BICUBIC,
                              expand=False, fillcolor='white')
        buffer = io.BytesIO()
        tilted.save(buffer, format='PNG')
    with fitz.open() as doc:
        if native_first:
            doc.new_page(width=600, height=400).insert_text(
                (50, 60), 'Native text', fontsize=15)
        doc.new_page(width=600, height=400).insert_image(
            fitz.Rect(0, 0, 600, 400), stream=buffer.getvalue())
        doc.save(path)


class OCRDeskewV534(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)
        self.source = self.root / 'tilted.pdf'
        self.destination = self.root / 'converted.xlsx'

    def test_invalid_angles_fail_before_publication(self):
        make_tilted_scan(self.source)
        for angle in ('8', True, None, float('nan'), float('inf'), 12.1, -12.1):
            with self.subTest(angle=angle), self.assertRaisesRegex(ValueError, 'Angle'):
                convert_scanned_office(self.source, self.destination,
                                       language='eng', deskew_angle_degrees=angle)
            self.assertFalse(self.destination.exists())

    def test_dispatcher_refuses_deskew_without_office_ocr(self):
        make_tilted_scan(self.source)
        for operation, options in [('xlsx', {'ocr_before_office': False}),
                                   ('compression', {'ocr_before_office': True})]:
            with self.subTest(operation=operation), self.assertRaisesRegex(ValueError, 'Angle'):
                operations.run_operation(operation, (str(self.source),),
                    str(self.destination), deskew_angle_degrees=-8, **options)
        self.assertFalse(self.destination.exists())

    def test_dispatcher_passes_explicit_angle(self):
        make_tilted_scan(self.source)
        seen = []
        def fake(source, destination, **kwargs):
            seen.append(kwargs)
            return 'routed'
        with patch('openpdf_ocr_office_pipeline_v511.convert_scanned_office',
                   side_effect=fake), patch.object(operations, 'write_event'):
            result = operations.run_operation('xlsx', (str(self.source),),
                str(self.destination), ocr_before_office=True, table_mode=True,
                language='eng', deskew_angle_degrees=-8)
        self.assertEqual(result, 'routed')
        self.assertEqual(seen, [{'language':'eng', 'table_mode':True,
                                 'deskew_angle_degrees':-8.0}])

    def test_gui_option_is_explicit(self):
        source = (ROOT / '03-Sources/openpdf_desktop.py').read_text(encoding='utf-8')
        self.assertIn("self.deskew_angle=tk.StringVar(value='0')", source)
        self.assertIn("show += ('Redresser scan (degrés)',)", source)
        self.assertIn("kwargs['deskew_angle_degrees']=angle", source)

    @unittest.skipUnless(english_available(), 'English offline Tesseract unavailable')
    def test_real_tilted_scan_recovers_editable_table(self):
        make_tilted_scan(self.source)
        original = hashlib.sha256(self.source.read_bytes()).hexdigest()
        with self.assertRaisesRegex(ValueError, 'Aucun tableau detecte'):
            convert_scanned_office(self.source, self.destination,
                                   language='eng', table_mode=True)
        self.assertFalse(self.destination.exists())
        receipt = convert_scanned_office(self.source, self.destination,
            language='eng', table_mode=True, deskew_angle_degrees=-8)
        self.assertEqual(receipt.table_cells_by_page, ((1, 12, 12),))
        self.assertEqual(receipt.table_worksheets, 1)
        self.assertEqual(receipt.deskew_angle_degrees, -8)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), original)
        self.assertEqual(hashlib.sha256(self.destination.read_bytes()).hexdigest(), receipt.sha256)
        workbook = load_workbook(self.destination)
        try:
            sheet = workbook['P0001-T001']
            self.assertEqual(tuple(tuple(sheet.cell(r, c).value for c in range(1, 4))
                                   for r in range(3, 7)), TABLE)
        finally:
            workbook.close()

    @unittest.skipUnless(english_available(), 'English offline Tesseract unavailable')
    def test_native_page_preserved_when_only_scanned_page_corrected(self):
        make_tilted_scan(self.source, native_first=True)
        receipt = convert_scanned_office(self.source, self.destination,
                                         language='eng', deskew_angle_degrees=-8)
        self.assertEqual(receipt.scanned_pages, (2,))
        self.assertEqual(receipt.ocr_pages, 1)
        self.assertTrue(self.destination.is_file())

    @unittest.skipUnless(english_available(), 'English offline Tesseract unavailable')
    def test_forced_ocr_on_native_text_refuses_rotation(self):
        with fitz.open() as doc:
            doc.new_page().insert_text((40, 80), 'Native content', fontsize=20)
            doc.save(self.source)
        output = self.root / 'forced.pdf'
        with self.assertRaisesRegex(ValueError, 'texte natif'):
            searchable_pdf(self.source, output, language='eng', skip_text=False,
                           deskew_angle_degrees=5)
        self.assertFalse(output.exists())


if __name__ == '__main__':
    unittest.main()
