"""V535 conservative auto-deskew regression with real synthetic OCR evidence.

These tests do not certify OCR accuracy or user-interface interaction.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations
import hashlib
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import fitz
from PIL import Image
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
sys.path.insert(0, str(ROOT / '04-Tests-et-Qualification'))
from openpdf_ocr_autodeskew_v535 import estimate_deskew
from openpdf_ocr_language_preflight_v486 import preflight_ocr_languages
from openpdf_ocr_office_pipeline_v511 import convert_scanned_office
from openpdf_ocr_service_v347 import searchable_pdf
from test_ocr_deskew_v534 import make_tilted_scan, TABLE
import openpdf_operations as operations


def english_available():
    try:
        preflight_ocr_languages('eng')
        return True
    except (RuntimeError, ValueError, OSError):
        return False


class AutoDeskewV535(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / 'scan.pdf'
        self.destination = self.root / 'result.xlsx'

    def image_from_scan(self, angle):
        make_tilted_scan(self.source, angle=angle)
        with fitz.open(self.source) as doc:
            pix = doc[0].get_pixmap(dpi=150, alpha=False)
            return Image.frombytes('RGB', (pix.width, pix.height), pix.samples)

    def test_projection_recovers_known_eight_degree_skew(self):
        decision = estimate_deskew(self.image_from_scan(8))
        self.assertTrue(decision.applied, decision)
        self.assertAlmostEqual(decision.angle_degrees, -8, delta=0.75)
        self.assertGreaterEqual(decision.projection_gain, 1.30)

    def test_already_horizontal_is_not_rotated(self):
        decision = estimate_deskew(self.image_from_scan(0))
        self.assertFalse(decision.applied, decision)
        self.assertEqual(decision.angle_degrees, 0.0)

    def test_blank_raster_is_not_rotated(self):
        decision = estimate_deskew(Image.new('RGB', (800, 500), 'white'))
        self.assertFalse(decision.applied)
        self.assertEqual(decision.reason, 'insufficient_text_like_pixels')

    def test_small_raster_is_not_rotated(self):
        decision = estimate_deskew(Image.new('RGB', (90, 90), 'black'))
        self.assertFalse(decision.applied)
        self.assertEqual(decision.reason, 'raster_too_small')

    def test_manual_auto_conflict_rejected_without_output(self):
        self.image_from_scan(8)
        with self.assertRaisesRegex(ValueError, 'automatique'):
            convert_scanned_office(self.source, self.destination,
                language='eng', auto_deskew=True, deskew_angle_degrees=-8)
        self.assertFalse(self.destination.exists())

    def test_invalid_boolean_rejected_without_output(self):
        self.image_from_scan(8)
        for bad in ('true', 1, None):
            with self.subTest(bad=bad), self.assertRaisesRegex(ValueError, 'automatique'):
                convert_scanned_office(self.source, self.destination,
                    language='eng', auto_deskew=bad)
            self.assertFalse(self.destination.exists())

    def test_dispatcher_rejects_unsupported_auto_deskew(self):
        self.image_from_scan(8)
        for operation, flags in [('compression', {'ocr_before_office': True}),
                                 ('xlsx', {'ocr_before_office': False}),
                                 ('xlsx', {'ocr_before_office': True, 'deskew_angle_degrees': 4})]:
            with self.subTest(operation=operation, flags=flags), self.assertRaisesRegex(ValueError, 'automatique'):
                operations.run_operation(operation, (str(self.source),),
                                         str(self.destination), auto_deskew=True, **flags)
        self.assertFalse(self.destination.exists())

    def test_dispatcher_routes_opt_in_to_pipeline(self):
        self.image_from_scan(8)
        with patch('openpdf_ocr_office_pipeline_v511.convert_scanned_office', return_value='routed') as mocked, \
             patch.object(operations, 'write_event'):
            result = operations.run_operation('xlsx', (str(self.source),), str(self.destination),
                                              ocr_before_office=True, table_mode=True,
                                              language='eng', auto_deskew=True)
        self.assertEqual(result, 'routed')
        self.assertEqual(mocked.call_args.kwargs,
                         {'language': 'eng', 'table_mode': True, 'auto_deskew': True})

    def test_gui_has_opt_in_control(self):
        source = (ROOT / '03-Sources/openpdf_desktop.py').read_text(encoding='utf-8')
        self.assertIn("self.auto_deskew=tk.BooleanVar(value=False)", source)
        self.assertIn("kwargs['auto_deskew']=True", source)
        self.assertIn('angle manuel OU redressement automatique', source)

    @unittest.skipUnless(english_available(), 'English OCR model unavailable')
    def test_real_auto_deskew_ocr_to_editable_xlsx(self):
        self.image_from_scan(8)
        original = hashlib.sha256(self.source.read_bytes()).hexdigest()
        result = convert_scanned_office(self.source, self.destination, language='eng',
                                        table_mode=True, auto_deskew=True)
        self.assertTrue(result.auto_deskew)
        self.assertEqual(result.deskew_pages[0][0], 1)
        self.assertAlmostEqual(result.deskew_pages[0][1], -8, delta=0.75)
        self.assertEqual(result.deskew_pages[0][2], 'projection_peak')
        self.assertEqual(result.table_cells_by_page, ((1, 12, 12),))
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), original)
        self.assertEqual(hashlib.sha256(self.destination.read_bytes()).hexdigest(), result.sha256)
        self.assertEqual(result.to_dict()['deskew_pages'][0]['page'], 1)
        workbook = load_workbook(self.destination)
        try:
            sheet = workbook['P0001-T001']
            self.assertEqual(tuple(tuple(sheet.cell(r, c).value for c in range(1, 4))
                                   for r in range(3, 7)), TABLE)
        finally:
            workbook.close()

    @unittest.skipUnless(english_available(), 'English OCR model unavailable')
    def test_mixed_pdf_native_page_is_not_rotated(self):
        make_tilted_scan(self.source, angle=8, native_first=True)
        result = convert_scanned_office(self.source, self.destination, language='eng',
                                        auto_deskew=True)
        self.assertEqual(result.scanned_pages, (2,))
        self.assertEqual(tuple(page for page, _, _ in result.deskew_pages), (2,))
        self.assertEqual(result.deskew_pages[0][2], 'projection_peak')

    @unittest.skipUnless(english_available(), 'English OCR model unavailable')
    def test_auto_deskew_refuses_native_text_force_mode(self):
        with fitz.open() as pdf:
            pdf.new_page().insert_text((40, 80), 'Native text', fontsize=18)
            pdf.save(self.source)
        destination = self.root / 'ocr.pdf'
        with self.assertRaisesRegex(ValueError, 'texte natif'):
            searchable_pdf(self.source, destination, language='eng',
                           skip_text=False, auto_deskew=True)
        self.assertFalse(destination.exists())


if __name__ == '__main__':
    unittest.main()
