"""Real PDF-to-Office integration: structural and textual assertions.
Author: Abdelhalim Talbi.
"""
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path
import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'03-Sources'))
from openpdf_operations import run_operation

class OfficeFunctionalTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.source=self.root/'source.pdf'
        with fitz.open() as doc:
            for label in ('FIRST-ALPHA','SECOND-BETA'):
                page=doc.new_page()
                page.insert_text((72,100),label,fontsize=18)
            doc.save(self.source)
    def test_docx_editable_text_and_digest(self):
        output=self.root/'output.docx'
        result=run_operation('docx',[str(self.source)],str(output))
        paragraphs=' '.join(p.text for p in Document(output).paragraphs)
        self.assertIn('FIRST-ALPHA',paragraphs)
        self.assertIn('SECOND-BETA',paragraphs)
        self.assertEqual(result.sha256,hashlib.sha256(output.read_bytes()).hexdigest())
    def test_xlsx_text_rows(self):
        output=self.root/'output.xlsx'
        result=run_operation('xlsx',[str(self.source)],str(output))
        wb=load_workbook(output,read_only=True,data_only=False)
        try:
            rows=list(wb.active.values)
            flattened=' '.join(str(cell) for row in rows for cell in row if cell is not None)
            self.assertIn('FIRST-ALPHA',flattened)
            self.assertIn('SECOND-BETA',flattened)
        finally:wb.close()
        self.assertEqual(result.output_format,'xlsx')
    def test_pptx_slide_count_and_text(self):
        output=self.root/'output.pptx'
        result=run_operation('pptx',[str(self.source)],str(output))
        presentation=Presentation(output)
        self.assertEqual(len(presentation.slides),2)
        slide_text=[' '.join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in presentation.slides]
        self.assertIn('FIRST-ALPHA',slide_text[0])
        self.assertIn('SECOND-BETA',slide_text[1])
        self.assertEqual(result.output_format,'pptx')
    def test_existing_office_output_preserved(self):
        output=self.root/'existing.docx';output.write_bytes(b'keep')
        with self.assertRaises(FileExistsError):
            run_operation('docx',[str(self.source)],str(output))
        self.assertEqual(output.read_bytes(),b'keep')
if __name__=='__main__':unittest.main()
