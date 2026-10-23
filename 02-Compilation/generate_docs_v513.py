"""Generate version-specific V51.3 Word/PDF and architecture diagram.
Author: Abdelhalim Talbi. No simulated Windows evidence.
"""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PDFImage, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'05-Documentation'
SECTIONS=[
 ('Objectif produit','Le portable et l installateur Windows doivent embarquer les modeles OCR eng et fra. Le moteur refuse un modele altere et ne telecharge jamais de donnees pendant le traitement PDF.'),
 ('Modele de confiance','Au build, les modeles officiels tessdata_fast sont telecharges en HTTPS et controles par taille et SHA-256. Le binaire gele controle encore les empreintes. La surcharge OPENPDF_TESSDATA_DIR est explicite et ne beneficie pas du SHA-256 du bundle.'),
 ('Compilation Windows','Depuis la racine, lancer BUILD-OPENPDF-WINDOWS.cmd. Python 3.12, PyInstaller et Inno Setup 6 sont geres par les scripts existants. PREPARE-OCR-MODELS.ps1 telecharge uniquement au build et conserve la licence Apache-2.0. PyInstaller embarque tessdata en onefile et onedir.'),
 ('Qualification fonctionnelle','--ocr-runtime-self-test cree un PDF raster temporaire, execute un vrai OCR fra+eng, verifie OPENPDF et 12345, puis controle le SHA-256 du PDF OCR. Les recus Windows doivent provenir des executables natifs et inclure le test dans le portable et l installation.'),
 ('Securite et publication','Aucune publication Office si l OCR echoue. Modeles integres absents, corrompus ou symboliques rejetes. Destinations existantes protegees par publication exclusive. Fichiers temporaires nettoyes.'),
 ('Limites','Le test OCR ne prouve ni la precision sur des documents metier, ni la fidelite des tableaux, polices, images ou de la mise en page Office. Il ne remplace pas la validation humaine de la GUI Windows ni la signature Authenticode.'),
]
PINS=[
 ('eng','4 113 088','7d4322bd2a7749724879683fc3912cb542f19906c83bcc1a52132556427170b2'),
 ('fra','1 130 365','ced037562e8c80c13122dece28dd477d399af80911a28791a66a63ac1e3445ca'),
]

def diagram(path):
    im=Image.new('RGB',(1440,400),'white');d=ImageDraw.Draw(im)
    try:font=ImageFont.truetype('DejaVuSans.ttf',25);small=ImageFont.truetype('DejaVuSans.ttf',19)
    except OSError:font=ImageFont.load_default();small=font
    labels=[('1. BUILD','HTTPS officiel\\nSHA-256 + licence'),('2. EMBARQUER','PyInstaller\\nonefile / onedir'),('3. RESOUDRE','tessdata interne\\naucun repli implicite'),('4. VERIFIER','empreintes + OCR reel\\nrecus de qualification')]
    for i,(title,desc) in enumerate(labels):
        x=20+i*360
        d.rounded_rectangle((x,75,x+310,300),radius=18,fill='#EAF2FA',outline='#315D89',width=4)
        d.text((x+24,104),title,fill='#183D60',font=font)
        for j,line in enumerate(desc.split('\\n')):d.text((x+24,171+j*34),line,fill='#25394A',font=small)
        if i<3:d.line((x+314,187,x+350,187),fill='#315D89',width=5)
    d.text((22,340),'Le test Linux ne constitue pas une qualification du binaire Windows.',fill='#7A3030',font=small)
    im.save(path)

def build():
    DOCS.mkdir(exist_ok=True)
    chart=DOCS/'Architecture-OCR-Autonome-V51.3.png';diagram(chart)
    doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.7)
    doc.add_heading('OpenPDF Studio V51.3',0)
    doc.add_paragraph('Modeles OCR autonomes, integrite et qualification Windows')
    doc.add_paragraph('Auteur : Abdelhalim Talbi | 24 octobre 2026')
    doc.add_heading('Architecture de livraison',1);doc.add_picture(str(chart),width=Inches(6.3))
    for heading,body in SECTIONS:doc.add_heading(heading,2);doc.add_paragraph(body)
    doc.add_heading('Modeles officiels epingles',1)
    table=doc.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
    for c,t in zip(table.rows[0].cells,['Langue','Octets','SHA-256']):c.text=t
    for lang,size,sha in PINS:
        row=table.add_row().cells;row[0].text=lang;row[1].text=size;row[2].text=sha
    doc.add_heading('Qualification et statut',1)
    doc.add_paragraph('12 tests locaux V51.3 PASS dont OCR reel Linux ; 11 tests OCR Office PASS ; fermeture statique de 93 modules. Aucune preuve Windows ne decoule de ces controles.')
    report=ROOT/'04-Tests-et-Qualification/Preuves-V51.3/windows-ci.json'
    if report.is_file():doc.add_paragraph('CI Windows : '+str(json.loads(report.read_text(encoding='utf-8')).get('status','non verifie')))
    else:doc.add_paragraph('CI Windows V51.3 : resultat non encore lie a ce guide ; aucun binaire declare qualifie ici.')
    doc.add_paragraph('Source officielle : github.com/tesseract-ocr/tessdata_fast (Apache-2.0).')
    docx=DOCS/'OpenPDF-Studio-V51.3-Guide.docx';doc.save(docx)
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='SmallV513',parent=styles['BodyText'],fontSize=9,leading=13,spaceAfter=8))
    pdf=DOCS/'OpenPDF-Studio-V51.3-Guide.pdf'
    story=[Paragraph('OpenPDF Studio V51.3 - Guide OCR autonome',styles['Title']),
           Paragraph('Abdelhalim Talbi | 24 octobre 2026',styles['SmallV513']),
           PDFImage(str(chart),width=490,height=136),Spacer(1,12)]
    for heading,body in SECTIONS:
        story.extend([Paragraph(heading,styles['Heading2']),Paragraph(body.replace('&','&amp;'),styles['SmallV513'])])
    story.append(PageBreak())
    story.append(Paragraph('Empreintes SHA-256 des modeles officiels',styles['Heading2']))
    for code,size,sha in PINS:story.append(Paragraph(code+' ('+size+' octets) : '+sha,styles['SmallV513']))
    story.append(Paragraph('Qualification locale : 12 tests V51.3 PASS, 11 tests OCR Office PASS, 93 modules statiques. Qualification Windows non inferee.',styles['SmallV513']))
    SimpleDocTemplate(str(pdf),pagesize=A4,rightMargin=48,leftMargin=48,topMargin=44,bottomMargin=44).build(story)
    return docx,pdf,chart

if __name__=='__main__':print(build())
