import tempfile
import unittest
from pathlib import Path
from docx import Document
from table_structure_audit import compare_table_structures, inspect_office_tables

class TableStructureTests(unittest.TestCase):
    def test_header_reorder(self):
        with tempfile.TemporaryDirectory() as tmp:
            a,b=Path(tmp)/'a.docx',Path(tmp)/'b.docx'
            for p,header in [(a,['Name','Age']),(b,['Age','Name'])]:
                d=Document(); t=d.add_table(rows=2, cols=2)
                for j,val in enumerate(header): t.cell(0,j).text=val
                t.cell(1,0).text='Alice'; t.cell(1,1).text='34'; d.save(p)
            result=compare_table_structures(a,b)
            self.assertIn('header_columns_reordered',[x['kind'] for x in result['findings']])
    def test_missing_table(self):
        with tempfile.TemporaryDirectory() as tmp:
            a,b=Path(tmp)/'a.docx',Path(tmp)/'b.docx'
            d=Document();d.add_table(rows=1,cols=1);d.save(a)
            Document().save(b)
            self.assertEqual(compare_table_structures(a,b)['findings'][0]['kind'],'missing_table')
    def test_no_tables(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'a.docx';Document().save(p)
            self.assertEqual(inspect_office_tables(p),[])

if __name__=='__main__':unittest.main()
