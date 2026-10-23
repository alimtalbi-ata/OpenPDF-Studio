"""V45.9: verify every extracted text segment before publishing editable Office.

Strictly verifies the OpenPDF Studio native XLSX and PPTX layouts, not a
third-party Office layout and not visual fidelity. PDF extraction uses PyMuPDF.
Author: Abdelhalim Talbi.
"""
from pathlib import Path

import fitz


class OfficeRoundtripError(ValueError):
    """Office output does not reproduce the PDF's extracted text faithfully."""


def _source_pages(source):
    with fitz.open(source) as pdf:
        if pdf.needs_pass:
            raise OfficeRoundtripError('PDF protégé : audit texte indisponible')
        for page in pdf:
            yield page.get_text(sort=True)


def verify_pptx_text(source, destination):
    """Check exact text, order, source page, continuation and slide count."""
    from pptx import Presentation
    from openpdf_pptx_pagination_v458 import chunks_for_page

    slides = Presentation(destination).slides
    cursor = 0
    for page_no, text in enumerate(_source_pages(source), 1):
        chunks = chunks_for_page(text)
        for segment_no, expected in enumerate(chunks, 1):
            if cursor >= len(slides):
                raise OfficeRoundtripError('Diapositives manquantes : texte incomplet')
            slide = slides[cursor]
            cursor += 1
            if len(slide.shapes) != 2 or any(not s.has_text_frame for s in slide.shapes):
                raise OfficeRoundtripError('Structure de diapositive inattendue')
            title = (f'Page {page_no} — partie {segment_no}/{len(chunks)}'
                     if len(chunks) > 1 else f'Page {page_no}')
            if slide.shapes[0].text != title:
                raise OfficeRoundtripError('Ordre ou provenance des diapositives invalide')
            if slide.shapes[1].text != expected:
                raise OfficeRoundtripError('Texte PowerPoint modifié, manquant ou réordonné')
    if cursor != len(slides):
        raise OfficeRoundtripError('Diapositives supplémentaires inattendues')
    return cursor


def verify_xlsx_text(source, destination):
    """Check exact cell text, order and segment metadata; no silent truncation."""
    from openpyxl import load_workbook
    from openpdf_core_v260 import _excel_text_fragments

    wb = load_workbook(destination, read_only=True, data_only=False)
    try:
        if len(wb.worksheets) != 1:
            raise OfficeRoundtripError('Nombre de feuilles Excel inattendu')
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        if next(rows, None) != ('Page', 'Ligne', 'Texte', 'Segment'):
            raise OfficeRoundtripError('En-tête Excel invalide')
        total = 0
        for page_no, text in enumerate(_source_pages(source), 1):
            for line_no, line in enumerate(text.splitlines(), 1):
                for segment_no, expected in enumerate(_excel_text_fragments(line), 1):
                    row = next(rows, None)
                    # OOXML stores an empty string as an empty cell (None).
                    # Treat it as the same empty extracted PDF line.
                    if row is not None and len(row) == 4 and expected == "" and row[2] is None:
                        row = (row[0], row[1], "", row[3])
                    if row != (page_no, line_no, expected, segment_no):
                        raise OfficeRoundtripError('Texte ou segments Excel modifiés/réordonnés')
                    total += 1
        if next(rows, None) is not None:
            raise OfficeRoundtripError('Lignes Excel supplémentaires inattendues')
        return total
    finally:
        wb.close()


def verify_editable_text(source, destination, fmt):
    """Verify exact extracted text for all three native Office exporters."""
    source, destination = Path(source), Path(destination)
    if fmt == 'pptx':
        return verify_pptx_text(source, destination)
    if fmt == 'xlsx':
        return verify_xlsx_text(source, destination)
    if fmt == 'docx':
        from openpdf_docx_roundtrip_v460 import verify_docx_text
        return verify_docx_text(source, destination)
    raise ValueError('Format Office invalide')
