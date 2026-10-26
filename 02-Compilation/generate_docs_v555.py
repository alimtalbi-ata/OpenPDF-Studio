"""Generate version-specific V555 ZIP integrity documentation. Author: Abdelhalim Talbi."""
from __future__ import annotations

import os
import re
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
AUTHOR = "Abdelhalim Talbi"
TITLE = "OpenPDF Studio V51.4 / V555 - ZIP local-record integrity"
DOCX_NAME = "OpenPDF-Studio-V51.4-V555-ZIP-Record-Integrity-Guide.docx"
PDF_NAME = "OpenPDF-Studio-V51.4-V555-ZIP-Record-Integrity-Guide.pdf"
PNG_NAME = "Architecture-V555-ZIP-Record-Guard.png"
SVG_NAME = "Architecture-V555-ZIP-Record-Guard.svg"


def build_context(*, commit_sha=None, pr_number=None, workflow_run_id=None):
    sha = commit_sha
    if sha is None:
        sha = os.environ.get("OPENPDF_HEAD_SHA") or os.environ.get("GITHUB_SHA", "")
    if not isinstance(sha, str) or not re.fullmatch(r"[0-9a-f]{40}", sha):
        sha = "NOT-VERIFIED"
    pr = str(pr_number if pr_number is not None else os.environ.get("OPENPDF_PULL_REQUEST_NUMBER", "")).strip()
    run_id = str(workflow_run_id if workflow_run_id is not None else os.environ.get("OPENPDF_RUN_ID", "")).strip()
    return {"commit_sha": sha, "pr_number": pr or "not supplied", "workflow_run_id": run_id or "not supplied"}


def build_sections(context):
    sha_verified = context["commit_sha"] != "NOT-VERIFIED"
    pr_url = ("https://github.com/alimtalbi-ata/OpenPDF-Studio/pull/" + context["pr_number"]
              if context["pr_number"].isdigit() else "not available")
    run_url = ("https://github.com/alimtalbi-ata/OpenPDF-Studio/actions/runs/" + context["workflow_run_id"]
               if context["workflow_run_id"].isdigit() else "not available")
    return [
        ("1. Risque traité",
         "Le vérificateur V554 contrôle déjà l'amplification de décompression. Il s'appuie sur le répertoire central du ZIP; des en-têtes locaux incohérents ou des plages de données qui se chevauchent peuvent rendre l'interprétation de l'archive dépendante de l'extracteur. V555 ajoute un précontrôle structurel des enregistrements locaux."),
        ("2. Contrôles V555",
         "Avant la lecture du manifeste ou d'un contenu, le contrôle rapproche les noms, indicateurs, méthode, CRC et tailles annoncées par les en-têtes locaux et le répertoire central. Il vérifie les champs additionnels de chemin Unicode et ZIP64, les descripteurs de données, les limites physiques et l'absence de chevauchement entre enregistrements. Les octets de préfixe ou d'intervalle non indexés sont refusés. Les entrées chiffrées ou à en-têtes masqués sont refusées."),
        ("3. Tests fonctionnels",
         "Les tests fabriquent des ZIP déterministes afin de couvrir les archives Stored/Deflated et ZIP64 valides, les noms NUL ou divergents, les champs Unicode falsifiés, les CRC/tailles/méthodes/flags différents, les descripteurs corrompus, le chiffrement et le chevauchement d'enregistrements. Ce sont des fixtures synthétiques, jamais des preuves de livraison."),
        ("4. Niveaux de qualification",
         "La suite unitaire du précontrôle est exécutée sur Ubuntu et Windows par GitHub Actions. La compilation portable Windows, la construction de l'installateur, les tests de démarrage/fonctionnement/OCR et l'installation/désinstallation sont des étapes distinctes. La dernière étape doit reconstruire le ZIP cumulatif, vérifier son CRC, l'inventaire et les SHA-256 des binaires et documents."),
        ("5. Limites",
         "La validation structurelle réduit les ambiguïtés usuelles mais ne garantit pas l'absence de toute vulnérabilité ZIP ou de toute différence entre extracteurs. Le SHA-256 protège contre les altérations après calcul mais n'authentifie pas l'éditeur. Aucune signature Authenticode, précision OCR métier, fidélité Office certifiée ou revue graphique humaine Windows n'est revendiquée."),
        ("6. Priorité suivante",
         "Consolider la chaîne cumulative V549-V555, résoudre tout échec du workflow sur le SHA exact, puis évaluer la fidélité PDF vers DOCX/XLSX/PPTX sur un corpus métier annoté indépendamment, avec résultats par document et par page."),
        ("Traçabilité",
         "Commit: " + context["commit_sha"] + "\nPR: " + pr_url + "\nWorkflow: " + run_url + "\n" +
         ("Le SHA a le format attendu; les reçus CI doivent encore établir la qualification." if sha_verified
          else "Le SHA source est absent/invalide; ce document est explicitement non qualifié.")),
    ]


def svg_diagram():
    return """<svg xmlns="http://www.w3.org/2000/svg" width="1120" height="330" viewBox="0 0 1120 330" role="img" aria-labelledby="title desc">
<title id="title">OpenPDF Studio V555 ZIP local record integrity</title>
<desc id="desc">Generated architecture diagram, not an application screenshot.</desc>
<defs><marker id="arrow" markerWidth="9" markerHeight="9" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#24566b"/></marker></defs>
<rect x="18" y="76" width="200" height="100" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="118" y="109" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#153747">ZIP central directory</text>
<text x="118" y="133" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">offset / sizes / CRC</text>
<path d="M218 126 L251 126" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="258" y="62" width="215" height="128" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="365" y="91" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#153747">V555 preflight</text>
<text x="365" y="116" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">local header vs central</text>
<text x="365" y="139" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">Unicode / ZIP64 / descriptor</text>
<text x="365" y="162" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">reject overlapping spans</text>
<path d="M473 126 L507 126" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="514" y="76" width="183" height="100" rx="12" fill="#fff3df" stroke="#a75c19" stroke-width="2"/>
<text x="605" y="108" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#75410d">Invalid layout?</text>
<text x="605" y="133" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#75410d">Reject before read</text>
<path d="M697 126 L729 126" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="736" y="76" width="169" height="100" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="820" y="108" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#153747">V554 ratio gate</text>
<text x="820" y="133" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">member + aggregate</text>
<path d="M905 126 L934 126" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="941" y="76" width="161" height="100" rx="12" fill="#e7f5e9" stroke="#2b7545" stroke-width="2"/>
<text x="1021" y="108" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#1d5833">Inventory SHA-256</text>
<text x="1021" y="133" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#1d5833">and CRC verification</text>
<rect x="258" y="225" width="844" height="64" rx="8" fill="#fbe7e7" stroke="#a33333" stroke-width="1.5"/>
<text x="680" y="250" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#8b2525">Raw metadata only before any payload read; release gates still required</text>
<text x="680" y="272" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#8b2525">Generated architecture diagram — not a real application screenshot or publisher signature.</text>
</svg>"""


def draw_diagram_png(path):
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1120, 330), "white")
    draw = ImageDraw.Draw(image)
    font_path = next((p for p in (
        "C:/Windows/Fonts/segoeui.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ) if Path(p).is_file()), None)
    font = ImageFont.truetype(font_path, 15) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(font_path, 12) if font_path else ImageFont.load_default()
    boxes = [
        (18, 76, 218, 176, "ZIP central directory", "offset / sizes / CRC", "#e8f3f5", "#24566b"),
        (258, 62, 473, 190, "V555 preflight", "local vs central; ZIP64", "#e8f3f5", "#24566b"),
        (514, 76, 697, 176, "Invalid layout?", "reject before read", "#fff3df", "#a75c19"),
        (736, 76, 905, 176, "V554 ratio gate", "member + aggregate", "#e8f3f5", "#24566b"),
        (941, 76, 1102, 176, "Inventory SHA-256", "CRC verification", "#e7f5e9", "#2b7545"),
    ]
    for x0, y0, x1, y1, title, sub, fill, outline in boxes:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=12, fill=fill, outline=outline, width=2)
        draw.text((x0 + 8, y0 + 28), title, font=small, fill="#153747")
        draw.text((x0 + 8, y0 + 56), sub, font=small, fill="#153747")
    for x0, x1 in ((218, 258), (473, 514), (697, 736), (905, 941)):
        y = 126
        draw.line((x0 + 2, y, x1 - 6, y), fill="#24566b", width=3)
        draw.polygon([(x1 - 6, y), (x1 - 16, y - 6), (x1 - 16, y + 6)], fill="#24566b")
    draw.rounded_rectangle((258, 225, 1102, 289), radius=8, fill="#fbe7e7", outline="#a33333", width=2)
    draw.text((275, 241), "Raw metadata only before any payload read; release gates still required", font=small, fill="#8b2525")
    draw.text((275, 263), "Architecture diagram only — not an application screenshot or publisher signature", font=small, fill="#8b2525")
    image.save(path, format="PNG", optimize=True)


def generate(root=ROOT, *, commit_sha=None, pr_number=None, workflow_run_id=None):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PdfImage, Table, TableStyle, KeepTogether

    root = Path(root)
    docs = root / "05-Documentation"
    docs.mkdir(parents=True, exist_ok=True)
    context = build_context(commit_sha=commit_sha, pr_number=pr_number, workflow_run_id=workflow_run_id)
    sections = build_sections(context)
    svg_path = docs / SVG_NAME
    png_path = docs / PNG_NAME
    docx_path = docs / DOCX_NAME
    pdf_path = docs / PDF_NAME
    svg_path.write_text(svg_diagram(), encoding="utf-8")
    draw_diagram_png(png_path)

    document = Document()
    document.core_properties.author = AUTHOR
    document.core_properties.title = TITLE
    document.core_properties.subject = "V555 ZIP local-record integrity and evidence levels"
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading(TITLE, level=0)
    p = document.add_paragraph()
    p.add_run("Auteur : ").bold = True
    p.add_run(AUTHOR + "   ")
    p.add_run("Version applicative : ").bold = True
    p.add_run("51.4   ")
    p.add_run("Jalon : ").bold = True
    p.add_run("V555")
    document.add_paragraph("Commit source : " + context["commit_sha"])
    document.add_paragraph("PR : " + context["pr_number"] + " | Workflow : " + context["workflow_run_id"])
    if context["commit_sha"] == "NOT-VERIFIED":
        note = document.add_paragraph("SHA absent ou invalide : guide non qualifié.")
        for run in note.runs:
            run.bold = True
    document.add_picture(str(png_path), width=Inches(6.35))
    caption = document.add_paragraph("Figure 1. Schéma généré pour expliquer le contrôle, pas une capture de l'interface.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, body in sections:
        document.add_heading(heading, level=1)
        for line in body.split("\n"):
            if line.strip():
                document.add_paragraph(line)
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("Niveau", "Contrôle", "Portée")):
        cell.text = text
    for row in (
        ("Précontrôle structurel", "En-têtes locaux + central", "Métadonnées avant lecture"),
        ("Tests fonctionnels", "Fixtures adversariales", "Pas une preuve de livraison"),
        ("Windows natif", "Portable / installateur", "Doit être prouvé par CI exacte"),
        ("Authenticité éditeur", "Non revendiquée", "SHA-256 ne signe pas l'éditeur"),
    ):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    document.save(docx_path)

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="V555Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12, spaceAfter=6))
    styles.add(ParagraphStyle(name="V555Small", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, leading=9, spaceAfter=5))
    story = [
        Paragraph(escape(TITLE), styles["Title"]),
        Spacer(1, 3 * mm),
        Paragraph(escape("Auteur : " + AUTHOR + " | Version applicative : 51.4 | Jalon : V555"), styles["V555Body"]),
        Paragraph(escape("Commit : " + context["commit_sha"] + " | PR : " + context["pr_number"] + " | Workflow : " + context["workflow_run_id"]), styles["V555Small"]),
        PdfImage(str(png_path), width=174 * mm, height=174 * mm * 330 / 1120),
        Paragraph("Schéma d'architecture généré — pas une capture de l'application.", styles["Italic"]),
        Spacer(1, 3 * mm),
    ]
    for heading, body in sections:
        block = [Paragraph(escape(heading), styles["Heading2"])]
        block.extend(
            Paragraph(escape(line), styles["V555Body"])
            for line in body.split("\n")
            if line.strip()
        )
        # Never leave the provenance heading stranded at a page boundary.
        if heading == "Traçabilité":
            story.append(KeepTogether(block))
        else:
            story.extend(block)
    rows = [
        ["Niveau", "Contrôle", "Portée"],
        ["Précontrôle", "Local + central", "Métadonnées avant lecture"],
        ["Tests fonctionnels", "Fixtures adversariales", "Pas une preuve de livraison"],
        ["Windows natif", "Portable / installateur", "CI exacte requise"],
        ["Authenticité", "Non revendiquée", "SHA-256 n'authentifie pas l'éditeur"],
    ]
    table = Table(rows, colWidths=[35 * mm, 53 * mm, 86 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24566b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9db3bd")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([Spacer(1, 3 * mm), KeepTogether([Paragraph("Niveaux de preuve", styles["Heading2"]), table])])
    SimpleDocTemplate(str(pdf_path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title=TITLE, author=AUTHOR).build(story)

    for path in (docx_path, pdf_path, png_path, svg_path):
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError("Generated file is missing/empty: " + str(path))
    return {"context": context, "docx": str(docx_path), "pdf": str(pdf_path), "png": str(png_path), "svg": str(svg_path)}


if __name__ == "__main__":
    for key, value in generate().items():
        print(str(key) + ": " + str(value))
