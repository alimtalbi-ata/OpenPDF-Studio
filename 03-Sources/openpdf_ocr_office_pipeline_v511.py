"""V51.1: transactional offline OCR -> editable Office text conversion.

OCR text coverage is measured, not a claim of OCR accuracy or layout fidelity.
The original PDF is never modified; no destination is published on failure.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations

import argparse
from dataclasses import asdict, dataclass
import hashlib
import json
import os
from pathlib import Path
import shutil
import tempfile

from openpdf_atomic_output import publish_exclusive
from openpdf_conversion_service_v351 import SUPPORTED_FORMATS, convert_office
from openpdf_office_scan_gate_v510 import assess_office_source, require_editable_source
from openpdf_ocr_quality_gate_v488 import OCRQualityPolicy
from openpdf_ocr_service_v347 import searchable_pdf
from openpdf_output_policy_v300 import validate_output
from openpdf_source_integrity_v430 import fingerprint_source, verify_source_unchanged


@dataclass(frozen=True)
class OCROfficeOutcome:
    schema: str
    output_format: str
    destination: str
    pages: int
    scanned_pages: tuple[int, ...]
    ocr_pages: int
    recognized_characters: int
    bytes_written: int
    sha256: str
    source_sha256: str
    warnings: tuple[str, ...]
    table_mode: bool = False
    table_worksheets: int = 0
    table_pages: tuple[int, ...] = ()
    scanned_pages_without_tables: tuple[int, ...] = ()
    table_cells_by_page: tuple[tuple[int, int, int], ...] = ()
    fidelity: str = 'editable_text_only_no_layout_or_recognition_accuracy_guarantee'

    def to_dict(self):
        data = asdict(self)
        data['scanned_pages'] = list(self.scanned_pages)
        data['table_pages'] = list(self.table_pages)
        data['scanned_pages_without_tables'] = list(self.scanned_pages_without_tables)
        data['table_cells_by_page'] = [
            {'page': page, 'cells': total, 'nonempty_cells': nonempty}
            for page, total, nonempty in self.table_cells_by_page]
        data['warnings'] = list(self.warnings)
        return data


def _sha256(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def convert_scanned_office(source, destination, *, language='fra+eng', dpi=200,
                           min_characters_per_scanned_page=1, tessdata=None,
                           table_mode=False, require_table_per_scanned_page=False,
                           min_nonempty_cells_per_scanned_page=0):
    """Convert scanned/mixed PDF into DOCX/XLSX/PPTX via offline Tesseract.

    The opt-in table_mode attempts editable table worksheets in XLSX only.
    It refuses output when no sufficiently reliable table is detected.
    Optional strict table coverage refuses an OCR result when any scanned
    page has no extractable table; the default preserves legacy behavior.
    An optional per-scanned-page minimum counts actual nonempty editable
    table cells, not OCR characters or guessed recognition confidence.
    Scanned pages must have at least min_characters_per_scanned_page
    extractable characters after OCR. This does not measure correctness.
    The staged Office file is audited by convert_office before exclusive commit.
    """
    source, destination = Path(source).expanduser(), Path(destination).expanduser()
    fmt = destination.suffix.lower().lstrip('.')
    if fmt not in SUPPORTED_FORMATS:
        raise ValueError('Format Office invalide : docx, xlsx ou pptx')
    if type(table_mode) is not bool or (table_mode and fmt != 'xlsx'):
        raise ValueError('Mode tableaux OCR reserve au format XLSX')
    if type(require_table_per_scanned_page) is not bool or (require_table_per_scanned_page and not table_mode):
        raise ValueError('Couverture stricte des pages scannees reservee au mode tableaux XLSX')
    if (type(min_nonempty_cells_per_scanned_page) is not int
            or not 0 <= min_nonempty_cells_per_scanned_page <= 10000
            or (min_nonempty_cells_per_scanned_page and not table_mode)):
        raise ValueError('Seuil de cellules OCR invalide ou mode tableaux XLSX absent')
    if source.is_symlink() or not source.is_file() or source.suffix.lower() != '.pdf':
        raise ValueError('Source PDF invalide ou lien symbolique interdit')
    if destination.is_symlink() or destination.exists():
        raise FileExistsError('Destination deja existante ou lien symbolique')
    if not destination.parent.is_dir():
        raise ValueError('Dossier de destination absent')
    if type(min_characters_per_scanned_page) is not int or not 1 <= min_characters_per_scanned_page <= 10000:
        raise ValueError('Seuil de caracteres OCR invalide')
    if type(dpi) is not int or not 100 <= dpi <= 400:
        raise ValueError('DPI OCR hors limites')
    validate_output([source], destination)
    original = fingerprint_source(source)
    preflight = assess_office_source(source)
    if original.sha256 != preflight.source_sha256:
        raise RuntimeError('Source modifiee pendant le precontrole')
    scanned = tuple(preflight.scanned_pages)
    stage = None
    try:
        with tempfile.TemporaryDirectory(prefix='.openpdf-ocr-office-', dir=destination.parent) as private:
            private = Path(private)
            pdf_for_office = source
            ocr_pages = 0
            recognized_characters = 0
            if scanned:
                pdf_for_office = private / 'recognized.pdf'
                receipt = searchable_pdf(
                    source, pdf_for_office, language=language, dpi=dpi,
                    skip_text=True, tessdata=tessdata,
                    quality_policy=OCRQualityPolicy(
                        min_covered_fraction=0.0,
                        min_characters_per_page=0,
                        max_replacement_characters=0))
                if receipt.sha256 != _sha256(pdf_for_office):
                    raise RuntimeError('Empreinte du PDF OCR incoherente')
                ocr_pages = receipt.ocr_pages
                if ocr_pages < len(scanned):
                    raise RuntimeError('Pages scannees non traitees par OCR')
                import pymupdf as fitz
                with fitz.open(pdf_for_office) as recognized:
                    if len(recognized) != preflight.page_count:
                        raise RuntimeError('OCR a modifie le nombre de pages')
                    for page_number in scanned:
                        extracted = recognized[page_number - 1].get_text('text')
                        count = sum(not ch.isspace() for ch in extracted)
                        if count < min_characters_per_scanned_page or '\ufffd' in extracted:
                            raise ValueError('OCR insuffisant sur page ' + str(page_number))
                        recognized_characters += count
                require_editable_source(assess_office_source(pdf_for_office))
            verify_source_unchanged(source, original)
            office_stage = private / ('converted.' + fmt)
            converted = convert_office(pdf_for_office, office_stage,
                                       output_format=fmt, table_mode=table_mode)
            verify_source_unchanged(source, original)
            if converted.sha256 != _sha256(office_stage):
                raise RuntimeError('Empreinte du document Office incoherente')
            table_worksheets = 0
            table_pages = ()
            scanned_pages_without_tables = ()
            table_cells_by_page = ()
            if table_mode:
                # The production converter already verifies every extracted cell.
                # Report table coverage, not a semantic OCR accuracy score.
                from openpdf_xlsx_tables_v514 import _expected
                # Reuse the audited table extractor; never infer coverage from
                # OCR character counts or workbook sheet names alone.
                _tables, evidence = _expected(pdf_for_office)
                table_pages = evidence.pages_with_tables
                # Count extracted nonempty values per page, across *all* tables.
                # A detected empty grid is not evidence of a usable conversion.
                counts = {}
                for page_number, _index, matrix, width, _merges in _tables:
                    total, nonempty = counts.get(page_number, (0, 0))
                    total += len(matrix) * width
                    nonempty += sum(value not in (None, '')
                                    for row in matrix for value in row)
                    counts[page_number] = (total, nonempty)
                table_cells_by_page = tuple(
                    (page_number, *counts[page_number])
                    for page_number in sorted(counts))
                insufficient_pages = tuple(
                    page_number for page_number in scanned
                    if counts.get(page_number, (0, 0))[1] < min_nonempty_cells_per_scanned_page)
                if insufficient_pages:
                    raise ValueError('Cellules OCR insuffisantes sur page scannee : ' +
                                     ', '.join(map(str, insufficient_pages)))
                scanned_pages_without_tables = tuple(
                    number for number in scanned if number not in table_pages)
                if require_table_per_scanned_page and scanned_pages_without_tables:
                    raise ValueError('Tableau OCR absent sur page scannee : ' +
                                     ', '.join(map(str, scanned_pages_without_tables)))
                from openpyxl import load_workbook
                workbook = load_workbook(office_stage, read_only=True, data_only=False)
                try:
                    table_worksheets = len(workbook.sheetnames) - 1
                    if table_worksheets != len(_tables):
                        raise RuntimeError('Nombre de feuilles de tableaux OCR incoherent')
                finally:
                    workbook.close()
            fd, stage_name = tempfile.mkstemp(prefix='.openpdf-office-commit-',
                                               suffix='.' + fmt, dir=destination.parent)
            os.close(fd)
            stage = Path(stage_name)
            shutil.copyfile(office_stage, stage)
            if _sha256(stage) != converted.sha256:
                raise RuntimeError('Copie du document Office alteree')
            verify_source_unchanged(source, original)
            if destination.is_symlink():
                raise FileExistsError('Destination devenue lien symbolique')
            publish_exclusive(stage, destination)
            warnings = tuple(converted.warnings) + (
                ('OCR coverage is not recognition accuracy; Office layout is not preserved',)
                if scanned else ('Office layout is not preserved',))
            if table_mode:
                warnings += ('OCR table geometry and values require human verification',)
                if scanned_pages_without_tables:
                    warnings += ('Scanned pages without detected editable tables: ' +
                                 ', '.join(map(str, scanned_pages_without_tables)),)
            return OCROfficeOutcome(
                'openpdf.ocr-office.v51.1', fmt, str(destination), preflight.page_count,
                scanned, ocr_pages, recognized_characters, stage.stat().st_size,
                converted.sha256, original.sha256, warnings,
                table_mode=table_mode, table_worksheets=table_worksheets,
                table_pages=table_pages,
                scanned_pages_without_tables=scanned_pages_without_tables,
                table_cells_by_page=table_cells_by_page)
    finally:
        if stage is not None:
            stage.unlink(missing_ok=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description='Offline OCR then editable-text Office conversion')
    parser.add_argument('source', type=Path)
    parser.add_argument('destination', type=Path)
    parser.add_argument('--language', default='fra+eng')
    parser.add_argument('--dpi', type=int, default=200)
    parser.add_argument('--min-chars', type=int, default=1)
    parser.add_argument('--editable-tables', action='store_true',
                        help='Opt-in editable table worksheets for XLSX; may refuse ambiguous OCR')
    parser.add_argument('--require-table-each-scanned-page', action='store_true',
                        help='Fail closed if any scanned page lacks a detected editable table')
    parser.add_argument('--min-nonempty-table-cells-per-scanned-page', type=int, default=0,
                        help='Fail closed if a scanned page has fewer than N nonempty editable table cells (0 disables)')
    args = parser.parse_args(argv)
    outcome = convert_scanned_office(args.source, args.destination,
        language=args.language, dpi=args.dpi,
        min_characters_per_scanned_page=args.min_chars,
        table_mode=args.editable_tables,
        require_table_per_scanned_page=args.require_table_each_scanned_page,
        min_nonempty_cells_per_scanned_page=args.min_nonempty_table_cells_per_scanned_page)
    print(json.dumps(outcome.to_dict(), ensure_ascii=False, indent=2))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
