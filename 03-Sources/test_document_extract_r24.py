import tempfile,unittest
from pathlib import Path
from document_extract import extract_document,compare_documents
class OfficeTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.folder=Path(self.tmp.name)
  from pptx import Presentation
  pres=Presentation();slide=pres.slides.add_slide(pres.slide_layouts[6]);shape=slide.shapes.add_textbox(100000,100000,3000000,600000);shape.text='Slide heading';pres.save(self.folder/'a.pptx')
  from openpyxl import Workbook
  wb=Workbook();ws=wb.active;ws['A1']='Heading';ws['B2']=42;wb.save(self.folder/'b.xlsx')
 def test_pptx_shape_geometry(self):
  r=extract_document(self.folder/'a.pptx');self.assertEqual(r['geometry'],'pptx_shape_bounds');self.assertEqual(len(r['pages']),1);self.assertEqual(r['pages'][0]['blocks'][0]['text'],'Slide heading')
 def test_xlsx_synthetic(self):
  r=extract_document(self.folder/'b.xlsx');self.assertEqual(r['geometry'],'synthetic_grid');self.assertEqual(len(r['pages'][0]['blocks']),2);self.assertTrue(r['warnings'])
 def test_pptx_pipeline(self):
  r=compare_documents(self.folder/'a.pptx',self.folder/'a.pptx');self.assertIn('report',r['diagnostic'])
 def test_xlsx_pipeline(self):
  r=compare_documents(self.folder/'b.xlsx',self.folder/'b.xlsx');self.assertIn('report',r['diagnostic'])
if __name__=='__main__':unittest.main(verbosity=2)
