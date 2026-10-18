import tempfile, unittest
from pathlib import Path
from document_extract import render_office_to_pdf, compare_documents, ExtractionError

class RenderingTests(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
  self.folder=Path(self.temp.name)
  from docx import Document
  doc=Document();doc.add_heading('Rendered document',0);doc.add_paragraph('Real coordinates via LibreOffice');doc.save(self.folder/'sample.docx')
 def test_render_docx(self):
  r=render_office_to_pdf(self.folder/'sample.docx',timeout=30)
  self.assertEqual(r['geometry'],'libreoffice_rendered_pdf')
  self.assertTrue(any('Rendered document' in b['text'] for p in r['pages'] for b in p['blocks']))
 def test_pdf_comparison_rendered(self):
  from reportlab.pdfgen import canvas
  path=self.folder/'sample.pdf';c=canvas.Canvas(str(path));c.drawString(50,750,'Rendered document');c.save()
  result=compare_documents(path,self.folder/'sample.docx',render_office=True,render_timeout=30)
  self.assertTrue(result['layout_comparable'])
  self.assertEqual(result['target']['geometry'],'libreoffice_rendered_pdf')
 def test_default_remains_synthetic(self):
  result=compare_documents(self.folder/'sample.docx',self.folder/'sample.docx')
  self.assertFalse(result['layout_comparable'])
 def test_explicit_failure(self):
  with self.assertRaisesRegex(ExtractionError,'not found'):
   render_office_to_pdf(self.folder/'sample.docx',executable='missing-openpdf-office-executable')
 def test_timeout_validation(self):
  with self.assertRaisesRegex(ExtractionError,'timeout'):
   render_office_to_pdf(self.folder/'sample.docx',timeout=0)

if __name__=='__main__':unittest.main(verbosity=2)
