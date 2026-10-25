"""V542 independent, bounded, privacy-preserving editable Office token coverage.
Actual editable DOCX/XLSX/PPTX content versus independently authored references.
Not a measure of layout, visual fidelity, reading order, or OCR character accuracy.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations
import argparse
from collections import Counter
import hashlib
import json
import math
from pathlib import Path
import re
import stat
import unicodedata
import zipfile
from openpdf_ocr_reference_corpus_v539 import _digest, _file, _json, ID_RE

SCHEMA = 'openpdf.office.reference.corpus.v1'
ANNOTATION_SCHEMA = 'openpdf.office.reference.v1'
REPORT_SCHEMA = 'openpdf.office.reference.report.v1'
FORMATS = {'.docx', '.xlsx', '.pptx'}
MAX_CASES = 32
MAX_CASE_TOKENS = 5000
MAX_TOTAL_TOKENS = 50000
MAX_EDITABLE_ELEMENTS = 100000
MAX_ZIP_BYTES = 32 * 1024 * 1024
MAX_ZIP_ENTRIES = 2048
MAX_MEMBER_BYTES = 8 * 1024 * 1024
MAX_UNPACKED_BYTES = 64 * 1024 * 1024
TOKEN_RE = re.compile(r'\w+', re.UNICODE)

def _tokens(text):
    return TOKEN_RE.findall(unicodedata.normalize('NFKC', text).casefold())

def _source_pdf_preflight(path):
    import fitz
    try:
        with fitz.open(path) as document:
            if document.needs_pass or not 1 <= document.page_count <= 64:
                raise ValueError('Invalid or encrypted source PDF')
    except (fitz.FileDataError, fitz.EmptyFileError) as exc:
        raise ValueError('Invalid source PDF') from exc

def _office_preflight(path):
    """Reject traversal, ZIP bombs, links, active content and external relationships."""
    if not zipfile.is_zipfile(path) or path.stat().st_size > MAX_ZIP_BYTES:
        raise ValueError('Invalid or oversized Office package')
    with zipfile.ZipFile(path) as package:
        members = package.infolist()
        if not 1 <= len(members) <= MAX_ZIP_ENTRIES:
            raise ValueError('Invalid Office package entry count')
        names = set()
        total = 0
        for member in members:
            name = member.filename
            if (name in names or not name or name.startswith('/') or '\\' in name or
                    ':' in name or any(part in ('', '.', '..') for part in name.split('/'))):
                raise ValueError('Unsafe or duplicate Office member')
            names.add(name)
            if member.flag_bits & 1 or stat.S_ISLNK(member.external_attr >> 16):
                raise ValueError('Encrypted or linked Office member')
            total += member.file_size
            if (member.file_size > MAX_MEMBER_BYTES or total > MAX_UNPACKED_BYTES or
                    member.file_size > max(1, member.compress_size) * 500):
                raise ValueError('Office decompression budget exceeded')
            lowered = name.lower()
            if (lowered.endswith(('.vbs', '.js', '.exe', '.dll')) or
                    'vbaproject' in lowered or '/activex/' in lowered or
                    '/externallinks/' in lowered or '/embeddings/' in lowered or
                    (lowered.endswith('.bin') and '/printersettings/' not in lowered)):
                raise ValueError('Active or embedded Office content refused')
            if lowered.endswith('.rels'):
                from xml.etree import ElementTree as ET
                try:
                    relations = ET.fromstring(package.read(member))
                except ET.ParseError as exc:
                    raise ValueError('Malformed Office relationship') from exc
                if any(el.attrib.get('TargetMode', '').lower() == 'external'
                       for el in relations.iter()):
                    raise ValueError('External Office relationships refused')
        if '[Content_Types].xml' not in names:
            raise ValueError('Office content types missing')

def _editable_text(path):
    """Count native editable elements, never embedded images."""
    suffix = path.suffix.lower()
    values = []
    elements = 0
    if suffix == '.docx':
        from docx import Document
        from docx.oxml.ns import qn
        document = Document(path)
        for node in document.element.body.iter(qn('w:p')):
            value = ''.join(item.text or '' for item in node.iter(qn('w:t')))
            if value.strip():
                values.append(value)
                elements += 1
                if elements > MAX_EDITABLE_ELEMENTS:
                    raise ValueError('Office editable-element budget exceeded')
    elif suffix == '.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        presentation = Presentation(path)
        def visit(shapes):
            nonlocal elements
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    visit(shape.shapes)
                if shape.has_text_frame and shape.text.strip():
                    values.append(shape.text)
                    elements += 1
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text.strip():
                                values.append(cell.text)
                                elements += 1
                if elements > MAX_EDITABLE_ELEMENTS:
                    raise ValueError('Office editable-element budget exceeded')
        for slide in presentation.slides:
            visit(slide.shapes)
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in workbook.worksheets:
                if sheet.max_row > MAX_EDITABLE_ELEMENTS or sheet.max_column > MAX_EDITABLE_ELEMENTS:
                    raise ValueError('Office worksheet dimensions exceed budget')
                for row in sheet.iter_rows(values_only=True):
                    for value in row:
                        if value is not None and str(value).strip():
                            values.append(str(value))
                            elements += 1
                            if elements > MAX_EDITABLE_ELEMENTS:
                                raise ValueError('Office editable-element budget exceeded')
        finally:
            workbook.close()
    else:
        raise ValueError('Unsupported Office format')
    return Counter(token for value in values for token in _tokens(value)), elements

def _thresholds(value):
    required = {'min_token_recall', 'min_case_token_recall', 'min_human_reviewed',
                'min_editable_elements_per_case'}
    if not isinstance(value, dict) or set(value) != required:
        raise ValueError('Invalid Office reference thresholds')
    for name in ('min_token_recall', 'min_case_token_recall'):
        x = value[name]
        if type(x) not in (int, float) or not math.isfinite(x) or not 0 <= x <= 1:
            raise ValueError('Invalid Office recall threshold')
    for name, upper in (('min_human_reviewed', MAX_CASES),
                        ('min_editable_elements_per_case', MAX_EDITABLE_ELEMENTS)):
        x = value[name]
        if type(x) is not int or not 0 <= x <= upper:
            raise ValueError('Invalid Office count threshold')
    return value

def evaluate_office_reference_corpus(manifest):
    """Unrounded threshold decisions; aggregated report without PDF text or paths."""
    manifest = Path(manifest)
    data = _json(manifest, 128 * 1024)
    if not isinstance(data, dict) or set(data) != {'schema', 'cases', 'thresholds'} or data['schema'] != SCHEMA:
        raise ValueError('Invalid Office reference manifest')
    cases = data['cases']
    if not isinstance(cases, list) or not 2 <= len(cases) <= MAX_CASES:
        raise ValueError('Office corpus requires 2 to 32 cases')
    thresholds = _thresholds(data['thresholds'])
    if thresholds['min_human_reviewed'] > len(cases):
        raise ValueError('Impossible human review threshold')
    manifest_hash = _digest(manifest, 128 * 1024)
    ids, source_hashes, office_hashes, evidence_hashes = set(), set(), set(), []
    total = matched = humans = failed_cases = missing_elements = editable_elements = 0
    worst_case_recall = 1.0
    formats = {fmt[1:]: {'cases': 0, 'reference_tokens': 0, 'matched_tokens': 0}
               for fmt in sorted(FORMATS)}
    for case in cases:
        if not isinstance(case, dict) or set(case) != {'id', 'source_pdf', 'office_file', 'annotations'}:
            raise ValueError('Invalid Office reference case')
        identifier = case['id']
        if type(identifier) is not str or not ID_RE.fullmatch(identifier) or identifier in ids:
            raise ValueError('Invalid or duplicate Office case ID')
        ids.add(identifier)
        root = manifest.parent
        source, office, annotation_path = (_file(root, case[key])
                                           for key in ('source_pdf', 'office_file', 'annotations'))
        if source.suffix.lower() != '.pdf' or office.suffix.lower() not in FORMATS:
            raise ValueError('Invalid source or Office format')
        if len({source, office, annotation_path}) != 3:
            raise ValueError('Duplicate Office evidence paths')
        source_sha = _digest(source, 32 * 1024 * 1024)
        office_sha = _digest(office, MAX_ZIP_BYTES)
        annotation_sha = _digest(annotation_path, 256 * 1024)
        if source_sha in source_hashes or office_sha in office_hashes:
            raise ValueError('Duplicate Office source or output')
        source_hashes.add(source_sha); office_hashes.add(office_sha)
        _source_pdf_preflight(source)
        annotation = _json(annotation_path, 256 * 1024)
        if (not isinstance(annotation, dict) or
                set(annotation) != {'schema', 'source_sha256', 'origin', 'expected_phrases'} or
                annotation['schema'] != ANNOTATION_SCHEMA or annotation['source_sha256'] != source_sha or
                annotation['origin'] not in ('human_reviewed', 'synthetic_known_at_generation')):
            raise ValueError('Invalid or unbound Office annotation')
        humans += annotation['origin'] == 'human_reviewed'
        phrases = annotation['expected_phrases']
        if (not isinstance(phrases, list) or not 1 <= len(phrases) <= 1000 or
                any(type(p) is not str or not p.strip() or len(p) > 5000 for p in phrases)):
            raise ValueError('Invalid Office reference phrases')
        expected = Counter(token for phrase in phrases for token in _tokens(phrase))
        expected_total = sum(expected.values())
        if not 1 <= expected_total <= MAX_CASE_TOKENS or total + expected_total > MAX_TOTAL_TOKENS:
            raise ValueError('Office reference token budget exceeded')
        _office_preflight(office)
        actual, elements = _editable_text(office)
        if _digest(source, 32 * 1024 * 1024) != source_sha or _digest(office, MAX_ZIP_BYTES) != office_sha:
            raise ValueError('Office evidence modified during evaluation')
        if _digest(annotation_path, 256 * 1024) != annotation_sha:
            raise ValueError('Office annotations modified during evaluation')
        found = sum(min(count, actual[token]) for token, count in expected.items())
        recall = found / expected_total
        worst_case_recall = min(worst_case_recall, recall)
        failed_cases += recall < thresholds['min_case_token_recall']
        missing_elements += elements < thresholds['min_editable_elements_per_case']
        editable_elements += elements
        total += expected_total; matched += found
        slot = formats[office.suffix.lower()[1:]]
        slot['cases'] += 1; slot['reference_tokens'] += expected_total; slot['matched_tokens'] += found
        evidence_hashes.extend((source_sha, office_sha, annotation_sha))
    if _digest(manifest, 128 * 1024) != manifest_hash:
        raise ValueError('Office manifest modified during evaluation')
    recall = matched / total
    failed = []
    if recall < thresholds['min_token_recall']:
        failed.append('min_token_recall')
    if failed_cases:
        failed.append('min_case_token_recall')
    if missing_elements:
        failed.append('min_editable_elements_per_case')
    if humans < thresholds['min_human_reviewed']:
        failed.append('min_human_reviewed')
    return {
        'schema': REPORT_SCHEMA, 'status': 'FAIL' if failed else 'PASS',
        'scope': 'editable_ooxml_text_token_multiset_recall_against_separate_reference',
        'visual_fidelity': 'not_assessed', 'reading_order': 'not_assessed',
        'token_precision': 'not_assessed',
        'annotation_provenance': 'declared_not_independently_attested',
        'manifest_sha256': manifest_hash,
        'evidence_set_sha256': hashlib.sha256('\n'.join(sorted(evidence_hashes)).encode('ascii')).hexdigest(),
        'cases': len(cases), 'human_reviewed_cases': humans,
        'synthetic_cases': len(cases) - humans,
        'editable_elements': editable_elements, 'cases_below_recall_floor': failed_cases,
        'cases_below_editable_floor': missing_elements,
        'reference_tokens': total, 'matched_tokens': matched,
        'token_recall': round(recall, 6),
        'worst_case_token_recall': round(worst_case_recall, 6),
        'by_format': {k: {**v, 'token_recall': round(v['matched_tokens'] / v['reference_tokens'], 6)
                            if v['reference_tokens'] else None} for k, v in formats.items()},
        'thresholds': thresholds, 'failed_thresholds': failed,
    }

def main(argv=None):
    parser = argparse.ArgumentParser(description='Independent editable Office token-recall qualification')
    parser.add_argument('manifest', type=Path)
    parser.add_argument('--receipt', type=Path, help='New JSON receipt path (never overwrite)')
    args = parser.parse_args(argv)
    try:
        report = evaluate_office_reference_corpus(args.manifest)
        serialized = json.dumps(report, sort_keys=True, indent=2, ensure_ascii=False) + '\n'
        if args.receipt:
            with args.receipt.open('x', encoding='utf-8') as stream:
                stream.write(serialized)
        print(serialized, end='')
        return 0 if report['status'] == 'PASS' else 2
    except Exception as exc:
        print(json.dumps({'schema': REPORT_SCHEMA, 'status': 'ERROR',
                          'reason': type(exc).__name__}, sort_keys=True))
        return 3

if __name__ == '__main__':
    raise SystemExit(main())
