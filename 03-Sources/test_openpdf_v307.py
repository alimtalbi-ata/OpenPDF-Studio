"""V30.7: real PDF -> editable Office E2E regression tests. Abdelhalim Talbi."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch, Mock
import fitz
from openpdf_conversion_controller_v301 import ConversionController
from openpdf_ui_bridge_v301 import connect as connect_conversion
from openpdf_core_v260 import pdf_to_docx

class ConversionEndToEnd(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.src=self.root/'input.pdf'
        with fitz.open() as doc:
            page=doc.new_page()
            page.insert_text((60,90),'OpenPDF TEST V307 alpha beta')
            doc.save(self.src)
        self.original=self.src.read_bytes()

    def test_word_editable(self):
        from docx import Document
        target=self.root/'output.docx'
        result=ConversionController().convert(self.src,target,'docx')
        self.assertGreater(result.bytes_written,0)
        text='\n'.join(p.text for p in Document(target).paragraphs)
        self.assertIn('OpenPDF TEST V307',text)

    def test_excel_editable(self):
        from openpyxl import load_workbook
        target=self.root/'output.xlsx'
        ConversionController().convert(self.src,target,'xlsx')
        book=load_workbook(target,read_only=True)
        try:
            cells=[str(cell) for row in book.active.iter_rows(values_only=True) for cell in row if cell is not None]
            self.assertTrue(any('OpenPDF TEST V307' in c for c in cells))
        finally: book.close()

    def test_powerpoint_editable(self):
        from pptx import Presentation
        target=self.root/'output.pptx'
        ConversionController().convert(self.src,target,'pptx')
        deck=Presentation(target)
        texts=[shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame]
        self.assertTrue(any('OpenPDF TEST V307' in t for t in texts))

    def test_txt_utf8(self):
        target=self.root/'output.txt'
        ConversionController().convert(self.src,target,'txt')
        self.assertIn('OpenPDF TEST V307',target.read_text(encoding='utf-8'))

    def test_failure_preserves_existing_output(self):
        target=self.root/'output.docx';target.write_bytes(b'USER DATA')
        with patch('openpdf_conversion_controller_v301.CONVERTERS',{'docx':lambda s,d: (_ for _ in ()).throw(RuntimeError('failure'))}):
            with self.assertRaises(RuntimeError): ConversionController().convert(self.src,target,'docx')
        self.assertEqual(target.read_bytes(),b'USER DATA')
        self.assertEqual(list(self.root.glob('.openpdf-convert-*')),[])

    def test_rejects_bad_extension(self):
        with self.assertRaises(ValueError): ConversionController().convert(self.src,self.root/'bad.pdf','docx')

    def test_source_unchanged(self):
        ConversionController().convert(self.src,self.root/'out.txt','txt')
        self.assertEqual(self.src.read_bytes(),self.original)

    def test_ui_dispatch_worker_receives_taskrunner_callbacks(self):
        app=type('FakeApp', (), {})(); app._run_pdf=Mock(return_value='scheduled')
        original=app._run_pdf
        connect_conversion(app)
        app._run_pdf(pdf_to_docx,self.src,self.root/'out.docx')
        worker=original.call_args.args[0]
        progress=[]
        result=worker(progress=lambda a,b:progress.append((a,b)),cancelled=lambda:False)
        self.assertTrue(Path(result).exists())
        self.assertEqual(progress,[(1,1)])

    def test_ui_cancel_before_conversion(self):
        app=type('FakeApp', (), {})(); app._run_pdf=Mock(return_value='scheduled')
        original=app._run_pdf
        connect_conversion(app)
        app._run_pdf(pdf_to_docx,self.src,self.root/'cancel.docx')
        worker=original.call_args.args[0]
        with self.assertRaises(RuntimeError):worker(cancelled=lambda:True)
        self.assertFalse((self.root/'cancel.docx').exists())

if __name__=='__main__':unittest.main()
