"""V558 functional and adversarial tests; all PDFs/Office inputs are synthetic fixtures.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

import fitz
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "03-Sources"))
from openpdf_office_visual_fidelity_v558 import (
    ANNOTATION_SCHEMA,
    DEFAULT_THRESHOLDS,
    REPORT_SCHEMA,
    SCHEMA,
    evaluate_office_visual_fidelity_corpus,
    main,
    safe_visual_fidelity_receipt,
    serialize_visual_fidelity_receipt,
)


class OfficeVisualFidelityV558Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.cases = []

    @staticmethod
    def digest(path):
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def make_pdf(self, path, text, *, x=72, pages=1):
        with fitz.open() as document:
            for index in range(pages):
                page = document.new_page(width=595, height=842)
                page.insert_text((x, 100), text if index == 0 else f"extra page {index + 1}", fontsize=12)
            document.save(path)
        return path

    def make_office(self, path, text):
        document = Document()
        document.add_heading("OpenPDF synthetic Office evidence", level=1)
        document.add_paragraph(text)
        document.save(path)
        return path

    def add_case(self, source_text, *, rendered_text=None, render_x=72,
                 source_pages=1, rendered_pages=None, origin="synthetic_known_at_generation"):
        index = len(self.cases) + 1
        source = self.root / f"source-{index}.pdf"
        rendered = self.root / f"rendered-{index}.pdf"
        office = self.root / f"office-{index}.docx"
        annotation = self.root / f"annotation-{index}.json"
        self.make_pdf(source, source_text, pages=source_pages)
        actual_text = source_text if rendered_text is None else rendered_text
        self.make_pdf(rendered, actual_text, x=render_x,
                      pages=source_pages if rendered_pages is None else rendered_pages)
        self.make_office(office, source_text)
        reference = {
            "schema": ANNOTATION_SCHEMA,
            "source_sha256": self.digest(source),
            "office_sha256": self.digest(office),
            "rendered_pdf_sha256": self.digest(rendered),
            "origin": origin,
            "renderer": "libreoffice",
            "renderer_version": "24.2.7.2",
        }
        annotation.write_text(json.dumps(reference, sort_keys=True), encoding="utf-8")
        case = {
            "id": f"case-{index}",
            "source_pdf": source.name,
            "office_file": office.name,
            "rendered_pdf": rendered.name,
            "annotations": annotation.name,
        }
        self.cases.append(case)
        return case

    def write_manifest(self, thresholds=None):
        path = self.root / "manifest.json"
        data = {
            "schema": SCHEMA,
            "cases": self.cases,
            "thresholds": dict(DEFAULT_THRESHOLDS if thresholds is None else thresholds),
        }
        path.write_text(json.dumps(data, sort_keys=True, allow_nan=True), encoding="utf-8")
        return path

    def refresh_annotation_hashes(self, case):
        annotation = self.root / case["annotations"]
        data = json.loads(annotation.read_text(encoding="utf-8"))
        for field, path_key in (
            ("source_sha256", "source_pdf"),
            ("office_sha256", "office_file"),
            ("rendered_pdf_sha256", "rendered_pdf"),
        ):
            data[field] = self.digest(self.root / case[path_key])
        annotation.write_text(json.dumps(data, sort_keys=True), encoding="utf-8")

    def test_identical_pdf_pages_text_and_geometry_pass(self):
        self.add_case("Contrat client Alpha 123")
        self.add_case("Budget annuel Beta 456")
        report = evaluate_office_visual_fidelity_corpus(self.write_manifest())
        self.assertEqual(report["status"], "PASS")
        self.assertEqual(report["mean_abs_pixel_error"], 0)
        self.assertEqual(report["changed_pixel_fraction"], 0)
        self.assertEqual(report["text_precision"], 1)
        self.assertEqual(report["text_recall"], 1)
        self.assertEqual(report["ordered_token_recall"], 1)
        self.assertEqual(report["mean_word_center_drift_pct"], 0)
        self.assertEqual(report["mean_word_box_iou"], 1)
        self.assertEqual(report["renderer_provenance"], "declared_not_attested")
        self.assertFalse(report["screenshot"])

    def test_shifted_render_fails_geometry_even_when_text_is_preserved(self):
        self.add_case("Contrat client Alpha 123", render_x=108)
        self.add_case("Deuxieme document Beta 456")
        report = evaluate_office_visual_fidelity_corpus(self.write_manifest())
        self.assertEqual(report["text_recall"], 1)
        self.assertEqual(report["text_precision"], 1)
        self.assertIn("max_mean_word_center_drift_pct", report["failed_thresholds"])
        self.assertIn("min_mean_word_box_iou", report["failed_thresholds"])
        self.assertEqual(report["status"], "FAIL")

    def test_missing_text_fails_recall_and_order(self):
        self.add_case("Contrat client Alpha 123", rendered_text="Contrat Alpha")
        self.add_case("Budget annuel Beta 456")
        report = evaluate_office_visual_fidelity_corpus(self.write_manifest())
        self.assertLess(report["text_recall"], 1)
        self.assertIn("min_text_recall", report["failed_thresholds"])
        self.assertIn("min_ordered_token_recall", report["failed_thresholds"])

    def test_unexpected_text_fails_precision(self):
        self.add_case("Contrat client Alpha 123", rendered_text="Contrat client Alpha 123 secretajoute")
        self.add_case("Budget annuel Beta 456")
        report = evaluate_office_visual_fidelity_corpus(self.write_manifest())
        self.assertLess(report["text_precision"], 1)
        self.assertIn("min_text_precision", report["failed_thresholds"])

    def test_extra_page_fails_page_count_threshold(self):
        self.add_case("Contrat client Alpha 123", rendered_pages=2)
        self.add_case("Budget annuel Beta 456")
        report = evaluate_office_visual_fidelity_corpus(self.write_manifest())
        self.assertEqual(report["max_page_count_delta"], 1)
        self.assertIn("max_page_count_delta", report["failed_thresholds"])

    def test_annotation_is_bound_to_rendered_pdf_hash(self):
        case = self.add_case("Contrat client Alpha 123")
        self.add_case("Budget annuel Beta 456")
        (self.root / case["rendered_pdf"]).write_bytes(b"changed evidence")
        with self.assertRaisesRegex(ValueError, "not bound"):
            evaluate_office_visual_fidelity_corpus(self.write_manifest())

    def test_duplicate_json_manifest_key_is_rejected(self):
        self.add_case("Contrat client Alpha 123")
        self.add_case("Budget annuel Beta 456")
        manifest = self.write_manifest()
        manifest.write_text('{"schema":"' + SCHEMA + '","schema":"' + SCHEMA + '","cases":[],"thresholds":{}}',
                            encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Duplicate JSON key"):
            evaluate_office_visual_fidelity_corpus(manifest)

    def test_source_symlink_is_refused(self):
        case = self.add_case("Contrat client Alpha 123")
        self.add_case("Budget annuel Beta 456")
        source = self.root / case["source_pdf"]
        target = self.root / "actual-source.pdf"
        source.replace(target)
        try:
            source.symlink_to(target)
        except (OSError, NotImplementedError) as exc:
            self.skipTest(f"symlink creation is not permitted on this runner: {type(exc).__name__}")
        with self.assertRaisesRegex(ValueError, "Symlink"):
            evaluate_office_visual_fidelity_corpus(self.write_manifest())

    def test_office_active_content_is_refused(self):
        case = self.add_case("Contrat client Alpha 123")
        self.add_case("Budget annuel Beta 456")
        office = self.root / case["office_file"]
        with zipfile.ZipFile(office, "a", compression=zipfile.ZIP_STORED) as archive:
            archive.writestr("word/vbaProject.bin", b"not executable fixture")
        self.refresh_annotation_hashes(case)
        with self.assertRaisesRegex(ValueError, "Active"):
            evaluate_office_visual_fidelity_corpus(self.write_manifest())

    def test_malformed_rendered_pdf_is_refused(self):
        case = self.add_case("Contrat client Alpha 123")
        self.add_case("Budget annuel Beta 456")
        (self.root / case["rendered_pdf"]).write_bytes(b"%PDF-1.7\ninvalid")
        self.refresh_annotation_hashes(case)
        with self.assertRaisesRegex(ValueError, "Invalid or unreadable PDF"):
            evaluate_office_visual_fidelity_corpus(self.write_manifest())

    def test_thresholds_reject_boolean_and_non_finite_values(self):
        self.add_case("Contrat client Alpha 123")
        self.add_case("Budget annuel Beta 456")
        for key, value in (
            ("max_page_count_delta", True),
            ("max_mean_abs_pixel_error", float("nan")),
            ("min_text_precision", float("inf")),
            ("min_mean_word_box_iou", 1.1),
        ):
            thresholds = dict(DEFAULT_THRESHOLDS)
            thresholds[key] = value
            with self.subTest(key=key), self.assertRaises(ValueError):
                evaluate_office_visual_fidelity_corpus(self.write_manifest(thresholds))

    def test_receipt_is_private_strict_and_never_overwrites(self):
        self.add_case("Phrase confidentielle jamais exposee 123")
        self.add_case("Autre phrase privee 456")
        manifest = self.write_manifest()
        receipt = self.root / "receipt.json"
        stdout = io.StringIO()
        with contextlib.redirect_stdout(stdout):
            self.assertEqual(main([str(manifest), "--receipt", str(receipt)]), 0)
        first_bytes = receipt.read_bytes()
        serialized = stdout.getvalue()
        self.assertNotIn("Phrase confidentielle", serialized)
        self.assertNotIn(str(self.root), serialized)
        self.assertEqual(json.loads(first_bytes)["schema"], REPORT_SCHEMA)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(main([str(manifest), "--receipt", str(receipt)]), 3)
        self.assertEqual(receipt.read_bytes(), first_bytes)

    def test_receipt_adapter_rejects_incomplete_or_forged_scope(self):
        self.add_case("Contrat client Alpha 123")
        self.add_case("Budget annuel Beta 456")
        report = evaluate_office_visual_fidelity_corpus(self.write_manifest())
        self.assertEqual(safe_visual_fidelity_receipt(report)["schema"], REPORT_SCHEMA)
        forged = dict(report)
        forged["screenshot"] = True
        with self.assertRaises(ValueError):
            safe_visual_fidelity_receipt(forged)
        self.assertIn('"renderer_provenance": "declared_not_attested"',
                      serialize_visual_fidelity_receipt(report))


if __name__ == "__main__":
    unittest.main(verbosity=2)
