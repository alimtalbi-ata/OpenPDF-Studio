"""V533: per-scanned-page OCR editable cell occupancy and strict gate.
Synthetic raster PDF tests only. Author: Abdelhalim Talbi.
"""
from __future__ import annotations
import hashlib
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest.mock import patch
import fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
from openpdf_ocr_office_pipeline_v511 import convert_scanned_office, main
from openpdf_ocr_language_preflight_v486 import preflight_ocr_languages

TABLE = (('ID', 'STOCK', 'PRICE'), ('A01', '102', '55'),
         ('A02', '103', '80'), ('A03', '104', '75'))

def english_ocr_available():
    try:
        preflight_ocr_languages('eng')
        return True
    except (OSError, RuntimeError, ValueError):
        return False

def make_pdf(path, *, scanned_pages=2, native_first=False):
    with fitz.open() as pdf:
        if native_first:
            native = pdf.new_page(width=600, height=400)
            for r, row in enumerate(TABLE):
                for c, value in enumerate(row):
                    native.insert_text((40 + 140*c, 80 + 45*r), value, fontsize=16)
        for _ in range(scanned_pages):
            with fitz.open() as inner:
                page = inner.new_page(width=600, height=400)
                for r, row in enumerate(TABLE):
                    for c, value in enumerate(row):
                        page.insert_text((40 + 140*c, 80 + 45*r), value, fontsize=16)
                png = page.get_pixmap(matrix=fitz.Matrix(2,2), alpha=False).tobytes('png')
            scanned = pdf.new_page(width=600, height=400)
            scanned.insert_image(scanned.rect, stream=png)
        pdf.save(path)

class OCRTableCellDensityV533(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.pdf = self.root / 'source.pdf'
        self.xlsx = self.root / 'converted.xlsx'

    def test_invalid_policy_rejected_before_ocr(self):
        make_pdf(self.pdf, scanned_pages=1)
        for invalid in (-1, 10001, True, '12', 1.5, None):
            with self.subTest(value=invalid), self.assertRaisesRegex(ValueError, 'Seuil de cellules OCR'):
                convert_scanned_office(self.pdf, self.xlsx, language='eng',
                    table_mode=True, min_nonempty_cells_per_scanned_page=invalid)
        with self.assertRaisesRegex(ValueError, 'mode tableaux XLSX absent'):
            convert_scanned_office(self.pdf, self.xlsx, language='eng',
                min_nonempty_cells_per_scanned_page=1)
        self.assertFalse(self.xlsx.exists())
        self.assertFalse(list(self.root.glob('.openpdf-*')))

    @unittest.skipUnless(english_ocr_available(), 'English OCR model unavailable')
    def test_two_page_ocr_reports_exact_nonempty_cells(self):
        make_pdf(self.pdf)
        before = hashlib.sha256(self.pdf.read_bytes()).hexdigest()
        receipt = convert_scanned_office(self.pdf, self.xlsx, language='eng',
            table_mode=True, require_table_per_scanned_page=True,
            min_nonempty_cells_per_scanned_page=12)
        self.assertEqual(receipt.table_cells_by_page, ((1,12,12), (2,12,12)))
        self.assertEqual(receipt.to_dict()['table_cells_by_page'], [
            {'page':1,'cells':12,'nonempty_cells':12},
            {'page':2,'cells':12,'nonempty_cells':12}])
        self.assertEqual(hashlib.sha256(self.pdf.read_bytes()).hexdigest(), before)
        self.assertEqual(hashlib.sha256(self.xlsx.read_bytes()).hexdigest(), receipt.sha256)
        book = load_workbook(self.xlsx)
        try:
            for sheet in book.worksheets[1:]:
                self.assertEqual(tuple(tuple(sheet.cell(r,c).value for c in range(1,4))
                    for r in range(3,7)), TABLE)
        finally:
            book.close()

    @unittest.skipUnless(english_ocr_available(), 'English OCR model unavailable')
    def test_excessive_minimum_fails_without_output(self):
        make_pdf(self.pdf, scanned_pages=1)
        with self.assertRaisesRegex(ValueError, 'Cellules OCR insuffisantes sur page scannee : 1'):
            convert_scanned_office(self.pdf, self.xlsx, language='eng', table_mode=True,
                min_nonempty_cells_per_scanned_page=13)
        self.assertFalse(self.xlsx.exists())
        self.assertFalse(list(self.root.glob('.openpdf-*')))

    @unittest.skipUnless(english_ocr_available(), 'English OCR model unavailable')
    def test_native_page_not_subject_to_scanned_minimum(self):
        make_pdf(self.pdf, scanned_pages=1, native_first=True)
        receipt = convert_scanned_office(self.pdf, self.xlsx, language='eng', table_mode=True,
            min_nonempty_cells_per_scanned_page=12)
        self.assertEqual(receipt.scanned_pages, (2,))
        self.assertEqual(receipt.table_cells_by_page, ((1,12,12), (2,12,12)))

    @unittest.skipUnless(english_ocr_available(), 'English OCR model unavailable')
    def test_cli_minimum_is_applied(self):
        make_pdf(self.pdf, scanned_pages=1)
        with self.assertRaisesRegex(ValueError, 'Cellules OCR insuffisantes'):
            main([str(self.pdf), str(self.xlsx), '--language', 'eng', '--editable-tables',
                  '--min-nonempty-table-cells-per-scanned-page', '13'])
        self.assertFalse(self.xlsx.exists())

class GUIAndDispatcherV533(unittest.TestCase):
    def test_gui_exposes_explicit_opt_in_and_validation(self):
        gui = (ROOT / '03-Sources/openpdf_desktop.py').read_text(encoding='utf-8')
        self.assertIn("'Cellules OCR min/page'", gui)
        self.assertIn("kwargs['min_nonempty_cells_per_scanned_page']=int(raw)", gui)
        self.assertIn('self.ocr_before_office.get() and self.xlsx_tables.get()', gui)

    def test_dispatcher_forwards_minimum_only_to_ocr_excel(self):
        import openpdf_operations as operations
        fake = types.ModuleType('openpdf_ocr_office_pipeline_v511')
        received = []
        def convert(*args, **kwargs):
            received.append(kwargs)
            return 'OK'
        fake.convert_scanned_office = convert
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'source.pdf'
            source.write_bytes(b'%PDF-synthetic-dispatch-fixture')
            destination = Path(directory) / 'output.xlsx'
            with patch.dict(sys.modules, {'openpdf_ocr_office_pipeline_v511':fake}), \
                 patch.object(operations, 'write_event'), \
                 patch.object(operations, '_legacy_run_operation', side_effect=AssertionError):
                self.assertEqual(operations.run_operation('xlsx', (str(source),), str(destination),
                    table_mode=True, ocr_before_office=True, language='eng',
                    min_nonempty_cells_per_scanned_page=12), 'OK')
        self.assertEqual(received, [{
            'language':'eng', 'table_mode':True,
            'min_nonempty_cells_per_scanned_page':12}])

    def test_dispatcher_rejects_non_ocr_minimum(self):
        import openpdf_operations as operations
        for opts in ({'table_mode':True, 'ocr_before_office':False},
                     {'table_mode':False, 'ocr_before_office':True}):
            with self.subTest(opts=opts), self.assertRaisesRegex(ValueError, 'Seuil de cellules OCR'):
                operations.run_operation('xlsx', ('no.pdf',), 'no.xlsx',
                    min_nonempty_cells_per_scanned_page=2, **opts)

if __name__ == '__main__':
    unittest.main()
