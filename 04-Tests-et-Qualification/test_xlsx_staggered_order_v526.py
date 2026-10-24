"""V526: real PDF staggered table columns and bounded reading bands.

Synthetic fixtures are not independently human annotated or GUI proof.
Author: Abdelhalim Talbi.
"""
from contextlib import closing
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
import fitz
from openpyxl import load_workbook
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_tables_v514 import _reading_order, _tables, verify_xlsx_tables


def staggered_pdf(path):
    pdf = fitz.open()
    page = pdf.new_page(width=900, height=700)
    # Three real 4x3 ruled tables. Neighboring boxes overlap vertically,
    # but the first and third do not overlap by half their height.
    for label, x, y in [('A', 620, 100), ('B', 330, 125), ('C', 40, 150)]:
        for row in range(5):
            page.draw_line((x, y + row * 20), (x + 240, y + row * 20))
        for col in range(4):
            page.draw_line((x + col * 80, y), (x + col * 80, y + 80))
        for row in range(4):
            for col in range(3):
                page.insert_text((x + col * 80 + 5, y + row * 20 + 13),
                                 f'{label}{row}{col}', fontsize=8)
    pdf.save(path)
    pdf.close()


class StaggeredReadingOrderV526(unittest.TestCase):
    def test_real_pdf_and_editable_excel_left_to_right(self):
        with tempfile.TemporaryDirectory() as temp:
            pdf = Path(temp) / 'stagger.pdf'
            xlsx = Path(temp) / 'stagger.xlsx'
            staggered_pdf(pdf)
            with fitz.open(pdf) as document:
                self.assertEqual(len(document[0].find_tables().tables), 3)
            self.assertEqual([rows[0][0] for _, _, rows, _, _ in _tables(pdf)],
                             ['C00', 'B00', 'A00'])
            convert_office(pdf, xlsx, table_mode=True)
            evidence = verify_xlsx_tables(pdf, xlsx)
            self.assertEqual((evidence.tables, evidence.cells), (3, 36))
            with closing(load_workbook(xlsx)) as workbook:
                self.assertEqual([workbook.worksheets[i]['A3'].value for i in (1, 2, 3)],
                                 ['C00', 'B00', 'A00'])
                self.assertEqual(workbook.worksheets[3]['C6'].value, 'A32')
            with closing(load_workbook(xlsx)) as workbook:
                workbook.worksheets[1]['B4'] = 'CORRUPTED'
                workbook.save(xlsx)
            with self.assertRaises(ValueError):
                verify_xlsx_tables(pdf, xlsx)

    def test_far_separated_rows_do_not_join_band(self):
        def candidate(label, box):
            return ([[label]], None, None, box)
        items = [candidate('top', (400, 100, 500, 180)),
                 candidate('bottom', (20, 280, 120, 360))]
        self.assertEqual([x[0][0][0] for x in _reading_order(items)],
                         ['top', 'bottom'])

    def test_diagonal_chain_has_bounded_height(self):
        def candidate(label, x, y):
            return ([[label]], None, None, (x, y, x + 90, y + 80))
        items = [candidate(str(i), 600 - i * 100, 100 + i * 25)
                 for i in range(6)]
        result = _reading_order(items)
        # An unbounded transitive band would reverse all six; the cap must
        # keep the bottommost column outside the first horizontal band.
        self.assertNotEqual(result[0][0][0][0], '5')
        self.assertEqual(sorted(x[0][0][0] for x in result), list('012345'))

    def test_invalid_geometry_fails_closed(self):
        with self.assertRaisesRegex(ValueError, 'Geometrie'):
            _reading_order([([['x']], None, None, (0, 20, 30, 20))])


if __name__ == '__main__':
    unittest.main()
