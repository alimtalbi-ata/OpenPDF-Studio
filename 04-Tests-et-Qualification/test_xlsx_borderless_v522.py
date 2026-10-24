"""V522 real-PDF regression for conservative borderless table extraction.

Synthetic source truth is known at PDF creation, not human annotated.
Author: Abdelhalim Talbi.
"""
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
import fitz
from openpyxl import load_workbook
from contextlib import closing
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_tables_v514 import verify_xlsx_tables


def make_pdf(path, *, pages=2, rows=4, columns=3, prose=False, injection=False):
    pdf = fitz.open()
    for page_number in range(pages):
        page = pdf.new_page(width=595, height=842)
        page.insert_text((50, 40), f'Inventaire stockage - page {page_number + 1}')
        if prose:
            for r in range(4):
                page.insert_text((50, 90 + 25*r), f'Paragraphe {r}: une phrase sans colonnes de tableau.')
        else:
            for r in range(rows):
                for c in range(columns):
                    value = f'P{page_number+1}-R{r}-C{c}'
                    if injection and r == rows - 1 and c == columns - 1:
                        value = '=1+1'
                    page.insert_text((50 + c*155, 100 + r*25), value, fontsize=10)
    pdf.save(path)
    pdf.close()


class TestBorderlessV522(unittest.TestCase):
    def test_two_page_borderless_tables_are_editable_without_spacer_rows(self):
        with tempfile.TemporaryDirectory() as d:
            pdf, xlsx = Path(d)/'input.pdf', Path(d)/'output.xlsx'
            make_pdf(pdf)
            convert_office(pdf, xlsx, table_mode=True)
            evidence = verify_xlsx_tables(pdf, xlsx)
            self.assertEqual((evidence.tables, evidence.pages_with_tables, evidence.cells), (2, (1, 2), 24))
            with closing(load_workbook(xlsx)) as wb:
                self.assertEqual(wb.sheetnames, ['Texte PDF', 'P0001-T001', 'P0002-T001'])
                self.assertEqual(wb.worksheets[1]['A3'].value, 'P1-R0-C0')
                self.assertEqual(wb.worksheets[2]['C6'].value, 'P2-R3-C2')
                self.assertEqual(wb.worksheets[1].max_row, 6)

    def test_borderless_formula_stays_literal(self):
        with tempfile.TemporaryDirectory() as d:
            pdf, xlsx = Path(d)/'input.pdf', Path(d)/'output.xlsx'
            make_pdf(pdf, pages=1, injection=True)
            convert_office(pdf, xlsx, table_mode=True)
            with closing(load_workbook(xlsx)) as wb:
                cell = wb.worksheets[1]['C6']
                self.assertEqual((cell.value, cell.data_type), ('=1+1', 's'))

    def test_no_borderless_table_does_not_publish(self):
        with tempfile.TemporaryDirectory() as d:
            pdf, xlsx = Path(d)/'input.pdf', Path(d)/'output.xlsx'
            make_pdf(pdf, pages=1, prose=True)
            with self.assertRaises(ValueError):
                convert_office(pdf, xlsx, table_mode=True)
            self.assertFalse(xlsx.exists())

    def test_two_row_alignment_is_not_sufficient(self):
        with tempfile.TemporaryDirectory() as d:
            pdf, xlsx = Path(d)/'input.pdf', Path(d)/'output.xlsx'
            make_pdf(pdf, pages=1, rows=2)
            with self.assertRaises(ValueError):
                convert_office(pdf, xlsx, table_mode=True)
            self.assertFalse(xlsx.exists())

    def test_legacy_mode_preserves_single_text_sheet(self):
        with tempfile.TemporaryDirectory() as d:
            pdf, xlsx = Path(d)/'input.pdf', Path(d)/'output.xlsx'
            make_pdf(pdf, pages=1)
            convert_office(pdf, xlsx)
            with closing(load_workbook(xlsx)) as wb:
                self.assertEqual(wb.sheetnames, ['Texte PDF'])

    def test_mutated_borderless_cell_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            pdf, xlsx = Path(d)/'input.pdf', Path(d)/'output.xlsx'
            make_pdf(pdf, pages=1)
            convert_office(pdf, xlsx, table_mode=True)
            wb = load_workbook(xlsx)
            wb.worksheets[1]['B4'] = 'FAUX'
            wb.save(xlsx)
            wb.close()
            with self.assertRaises(ValueError):
                verify_xlsx_tables(pdf, xlsx)


if __name__ == '__main__':
    unittest.main()
