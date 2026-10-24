"""V51.5 real PDF/XLSX fidelity and adversarial archive tests.
Author: Abdelhalim Talbi. No Windows UI qualification claimed.
"""
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
sys.path.insert(0, str(ROOT / '04-Tests-et-Qualification'))
from openpyxl import load_workbook
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_fidelity_v515 import audit_xlsx_fidelity, write_report_new, main
from test_xlsx_tables_v514 import fixture


class TestXlsxFidelityV515(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.pdf = Path(self.tmp.name) / 'input.pdf'
        self.xlsx = Path(self.tmp.name) / 'output.xlsx'
        fixture(self.pdf)
        convert_office(self.pdf, self.xlsx, table_mode=True)

    def test_real_conversion_evidence(self):
        result = audit_xlsx_fidelity(self.pdf, self.xlsx)
        self.assertEqual(result['status'], 'PASS')
        self.assertEqual(result['tables'], 1)
        self.assertEqual(result['grid_cells'], 6)
        self.assertEqual(result['nonempty_cells'], 6)
        self.assertEqual(result['page_table_coverage'], 1.0)
        self.assertEqual(result['semantic_fidelity'], 'not_assessed')
        self.assertEqual(len(result['source_sha256']), 64)
        self.assertNotIn('SAN-A', json.dumps(result))

    def test_tampered_cell_rejected(self):
        wb = load_workbook(self.xlsx)
        wb.worksheets[1]['A4'] = 'altered'
        wb.save(self.xlsx)
        wb.close()
        with self.assertRaises(ValueError):
            audit_xlsx_fidelity(self.pdf, self.xlsx)

    def test_formula_in_legacy_sheet_rejected(self):
        wb = load_workbook(self.xlsx)
        wb.worksheets[0]['A20'] = '=WEBSERVICE("example")'
        wb.save(self.xlsx)
        wb.close()
        with self.assertRaisesRegex(ValueError, 'Formule'):
            audit_xlsx_fidelity(self.pdf, self.xlsx)

    def test_hyperlink_in_legacy_sheet_rejected(self):
        wb = load_workbook(self.xlsx)
        wb.worksheets[0]['A20'].hyperlink = 'https://example.invalid'
        wb.save(self.xlsx)
        wb.close()
        with self.assertRaisesRegex(ValueError, 'hyperlien|externe'):
            audit_xlsx_fidelity(self.pdf, self.xlsx)

    def test_external_link_archive_rejected(self):
        with zipfile.ZipFile(self.xlsx, 'a') as archive:
            archive.writestr('xl/externalLinks/externalLink1.xml', '<x/>')
        with self.assertRaisesRegex(ValueError, 'externe'):
            audit_xlsx_fidelity(self.pdf, self.xlsx)

    def test_duplicate_zip_member_rejected(self):
        with zipfile.ZipFile(self.xlsx, 'a') as archive:
            archive.writestr('xl/workbook.xml', '<x/>')
        with self.assertRaisesRegex(ValueError, 'duplique'):
            audit_xlsx_fidelity(self.pdf, self.xlsx)

    def test_source_symlink_rejected(self):
        linked = Path(self.tmp.name) / 'linked.pdf'
        linked.symlink_to(self.pdf)
        with self.assertRaisesRegex(ValueError, 'symbolique'):
            audit_xlsx_fidelity(linked, self.xlsx)

    def test_failed_audit_blocks_publication(self):
        destination = Path(self.tmp.name) / 'second.xlsx'
        with patch('openpdf_xlsx_fidelity_v515.audit_xlsx_fidelity',
                   side_effect=ValueError('Audit refuse')):
            with self.assertRaisesRegex(ValueError, 'Audit refuse'):
                convert_office(self.pdf, destination, table_mode=True)
        self.assertFalse(destination.exists())

    def test_sparse_legacy_sheet_does_not_expand_rows(self):
        wb = load_workbook(self.xlsx)
        wb.worksheets[0]['A1048576'] = 'last row'
        wb.save(self.xlsx)
        wb.close()
        self.assertEqual(audit_xlsx_fidelity(self.pdf, self.xlsx)['status'], 'PASS')

    def test_atomic_report_and_no_overwrite(self):
        evidence = audit_xlsx_fidelity(self.pdf, self.xlsx)
        report = Path(self.tmp.name) / 'audit.json'
        write_report_new(report, evidence)
        self.assertEqual(json.loads(report.read_text())['tables'], 1)
        with self.assertRaises(FileExistsError):
            write_report_new(report, {'status': 'modified'})
        self.assertEqual(json.loads(report.read_text())['status'], 'PASS')

    def test_cli_success_and_fail_exit(self):
        report = Path(self.tmp.name) / 'audit.json'
        self.assertEqual(main(['--pdf', str(self.pdf), '--xlsx', str(self.xlsx),
                               '--report', str(report)]), 0)
        self.assertEqual(main(['--pdf', str(self.pdf), '--xlsx', str(self.xlsx),
                               '--report', str(report)]), 2)


if __name__ == '__main__':
    unittest.main()
