"""V524: mixed PDF table worksheet ordering and mutation checks.

Synthetic PDFs only, not human-annotated production documents.
Author: Abdelhalim Talbi.
"""
from contextlib import closing
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
import fitz
from openpyxl import load_workbook
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_tables_v514 import _tables, verify_xlsx_tables


def make_reverse_mixed_pdf(path, *, side_by_side=False):
    doc = fitz.open()
    page = doc.new_page(width=842 if side_by_side else 595, height=842)
    # The text-only table precedes the ruled table in visual reading order.
    bx, by = 35, 105
    rx, ry = (400, 90) if side_by_side else (35, 380)
    for r in range(4):
        for c in range(3):
            page.insert_text((bx + c * 95, by + r * 27),
                             f'TEXT-{r}-{c}', fontsize=8)
    for r in range(5):
        page.draw_line((rx, ry + r * 27), (rx + 285, ry + r * 27))
    for c in range(4):
        page.draw_line((rx + c * 95, ry), (rx + c * 95, ry + 108))
    for r in range(4):
        for c in range(3):
            page.insert_text((rx + c * 95 + 5, ry + r * 27 + 17),
                             f'RULE-{r}-{c}', fontsize=9)
    doc.save(path)
    doc.close()


class ReadingOrderV524(unittest.TestCase):
    def _check(self, side_by_side):
        with tempfile.TemporaryDirectory() as tmp:
            pdf, xlsx = Path(tmp) / 'mixed.pdf', Path(tmp) / 'mixed.xlsx'
            make_reverse_mixed_pdf(pdf, side_by_side=side_by_side)
            with fitz.open(pdf) as doc:
                self.assertEqual(len(doc[0].find_tables().tables), 1)
            actual = list(_tables(pdf))
            self.assertEqual([(p, i, rows[0][0]) for p, i, rows, _, _ in actual],
                             [(1, 1, 'TEXT-0-0'), (1, 2, 'RULE-0-0')])
            convert_office(pdf, xlsx, table_mode=True)
            self.assertEqual((verify_xlsx_tables(pdf, xlsx).tables,
                              verify_xlsx_tables(pdf, xlsx).cells), (2, 24))
            with closing(load_workbook(xlsx)) as wb:
                self.assertEqual(wb.sheetnames,
                                 ['Texte PDF', 'P0001-T001', 'P0001-T002'])
                self.assertEqual(wb.worksheets[1]['A3'].value, 'TEXT-0-0')
                self.assertEqual(wb.worksheets[2]['A3'].value, 'RULE-0-0')
            with closing(load_workbook(xlsx)) as wb:
                wb.worksheets[1]['A3'] = 'CORRUPTED'
                wb.save(xlsx)
            with self.assertRaises(ValueError):
                verify_xlsx_tables(pdf, xlsx)

    def test_text_above_ruled_table(self):
        self._check(side_by_side=False)

    def test_text_left_of_ruled_table(self):
        self._check(side_by_side=True)


if __name__ == '__main__':
    unittest.main()
