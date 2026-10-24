"""V532: real OCR page coverage and strict fail-closed mode, synthetic fixtures.

Author: Abdelhalim Talbi. Does not claim recognition accuracy or human QA.
"""
from __future__ import annotations

import hashlib
from pathlib import Path
import sys
import tempfile
import unittest

import fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
from openpdf_ocr_office_pipeline_v511 import convert_scanned_office
from test_ocr_office_tables_v531 import english_ocr_available, TABLE


def scanned_pdf(path, *, prose_second=False, native_first=False):
    with fitz.open() as output:
        if native_first:
            output.new_page(width=600, height=400).insert_text(
                (50, 70), 'Native editable introduction', fontsize=16)
        for page_number in range(2):
            with fitz.open() as native:
                page = native.new_page(width=600, height=400)
                if prose_second and page_number == 1:
                    page.insert_text((40, 80), 'INVOICE TOTAL 12345', fontsize=16)
                else:
                    for row, values in enumerate(TABLE):
                        for col, value in enumerate(values):
                            page.insert_text((40 + col * 140, 80 + row * 45),
                                             value, fontsize=16)
                png = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes('png')
            scanned = output.new_page(width=600, height=400)
            scanned.insert_image(scanned.rect, stream=png)
        output.save(path)


class TableCoverageV532(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.pdf = self.root / 'source.pdf'
        self.xlsx = self.root / 'output.xlsx'

    def test_invalid_coverage_policy_rejected_before_output(self):
        scanned_pdf(self.pdf)
        for kwargs in ({'require_table_per_scanned_page': 1},
                       {'require_table_per_scanned_page': 'yes'},
                       {'require_table_per_scanned_page': True}):
            with self.subTest(kwargs=kwargs), self.assertRaisesRegex(ValueError, 'Couverture stricte'):
                convert_scanned_office(self.pdf, self.xlsx, table_mode=False, **kwargs)
        self.assertFalse(self.xlsx.exists())

    @unittest.skipUnless(english_ocr_available(), 'Offline English OCR unavailable')
    def test_two_scanned_pages_two_editable_tables(self):
        scanned_pdf(self.pdf)
        before = hashlib.sha256(self.pdf.read_bytes()).hexdigest()
        receipt = convert_scanned_office(self.pdf, self.xlsx, language='eng', table_mode=True,
                                         require_table_per_scanned_page=True)
        self.assertEqual(receipt.scanned_pages, (1, 2))
        self.assertEqual(receipt.table_pages, (1, 2))
        self.assertEqual(receipt.scanned_pages_without_tables, ())
        self.assertEqual(receipt.table_worksheets, 2)
        self.assertEqual(receipt.to_dict()['table_pages'], [1, 2])
        self.assertEqual(hashlib.sha256(self.pdf.read_bytes()).hexdigest(), before)
        self.assertEqual(hashlib.sha256(self.xlsx.read_bytes()).hexdigest(), receipt.sha256)
        wb = load_workbook(self.xlsx)
        try:
            self.assertEqual(wb.sheetnames, ['Texte PDF', 'P0001-T001', 'P0002-T001'])
            for sheet in wb.worksheets[1:]:
                self.assertEqual(tuple(tuple(sheet.cell(r, c).value for c in range(1, 4))
                                       for r in range(3, 7)), TABLE)
        finally:
            wb.close()

    @unittest.skipUnless(english_ocr_available(), 'Offline English OCR unavailable')
    def test_prose_scanned_page_strictly_rejected_without_output(self):
        scanned_pdf(self.pdf, prose_second=True)
        with self.assertRaisesRegex(ValueError, 'page scannee : 2'):
            convert_scanned_office(self.pdf, self.xlsx, language='eng', table_mode=True,
                                   require_table_per_scanned_page=True)
        self.assertFalse(self.xlsx.exists())
        self.assertEqual(list(self.root.glob('.openpdf-*')), [])

    @unittest.skipUnless(english_ocr_available(), 'Offline English OCR unavailable')
    def test_legacy_mode_warns_on_uncovered_scanned_page(self):
        scanned_pdf(self.pdf, prose_second=True, native_first=True)
        receipt = convert_scanned_office(self.pdf, self.xlsx, language='eng', table_mode=True)
        self.assertEqual(receipt.scanned_pages, (2, 3))
        self.assertEqual(receipt.table_pages, (2,))
        self.assertEqual(receipt.scanned_pages_without_tables, (3,))
        self.assertEqual(receipt.table_worksheets, 1)
        self.assertTrue(any('3' in warning for warning in receipt.warnings))
        self.assertTrue(self.xlsx.is_file())


if __name__ == '__main__':
    unittest.main()
