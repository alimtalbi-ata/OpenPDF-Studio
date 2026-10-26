"""Generate an exact-commit V550 Word/PDF guide; never reuse stale CI counts.
Author: Abdelhalim Talbi.
"""
from datetime import datetime, timezone
import json
import os
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
REPORT = "04-Tests-et-Qualification/V550-CI-QUALIFICATION.json"
SCHEMA = "openpdf.v550.ci-qualification.v1"
ALLOWED_WINDOWS_SKIPS = {
    "test_final_symlink_staging_entry_is_rejected",
    "test_link_in_parent_chain_is_rejected_before_publication",
    "test_dotdot_does_not_hide_a_symlink_component",
    "test_parent_replacement_during_link_does_not_redirect_write",
    "test_temporary_entry_swap_cannot_change_inode_published_on_linux",
}


def load_current_report(root=ROOT, expected_head_sha=None):
    path = Path(root) / REPORT
    if not path.is_file() or path.is_symlink():
        return None
    try:
        report = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        return None
    head = expected_head_sha or os.environ.get("OPENPDF_HEAD_SHA", "").strip() or os.environ.get("GITHUB_SHA", "").strip()
    if (not isinstance(report, dict) or report.get("schema") != SCHEMA
            or report.get("version") != "51.4" or report.get("milestone") != "V550"
            or report.get("qualified") is not True or not report.get("commit_sha")):
        return None
    if head and report.get("commit_sha") != head:
        return None
    page = report.get("page_order")
    if not isinstance(page, dict) or page.get("head_sha") != report.get("commit_sha"):
        return None
    if page.get("status") != "completed" or page.get("conclusion") != "success":
        return None
    tests = page.get("tests")
    if not isinstance(tests, dict) or set(tests) != {"ubuntu", "windows"}:
        return None
    u, w = tests["ubuntu"], tests["windows"]
    if (not isinstance(u, dict) or u.get("passed") != u.get("expected") or
            u.get("expected") != 11 or u.get("skipped") != 0 or u.get("failed") != 0):
        return None
    if (not isinstance(w, dict) or w.get("expected") != 11 or w.get("passed") != 6 or
            w.get("skipped") != 5 or w.get("failed") != 0 or
            set(w.get("skipped_names", [])) != ALLOWED_WINDOWS_SKIPS):
        return None
    parent = report.get("parent_v549")
    if not isinstance(parent, dict) or parent.get("qualified") is not True or parent.get("commit_sha") != report.get("commit_sha"):
        return None
    native = parent.get("native_windows")
    if (not isinstance(native, dict) or native.get("portable", {}).get("smoke_passed") is not True
            or native.get("installer", {}).get("smoke_passed") is not True
            or native.get("installation", {}).get("passed") is not True):
        return None
    return report


def build_sections(report):
    if report is None:
        qualification = "NON VÉRIFIÉE : aucun reçu V550 qualifié correspondant exactement à ce commit. Aucun résultat historique n'est repris."
    else:
        u, w = report["page_order"]["tests"]["ubuntu"], report["page_order"]["tests"]["windows"]
        qualification = (
            f"Commit qualifié : {report['commit_sha']} ; run CI {report['page_order']['url']}. "
            f"Ubuntu : {u['passed']}/{u['expected']} réussis, {u['skipped']} ignoré, {u['failed']} échec. "
            f"Windows : {w['passed']}/{w['expected']} réussis, {w['skipped']} ignorés attendus pour les primitives POSIX/Linux, {w['failed']} échec. "
            f"La qualification Windows native des binaires provient du reçu V549 exact du même SHA."
        )
    return [
        ("Résumé", "V550 remplace la publication atomique basée uniquement sur des chemins par une publication ancrée sur des handles POSIX lorsqu'elle est prise en charge."),
        ("Détails techniques", "Les composants du répertoire sont ouverts un à un avec O_DIRECTORY et O_NOFOLLOW. Le descripteur du parent est conservé pendant le lien dur exclusif. Le temporaire est ouvert sans suivre les liens ; Linux utilise l'inode déjà ouvert via /proc/self/fd si cette primitive est disponible. L'inode publié est comparé et le répertoire synchronisé. Le chemin est ensuite réouvert pour vérifier qu'il désigne encore le répertoire épinglé."),
        ("Qualification liée au commit", qualification),
        ("Portée Windows", "Le chemin Windows conserve le précontrôle des reparse points de V549 et un lien dur exclusif avec vérification d'identité après publication. Ce chemin de compatibilité n'utilise pas un parcours parent équivalent par handles : la suppression complète de TOCTOU sous Windows n'est pas revendiquée."),
        ("Gestion d'une course", "Si le dossier est déplacé ou remplacé, le primitive signale une ambiguïté et préserve la sortie publiée pour revue manuelle lorsque son retrait ne peut pas être effectué de façon sûre. Il ne supprime pas à l'aveugle un nom qui pourrait désormais désigner un autre fichier."),
        ("Limites", "Les tests contrôlent les courses par injection déterministe et manipulation temporaire du système de fichiers sous Linux ; ils ne représentent pas un test d'intrusion exhaustif. Ils ne certifient ni la précision OCR, ni la fidélité visuelle Office, ni une revue graphique humaine Windows, ni une signature Authenticode."),
        ("Prochaine priorité", "Implémenter une variante de publication Windows fondée sur des handles natifs et des opérations relatives au handle, puis la qualifier avec des tests de course sur runner Windows. Conserver des résultats explicitement non certifiés jusqu'à ce que cette preuve existe."),
    ]


def build(root=ROOT):
    root = Path(root).resolve()
    report = load_current_report(root)
    sections = build_sections(report)
    version_title = "OpenPDF Studio V51.4 — V550 : publication ancrée"
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    out_docx = root / "05-Documentation" / "OpenPDF-Studio-V51.4-V550-Guide.docx"
    out_pdf = root / "05-Documentation" / "OpenPDF-Studio-V51.4-V550-Guide.pdf"

    from docx import Document
    doc = Document()
    doc.add_heading(version_title, 0)
    doc.add_paragraph("Auteur : Abdelhalim Talbi")
    doc.add_paragraph("Généré : " + generated)
    doc.add_paragraph("Version applicative : 51.4 ; jalon : V550.")
    if report:
        doc.add_paragraph("Commit : " + report["commit_sha"])
        if report.get("pull_request"):
            doc.add_paragraph("PR : https://github.com/" + report["repository"] + "/pull/" + str(report["pull_request"]))
    else:
        doc.add_paragraph("Qualification CI de cette révision : NON VÉRIFIÉE.")
    doc.add_paragraph("Le diagramme SVG associé décrit une architecture technique et n'est pas une capture d'écran.")
    for heading, body in sections:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)
    doc.save(out_docx)

    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="V550Body", parent=styles["BodyText"], leading=14, spaceAfter=8))
    flow = [Paragraph(escape(version_title), styles["Title"]), Spacer(1, 8),
            Paragraph(escape("Auteur : Abdelhalim Talbi | " + generated), styles["BodyText"]),
            Paragraph(escape("Version applicative 51.4 — jalon V550"), styles["BodyText"]), Spacer(1, 12)]
    if report:
        rows = [
            ["Commit", report["commit_sha"]],
            ["CI Ubuntu", f"{report['page_order']['tests']['ubuntu']['passed']}/11 PASS, 0 skip"],
            ["CI Windows", f"{report['page_order']['tests']['windows']['passed']}/11 PASS, 5 skips POSIX/Linux attendus"],
            ["Run", report["page_order"]["url"] or "URL absente"],
        ]
        table = Table([[Paragraph(escape(str(c)), styles["BodyText"]) for c in row] for row in rows], colWidths=[95, 390])
        table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.5,colors.grey),("BACKGROUND",(0,0),(0,-1),colors.whitesmoke),("VALIGN",(0,0),(-1,-1),"TOP"),("PADDING",(0,0),(-1,-1),6)]))
        flow.extend([table, Spacer(1, 12)])
    for heading, body in sections:
        flow.append(Paragraph(escape(heading), styles["Heading2"]))
        flow.append(Paragraph(escape(body), styles["V550Body"]))
    SimpleDocTemplate(str(out_pdf), pagesize=A4, title=version_title, author="Abdelhalim Talbi").build(flow)
    return {"docx": str(out_docx), "pdf": str(out_pdf), "qualified": report is not None,
            "commit_sha": report.get("commit_sha") if report else None}


if __name__ == "__main__":
    print(json.dumps(build(), ensure_ascii=False, indent=2))
