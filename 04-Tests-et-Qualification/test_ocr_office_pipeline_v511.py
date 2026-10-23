"""V51.1: functional OCR->Office tests; synthetic documents, no Windows GUI claim.
Author: Abdelhalim Talbi.
"""
import hashlib
import io
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import fitz
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
import openpdf_ocr_office_pipeline_v511 as pipeline
from openpdf_ocr_language_preflight_v486 import preflight_ocr_languages


def ocr_ready():
    try:
        preflight_ocr_languages('eng')
        return True
    except (OSError, RuntimeError, ValueError):
        return False


class OCRToOfficeV511(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.source = self.root / 'input.pdf'
        self.destination = self.root / 'output.docx'

    def make_pdf(self, *, scanned=True, mixed=False):
        with fitz.open() as doc:
            if mixed:
                doc.new_page().insert_text((72, 100), 'Native editable source text')
            page = doc.new_page(width=612, height=792)
            if scanned:
                img = Image.new('RGB', (1200, 1500), 'white')
                d = ImageDraw.Draw(img)
                try:
                    font = ImageFont.truetype('DejaVuSans.ttf', 56)
                except OSError:
                    font = ImageFont.load_default()
                d.text((80, 180), 'INVOICE NUMBER 12345', font=font, fill='black')
                d.text((80, 300), 'TOTAL AMOUNT 67890', font=font, fill='black')
                data = io.BytesIO()
                img.save(data, 'PNG')
                page.insert_image(page.rect, stream=data.getvalue())
            else:
                page.insert_text((72, 100), 'Native editable source text')
            doc.save(self.source)
        return hashlib.sha256(self.source.read_bytes()).hexdigest()

    def test_native_pdf_converts_without_ocr(self):
        before = self.make_pdf(scanned=False)
        with patch.object(pipeline, 'searchable_pdf', side_effect=AssertionError('OCR called')):
            receipt = pipeline.convert_scanned_office(self.source, self.destination)
        self.assertEqual(receipt.ocr_pages, 0)
        self.assertEqual(receipt.scanned_pages, ())
        self.assertEqual(receipt.sha256, hashlib.sha256(self.destination.read_bytes()).hexdigest())
        self.assertEqual(before, hashlib.sha256(self.source.read_bytes()).hexdigest())

    @unittest.skipUnless(ocr_ready(), "local eng Tesseract model unavailable")
    def test_scanned_pdf_real_ocr_docx(self):
        before = self.make_pdf()
        receipt = pipeline.convert_scanned_office(self.source, self.destination, language='eng')
        self.assertEqual(receipt.scanned_pages, (1,))
        self.assertEqual(receipt.ocr_pages, 1)
        self.assertGreater(receipt.recognized_characters, 10)
        self.assertEqual(receipt.sha256, hashlib.sha256(self.destination.read_bytes()).hexdigest())
        from docx import Document
        paragraphs = '\n'.join(p.text for p in Document(self.destination).paragraphs)
        self.assertIn('INVOICE', paragraphs.upper())
        self.assertEqual(before, hashlib.sha256(self.source.read_bytes()).hexdigest())

    @unittest.skipUnless(ocr_ready(), "local eng Tesseract model unavailable")
    def test_mixed_pdf_real_ocr(self):
        self.make_pdf(mixed=True)
        receipt = pipeline.convert_scanned_office(self.source, self.destination, language='eng')
        self.assertEqual(receipt.scanned_pages, (2,))
        self.assertEqual(receipt.ocr_pages, 1)
        from docx import Document
        text = '\n'.join(p.text for p in Document(self.destination).paragraphs)
        self.assertIn('Native editable source text', text)
        self.assertIn('INVOICE', text.upper())

    def test_existing_destination_preserved(self):
        self.make_pdf()
        self.destination.write_bytes(b'KEEP')
        with self.assertRaises(FileExistsError):
            pipeline.convert_scanned_office(self.source, self.destination)
        self.assertEqual(self.destination.read_bytes(), b'KEEP')

    def test_missing_ocr_engine_fails_closed(self):
        self.make_pdf()
        with patch.object(pipeline, 'searchable_pdf', side_effect=RuntimeError('missing tessdata')):
            with self.assertRaisesRegex(RuntimeError, 'missing tessdata'):
                pipeline.convert_scanned_office(self.source, self.destination)
        self.assertFalse(self.destination.exists())
        self.assertEqual(list(self.root.glob('.openpdf-*')), [])

    def test_empty_ocr_is_rejected_without_output(self):
        self.make_pdf()
        def fake_ocr(source, destination, **kwargs):
            with fitz.open() as doc:
                doc.new_page(width=612, height=792)
                doc.save(destination)
            from types import SimpleNamespace
            return SimpleNamespace(sha256=hashlib.sha256(Path(destination).read_bytes()).hexdigest(), ocr_pages=1)
        with patch.object(pipeline, 'searchable_pdf', side_effect=fake_ocr):
            with self.assertRaisesRegex(ValueError, 'OCR insuffisant'):
                pipeline.convert_scanned_office(self.source, self.destination)
        self.assertFalse(self.destination.exists())
        self.assertEqual(list(self.root.glob('.openpdf-*')), [])

    def test_office_conversion_failure_fails_closed(self):
        self.make_pdf(scanned=False)
        with patch.object(pipeline, 'convert_office', side_effect=RuntimeError('converter failed')):
            with self.assertRaisesRegex(RuntimeError, 'converter failed'):
                pipeline.convert_scanned_office(self.source, self.destination)
        self.assertFalse(self.destination.exists())
        self.assertEqual(list(self.root.glob('.openpdf-*')), [])

    def test_invalid_parameters(self):
        self.make_pdf(scanned=False)
        for kwargs in ({'dpi': 999}, {'dpi': True}, {'min_characters_per_scanned_page': 0},
                       {'min_characters_per_scanned_page': True}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                pipeline.convert_scanned_office(self.source, self.destination, **kwargs)
        self.assertFalse(self.destination.exists())

    def test_symlink_source_rejected(self):
        self.make_pdf(scanned=False)
        link = self.root / 'alias.pdf'
        try:
            link.symlink_to(self.source)
        except OSError:
            self.skipTest('Symlinks unavailable')
        with self.assertRaises(ValueError):
            pipeline.convert_scanned_office(link, self.destination)

    def test_reject_tampered_ocr_receipt(self):
        self.make_pdf()
        from types import SimpleNamespace
        def tampered(source, destination, **kwargs):
            Path(destination).write_bytes(b'not a valid OCR PDF')
            return SimpleNamespace(sha256='0'*64, ocr_pages=1)
        with patch.object(pipeline, 'searchable_pdf', side_effect=tampered):
            with self.assertRaisesRegex(RuntimeError, 'Empreinte'):
                pipeline.convert_scanned_office(self.source, self.destination)
        self.assertFalse(self.destination.exists())

    def test_supported_formats_native(self):
        self.make_pdf(scanned=False)
        for ext in ('docx', 'xlsx', 'pptx'):
            dest = self.root / ('native.' + ext)
            receipt = pipeline.convert_scanned_office(self.source, dest)
            self.assertEqual(receipt.output_format, ext)
            self.assertTrue(dest.stat().st_size > 0)


if __name__ == '__main__':
    unittest.main()
