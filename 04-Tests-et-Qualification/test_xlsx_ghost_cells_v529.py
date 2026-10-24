"""V529: real PDF -> XLSX ghost-cell and hidden-table fail-closed regression.

All PDF fixtures are synthetic, not independently annotated business documents.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations

from contextlib import closing
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
from openpyxl import load_workbook
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_tables_v514 import verify_xlsx_tables
from openpdf_xlsx_fidelity_v515 import audit_xlsx_fidelity
from test_xlsx_tables_v514 import fixture


class GhostCellsV529(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.pdf = Path(temp.name) / 'table.pdf'
        self.xlsx = Path(temp.name) / 'table.xlsx'
        fixture(self.pdf)
        convert_office(self.pdf, self.xlsx, table_mode=True)
        with closing(load_workbook(self.xlsx)) as wb:
            self.assertEqual(wb.worksheets[1].max_column, 4)
            self.assertEqual(wb.worksheets[1]['B5'].value, '50 To')

    def mutate(self, cell=None, value='PHANTOM', hidden=False):
        with closing(load_workbook(self.xlsx)) as wb:
            sheet = wb.worksheets[1]
            if cell:
                sheet[cell] = value
                if isinstance(value, str) and value.startswith('='):
                    sheet[cell].data_type = 's'
            if hidden:
                sheet.sheet_state = 'hidden'
            wb.save(self.xlsx)

    def test_clean_conversion_passes_both_audits(self):
        self.assertEqual(verify_xlsx_tables(self.pdf, self.xlsx).cells, 6)
        self.assertEqual(audit_xlsx_fidelity(self.pdf, self.xlsx)['status'], 'PASS')

    def test_phantom_metadata_spacer_row_is_rejected(self):
        self.mutate('A2')
        with self.assertRaisesRegex(ValueError, 'fantome|inattendue'):
            verify_xlsx_tables(self.pdf, self.xlsx)

    def test_phantom_cell_beyond_two_column_table_is_rejected(self):
        self.mutate('C4')
        with self.assertRaisesRegex(ValueError, 'fantome|inattendue'):
            verify_xlsx_tables(self.pdf, self.xlsx)

    def test_formula_like_text_outside_grid_is_rejected(self):
        self.mutate('D5', '=HYPERLINK("http://example.invalid", "click")')
        with self.assertRaisesRegex(ValueError, 'fantome|inattendue'):
            verify_xlsx_tables(self.pdf, self.xlsx)

    def test_hidden_table_is_rejected(self):
        self.mutate(hidden=True)
        with self.assertRaisesRegex(ValueError, 'masquee|visible'):
            verify_xlsx_tables(self.pdf, self.xlsx)

    def test_independent_audit_rejects_phantom_cell(self):
        self.mutate('D3')
        with self.assertRaises(ValueError):
            audit_xlsx_fidelity(self.pdf, self.xlsx)


if __name__ == '__main__':
    unittest.main()
