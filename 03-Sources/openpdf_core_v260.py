"""OpenPDF Studio V26.0: offline PDF operations. Author: Abdelhalim Talbi."""
from pathlib import Path
import fitz

class PDFOperationError(ValueError): pass

def _inputs(paths):
    paths=[Path(p) for p in paths]
    if not paths or any(not p.is_file() or p.suffix.lower()!='.pdf' for p in paths):
        raise PDFOperationError('Sélectionnez des fichiers PDF existants.')
    return paths

def _output(output, sources):
    output=Path(output)
    if output.suffix.lower()!='.pdf': raise PDFOperationError('Le fichier de sortie doit être un PDF.')
    if output.resolve() in {p.resolve() for p in sources}: raise PDFOperationError('La sortie ne peut pas écraser un document source.')
    output.parent.mkdir(parents=True,exist_ok=True)
    return output

def merge_pdfs(paths, output):
    paths=_inputs(paths); output=_output(output,paths)
    with fitz.open() as result:
        for path in paths:
            with fitz.open(path) as doc:
                if doc.needs_pass: raise PDFOperationError(f'PDF protégé : {path.name}')
                result.insert_pdf(doc)
        if not result.page_count: raise PDFOperationError('Aucune page à fusionner.')
        result.save(output,garbage=4,deflate=True)
    return [output]

def split_pdf(source, folder, mode='pages', chunk_size=1):
    source=_inputs([source])[0]; folder=Path(folder)
    if mode not in ('pages','chunks'): raise PDFOperationError('Mode de découpage invalide.')
    if not isinstance(chunk_size,int) or chunk_size<1: raise PDFOperationError('Taille de lot invalide.')
    folder.mkdir(parents=True,exist_ok=True)
    outputs=[]
    with fitz.open(source) as doc:
        if doc.needs_pass: raise PDFOperationError('PDF protégé.')
        step=1 if mode=='pages' else chunk_size
        for start in range(0,len(doc),step):
            target=folder/f'{source.stem}_pages_{start+1:04d}-{min(start+step,len(doc)):04d}.pdf'
            if target.resolve()==source.resolve(): raise PDFOperationError('Écrasement source interdit.')
            with fitz.open() as part:
                part.insert_pdf(doc,from_page=start,to_page=min(start+step,len(doc))-1)
                part.save(target,garbage=4,deflate=True)
            outputs.append(target)
    return outputs

def compress_pdf(source,output):
    source=_inputs([source])[0]; output=_output(output,[source])
    with fitz.open(source) as doc:
        if doc.needs_pass: raise PDFOperationError('PDF protégé.')
        doc.save(output,garbage=4,deflate=True,deflate_images=True,deflate_fonts=True,clean=True)
    return [output]

def extract_text(source,output):
    source=_inputs([source])[0]; output=Path(output)
    if output.suffix.lower()!='.txt': raise PDFOperationError('La sortie doit être un TXT.')
    if output.resolve()==source.resolve(): raise PDFOperationError('Écrasement source interdit.')
    with fitz.open(source) as doc:
        if doc.needs_pass: raise PDFOperationError('PDF protégé.')
        content='\n\n'.join(page.get_text(sort=True) for page in doc)
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(content,encoding='utf-8')
    return [output]

def pdf_to_docx(source,output):
    from docx import Document
    source=_inputs([source])[0]; output=Path(output)
    if output.suffix.lower()!='.docx': raise PDFOperationError('La sortie doit être un DOCX.')
    document=Document()
    with fitz.open(source) as doc:
        if doc.needs_pass: raise PDFOperationError('PDF protégé.')
        for index,page in enumerate(doc):
            if index: document.add_page_break()
            # Preserve images as editable Word drawing objects, not only text.
            # PDF positioning and complex tables still require dedicated layout reconstruction.
            from io import BytesIO
            from docx.shared import Inches
            blocks=page.get_text('dict',sort=True).get('blocks',[])
            for block in blocks:
                if block.get('type')==0:
                    for line in block.get('lines',[]):
                        text=''.join(span.get('text','') for span in line.get('spans',[]))
                        if text.strip():
                            document.add_paragraph(text)
                elif block.get('type')==1 and block.get('image'):
                    width=max(0.5,min(6.0,block.get('width',400)*72/max(block.get('xres',72),1)/72))
                    try:
                        document.add_picture(BytesIO(block['image']),width=Inches(width))
                    except (ValueError, OSError):
                        # Unsupported embedded image format: conversion continues,
                        # but fidelity checks must report missing media.
                        document.add_paragraph('[Image PDF non convertible]')
    output.parent.mkdir(parents=True,exist_ok=True)
    document.save(output)
    return [output]

def pdf_to_xlsx(source,output):
    from openpyxl import Workbook
    source=_inputs([source])[0]; output=Path(output)
    if output.suffix.lower()!='.xlsx': raise PDFOperationError('La sortie doit être un XLSX.')
    wb=Workbook(); ws=wb.active; ws.title='Texte PDF'; ws.append(['Page','Ligne','Texte'])
    with fitz.open(source) as doc:
        if doc.needs_pass: raise PDFOperationError('PDF protégé.')
        for page_num,page in enumerate(doc,1):
            for line_num,line in enumerate(page.get_text(sort=True).splitlines(),1):
                # Prevent spreadsheet formula execution when opening untrusted PDF text.
                if line.startswith(('=','+','-','@')): line="'"+line
                ws.append([page_num,line_num,line[:32767]])
    output.parent.mkdir(parents=True,exist_ok=True); wb.save(output)
    return [output]

def pdf_to_pptx(source,output):
    from pptx import Presentation
    from pptx.util import Inches
    source=_inputs([source])[0]; output=Path(output)
    if output.suffix.lower()!='.pptx': raise PDFOperationError('La sortie doit être un PPTX.')
    prs=Presentation()
    with fitz.open(source) as doc:
        if doc.needs_pass: raise PDFOperationError('PDF protégé.')
        for index,page in enumerate(doc,1):
            slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(9), Inches(.5)).text=f'Page {index}'
            slide.shapes.add_textbox(Inches(.5), Inches(1), Inches(9), Inches(6)).text=page.get_text(sort=True)[:15000]
    output.parent.mkdir(parents=True,exist_ok=True); prs.save(output)
    return [output]
