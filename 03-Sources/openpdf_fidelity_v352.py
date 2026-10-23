"""Indicative text preservation audit, never a visual fidelity score. Abdelhalim Talbi."""
from dataclasses import dataclass
import re
from pathlib import Path
import unicodedata
import fitz

@dataclass(frozen=True)
class TextCoverage:
    source_words: int
    matched_words: int
    coverage: float | None
    note: str

def _words(text):
    normalized = unicodedata.normalize('NFKC', text).casefold()
    return re.findall(r"[^\W_]+(?:['’-][^\W_]+)*", normalized, flags=re.UNICODE)

def office_text(path):
    ext=Path(path).suffix.lower()
    if ext=='.docx':
        from docx import Document
        d=Document(path)
        chunks=[p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                chunks.extend(c.text for c in row.cells)
        return '\n'.join(chunks)
    if ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(path,read_only=True,data_only=False)
        try:
            return '\n'.join(str(c) for sheet in wb for row in sheet.iter_rows(values_only=True) for c in row if c is not None)
        finally: wb.close()
    if ext=='.pptx':
        from pptx import Presentation
        prs=Presentation(path)
        return '\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    raise ValueError('Format Office non pris en charge')

def measure_text_coverage(source, destination):
    from collections import Counter
    with fitz.open(source) as pdf:
        if pdf.needs_pass: raise ValueError('PDF protégé')
        src=Counter(_words('\n'.join(p.get_text() for p in pdf)))
    if not src:
        return TextCoverage(0,0,None,'PDF sans texte extractible : OCR ou inspection visuelle nécessaire')
    dst=Counter(_words(office_text(destination)))
    matched=sum(min(n,dst[word]) for word,n in src.items())
    total=sum(src.values())
    return TextCoverage(total,matched,round(matched/total,4),'Couverture lexicale indicative ; ni géométrie ni mise en page mesurées')
