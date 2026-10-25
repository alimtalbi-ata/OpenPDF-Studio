"""V542 functional qualification: actual editable OOXML, real converter and security negatives."""
import contextlib
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
from openpdf_office_reference_v542 import (evaluate_office_reference_corpus, main,
    SCHEMA, ANNOTATION_SCHEMA, REPORT_SCHEMA)

class OfficeReferenceV542Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.cases = []
        self.thresholds = {'min_token_recall': .95, 'min_case_token_recall': .9,
                           'min_human_reviewed': 0, 'min_editable_elements_per_case': 1}

    def case(self, fmt, phrase, *, actual=None, origin='synthetic_known_at_generation'):
        import fitz
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        idx = len(self.cases) + 1
        pdf = self.root / f'source-{idx}.pdf'
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_text((40, 80), phrase, fontsize=11)
            doc.save(pdf)
        office = self.root / f'output-{idx}.{fmt}'
        actual = phrase if actual is None else actual
        if fmt == 'docx':
            doc = Document()
            if actual:
                doc.add_paragraph(actual)
            else:
                doc.add_picture(self._png(), width=Inches(1))
            doc.save(office)
        elif fmt == 'pptx':
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            if actual:
                textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
                textbox.text = actual
            else:
                slide.shapes.add_picture(str(self._png()), Inches(1), Inches(1))
            presentation.save(office)
        elif fmt == 'xlsx':
            workbook = Workbook()
            if actual:
                workbook.active['A1'] = actual
            else:
                workbook.active.add_image(self._xl_image())
            workbook.save(office)
        ann = self.root / f'ref-{idx}.json'
        ann.write_text(json.dumps({'schema': ANNOTATION_SCHEMA,
            'source_sha256': hashlib.sha256(pdf.read_bytes()).hexdigest(),
            'origin': origin, 'expected_phrases': [phrase]}), encoding='utf-8')
        self.cases.append({'id': f'case-{idx}', 'source_pdf': pdf.name,
                           'office_file': office.name, 'annotations': ann.name})
        return office, ann

    def _png(self):
        from PIL import Image
        file = self.root / 'test-image.png'
        Image.new('RGB', (8, 8), 'white').save(file)
        return file

    def _xl_image(self):
        from openpyxl.drawing.image import Image
        return Image(str(self._png()))

    def manifest(self):
        path = self.root / 'manifest.json'
        path.write_text(json.dumps({'schema': SCHEMA, 'cases': self.cases,
                                    'thresholds': self.thresholds}), encoding='utf-8')
        return path

    def test_real_editable_docx_xlsx_pptx_all_pass(self):
        for fmt, text in [('docx', 'Contrat client Alpha 123'),
                          ('xlsx', 'Budget annuel Beta 456'),
                          ('pptx', 'Plan projet Gamma 789')]:
            self.case(fmt, text)
        report = evaluate_office_reference_corpus(self.manifest())
        self.assertEqual(report['status'], 'PASS')
        self.assertEqual(report['token_recall'], 1)
        self.assertEqual(report['cases'], 3)
        self.assertTrue(all(report['by_format'][x]['cases'] == 1 for x in ('docx','xlsx','pptx')))
        self.assertNotIn('Contrat', json.dumps(report))
        self.assertNotIn(str(self.root), json.dumps(report))

    def test_missing_text_fails_without_raster_false_positive(self):
        self.case('docx', 'texte reference complet')
        self.case('pptx', 'informations indispensables', actual='')
        report = evaluate_office_reference_corpus(self.manifest())
        self.assertEqual(report['status'], 'FAIL')
        self.assertEqual(report['cases_below_recall_floor'], 1)
        self.assertEqual(report['cases_below_editable_floor'], 1)
        self.assertIn('min_case_token_recall', report['failed_thresholds'])

    def test_duplicate_tokens_not_counted_twice(self):
        self.case('docx', 'Alpha Alpha Beta', actual='Alpha Beta')
        self.case('xlsx', 'autre exemple')
        report = evaluate_office_reference_corpus(self.manifest())
        self.assertEqual(report['status'], 'FAIL')
        self.assertEqual(report['matched_tokens'], 4)
        self.assertEqual(report['reference_tokens'], 5)

    def test_annotation_is_bound_to_source_sha(self):
        self.case('docx', 'test unique')
        _, annotation = self.case('xlsx', 'autre test')
        data = json.loads(annotation.read_text())
        data['source_sha256'] = '0' * 64
        annotation.write_text(json.dumps(data))
        with self.assertRaisesRegex(ValueError, 'unbound'):
            evaluate_office_reference_corpus(self.manifest())

    def test_strict_human_provenance_gate(self):
        self.case('docx', 'premier texte')
        self.case('xlsx', 'second texte')
        self.thresholds['min_human_reviewed'] = 1
        report = evaluate_office_reference_corpus(self.manifest())
        self.assertEqual(report['status'], 'FAIL')
        self.assertIn('min_human_reviewed', report['failed_thresholds'])

    def test_macro_member_refused(self):
        office, _ = self.case('docx', 'test un')
        self.case('xlsx', 'test deux')
        with zipfile.ZipFile(office, 'a') as package:
            package.writestr('word/vbaProject.bin', b'evil')
        with self.assertRaisesRegex(ValueError, 'Active'):
            evaluate_office_reference_corpus(self.manifest())

    def test_zip_traversal_member_refused(self):
        office, _ = self.case('docx', 'test un')
        self.case('xlsx', 'test deux')
        with zipfile.ZipFile(office, 'a') as package:
            package.writestr('../outside.xml', b'unsafe')
        with self.assertRaisesRegex(ValueError, 'Unsafe'):
            evaluate_office_reference_corpus(self.manifest())

    def test_external_relationship_refused(self):
        office, _ = self.case('docx', 'test un')
        self.case('xlsx', 'test deux')
        with zipfile.ZipFile(office, 'a') as package:
            package.writestr('word/_rels/custom.rels',
                b'<Relationships><Relationship TargetMode="External" Target="https://invalid.test"/></Relationships>')
        with self.assertRaisesRegex(ValueError, 'External'):
            evaluate_office_reference_corpus(self.manifest())

    def test_cli_refuses_overwrite_and_does_not_leak(self):
        self.case('docx', 'secretmetier tresconfidentiel')
        self.case('xlsx', 'donnee autre')
        receipt = self.root / 'receipt.json'
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            self.assertEqual(main([str(self.manifest()), '--receipt', str(receipt)]), 0)
        self.assertEqual(json.loads(receipt.read_text())['schema'], REPORT_SCHEMA)
        self.assertNotIn('secretmetier', out.getvalue())
        with contextlib.redirect_stdout(out):
            self.assertEqual(main([str(self.manifest()), '--receipt', str(receipt)]), 3)
        self.assertNotIn(str(self.root), out.getvalue())

    def test_actual_openpdf_converter_three_formats(self):
        from openpdf_conversion_service_v351 import convert_office
        for fmt, phrase in [('docx', 'Contrat client Alpha 123'),
                            ('xlsx', 'Budget annuel Beta 456'),
                            ('pptx', 'Plan projet Gamma 789')]:
            office, _ = self.case(fmt, phrase)
            office.unlink()
            source = self.root / f'source-{len(self.cases)}.pdf'
            result = convert_office(source, office)
            self.assertEqual(hashlib.sha256(office.read_bytes()).hexdigest(), result.sha256)
        report = evaluate_office_reference_corpus(self.manifest())
        self.assertEqual(report['status'], 'PASS')
        self.assertEqual(report['matched_tokens'], 12)
        self.assertEqual(report['reference_tokens'], 12)

    def test_invalid_source_pdf_refused(self):
        self.case('docx', 'premier test')
        self.case('xlsx', 'second test')
        pdf = self.root / 'source-1.pdf'
        pdf.write_bytes(b'%PDF-1.7\nnot a real document')
        with self.assertRaises(ValueError):
            evaluate_office_reference_corpus(self.manifest())

    def test_thresholds_reject_nan_and_bool(self):
        self.case('docx', 'test un')
        self.case('xlsx', 'test deux')
        for value in [float('nan'), True, -0.1, 1.1]:
            self.thresholds['min_token_recall'] = value
            with self.subTest(value=value), self.assertRaises(ValueError):
                evaluate_office_reference_corpus(self.manifest())

if __name__ == '__main__':
    unittest.main()
