"""V51.4 functional tests: real PDF tables, Excel cells, formula safety and rollback."""
import tempfile
import unittest
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))

import fitz
from openpyxl import load_workbook
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_tables_v514 import verify_xlsx_tables


def fixture(path, *, table=True, injection=False):
    pdf = fitz.open()
    page = pdf.new_page(width=595, height=842)
    page.insert_text((50, 40), 'Rapport inventaire stockage')
    if table:
        x, y, w, h = 50, 80, 400, 120
        shape = page.new_shape()
        for j in range(4):
            shape.draw_line((x, y + j * h / 3), (x + w, y + j * h / 3))
        for j in range(3):
            shape.draw_line((x + j * w / 2, y), (x + j * w / 2, y + h))
        shape.finish(color=(0, 0, 0), width=1)
        shape.commit()
        values = [['Volume', 'Capacite'], ['SAN-A', '20 To'], ['SAN-B', '=1+1' if injection else '50 To']]
        for i, row in enumerate(values):
            for j, value in enumerate(row):
                page.insert_text((x + 10 + j * w / 2, y + 25 + i * h / 3), value, fontsize=10)
    pdf.save(path)
    pdf.close()


class TestTablesV514(unittest.TestCase):
    def test_table_export_and_exact_cells(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d)/'a.pdf', Path(d)/'a.xlsx'
            fixture(src)
            result = convert_office(src, dst, table_mode=True)
            self.assertEqual(result.output_format, 'xlsx')
            self.assertEqual(verify_xlsx_tables(src, dst).tables, 1)
            wb = load_workbook(dst)
            try:
                self.assertEqual(len(wb.worksheets), 2)
                self.assertEqual(wb.worksheets[1]['A4'].value, 'SAN-A')
                self.assertEqual(wb.worksheets[1]['B5'].value, '50 To')
            finally:
                wb.close()

    def test_legacy_export_unchanged_by_default(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d)/'a.pdf', Path(d)/'a.xlsx'
            fixture(src)
            convert_office(src, dst)
            wb = load_workbook(dst)
            try:
                self.assertEqual(wb.sheetnames, ['Texte PDF'])
            finally:
                wb.close()

    def test_missing_sheet_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d)/'a.pdf', Path(d)/'a.xlsx'
            fixture(src)
            convert_office(src, dst, table_mode=True)
            wb = load_workbook(dst)
            wb.remove(wb.worksheets[-1])
            wb.save(dst)
            wb.close()
            with self.assertRaises(ValueError):
                verify_xlsx_tables(src, dst)

    def test_formula_is_literal(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d)/'a.pdf', Path(d)/'a.xlsx'
            fixture(src, injection=True)
            convert_office(src, dst, table_mode=True)
            wb = load_workbook(dst)
            try:
                self.assertEqual(wb.worksheets[1]['B5'].value, '=1+1')
                self.assertEqual(wb.worksheets[1]['B5'].data_type, 's')
            finally:
                wb.close()

    def test_no_table_fails_closed(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d)/'a.pdf', Path(d)/'a.xlsx'
            fixture(src, table=False)
            with self.assertRaises(ValueError):
                convert_office(src, dst, table_mode=True)
            self.assertFalse(dst.exists())

    def test_existing_destination_untouched(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d)/'a.pdf', Path(d)/'a.xlsx'
            fixture(src)
            dst.write_bytes(b'important')
            with self.assertRaises(FileExistsError):
                convert_office(src, dst, table_mode=True)
            self.assertEqual(dst.read_bytes(), b'important')

    def test_opt_in_rejects_other_formats(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d)/'a.pdf', Path(d)/'a.docx'
            fixture(src)
            with self.assertRaises(ValueError):
                convert_office(src, dst, table_mode=True)
            self.assertFalse(dst.exists())

    def test_mutation_detected(self):
        with tempfile.TemporaryDirectory() as d:
            src, dst = Path(d)/'a.pdf', Path(d)/'a.xlsx'
            fixture(src)
            convert_office(src, dst, table_mode=True)
            wb = load_workbook(dst)
            wb.worksheets[1]['B4'] = 'TAMPERED'
            wb.save(dst)
            wb.close()
            with self.assertRaises(ValueError):
                verify_xlsx_tables(src, dst)


if __name__ == '__main__':
    unittest.main()
