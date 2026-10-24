"""V523 mixed-page PDF table extraction, real synthetic PDFs; not human proof.
Author: Abdelhalim Talbi.
"""
from contextlib import closing
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
import fitz
from openpyxl import load_workbook
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_tables_v514 import (
    _rect_overlap, _unruled_regions, verify_xlsx_tables,
)


def make_mixed_pdf(path, *, side_by_side=False, prose=False, formula=False):
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    # Continuous grid, unlike disjoint cell rectangles that PyMuPDF may not
    # recognize as a geometrically ruled table.
    top, left, col_width, row_height = 90, 40, 95, 27
    for r in range(5):
        page.draw_line((left, top + r * row_height),
                       (left + 3 * col_width, top + r * row_height))
    for c in range(4):
        page.draw_line((left + c * col_width, top),
                       (left + c * col_width, top + 4 * row_height))
    for r in range(4):
        for c in range(3):
            page.insert_text((left + c * col_width + 5, top + r * row_height + 17),
                             f'RULE-{r}-{c}', fontsize=9)
    bx, by = (365, 100) if side_by_side else (40, 360)
    for r in range(4):
        for c in range(3):
            value = '=1+1' if formula and r == 3 and c == 2 else f'TEXT-{r}-{c}'
            page.insert_text((bx + c * (67 if side_by_side else 95), by + r * 27),
                             value, fontsize=8)
    if prose:
        for r in range(4):
            page.insert_text((40, 620 + r * 25),
                             f'Paragraphe {r}: texte narratif sans colonnes.', fontsize=9)
    doc.save(path)
    doc.close()


class MixedTablesV523(unittest.TestCase):
    def test_overlap_geometry(self):
        self.assertTrue(_rect_overlap((0, 0, 10, 10), (5, 5, 15, 15)))
        self.assertFalse(_rect_overlap((0, 0, 10, 10), (10, 0, 20, 10)))

    def test_regions_disjoint_and_exclude_ruled_area(self):
        page = fitz.Rect(0, 0, 500, 700)
        boxes = [(20, 20, 200, 200), (300, 300, 400, 400)]
        regions = _unruled_regions(page, boxes)
        self.assertGreater(len(regions), 2)
        for i, region in enumerate(regions):
            for box in boxes:
                self.assertFalse(_rect_overlap(region, box))
            for other in regions[i+1:]:
                self.assertFalse(_rect_overlap(region, other))

    def test_fragmentation_fails_closed(self):
        boxes = [(x, y, x + 15, y + 15) for x in range(20, 480, 50)
                 for y in range(20, 480, 50)]
        with self.assertRaisesRegex(ValueError, 'fragmentee'):
            _unruled_regions(fitz.Rect(0, 0, 500, 500), boxes, limit=10)

    def _verify(self, side_by_side=False, prose=False, formula=False):
        with tempfile.TemporaryDirectory() as tmp:
            pdfpath, xlspath = Path(tmp)/'mixed.pdf', Path(tmp)/'mixed.xlsx'
            make_mixed_pdf(pdfpath, side_by_side=side_by_side, prose=prose, formula=formula)
            with fitz.open(pdfpath) as doc:
                self.assertEqual(len(doc[0].find_tables().tables), 1,
                                 'Test fixture must contain a detected ruled table')
            convert_office(pdfpath, xlspath, table_mode=True)
            ev = verify_xlsx_tables(pdfpath, xlspath)
            self.assertEqual((ev.tables, ev.cells), (2, 24))
            with closing(load_workbook(xlspath)) as wb:
                self.assertEqual(wb.sheetnames,
                                 ['Texte PDF', 'P0001-T001', 'P0001-T002'])
                self.assertEqual(wb.worksheets[1]['A3'].value, 'RULE-0-0')
                self.assertEqual(wb.worksheets[2]['A3'].value, 'TEXT-0-0')
                self.assertEqual(wb.worksheets[2]['C6'].value,
                                 '=1+1' if formula else 'TEXT-3-2')
                if formula:
                    self.assertEqual(wb.worksheets[2]['C6'].data_type, 's')
            wb = load_workbook(xlspath)
            wb.worksheets[2]['A3'] = 'CORRUPTED'
            wb.save(xlspath)
            wb.close()
            with self.assertRaises(ValueError):
                verify_xlsx_tables(pdfpath, xlspath)

    def test_stacked_mixed_tables(self):
        self._verify()

    def test_side_by_side_mixed_tables(self):
        self._verify(side_by_side=True)

    def test_mixed_page_prose_not_table(self):
        self._verify(prose=True)

    def test_mixed_page_formula_is_literal(self):
        self._verify(formula=True)


if __name__ == '__main__':
    unittest.main()
