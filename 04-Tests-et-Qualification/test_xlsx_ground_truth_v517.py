"""Independent ground-truth PDF/XLSX tests; known synthetic truth, not extractor-derived.
Author: Abdelhalim Talbi. No Windows GUI qualification is implied.
"""
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
from openpyxl import load_workbook
from openpdf_conversion_service_v351 import convert_office
from openpdf_xlsx_ground_truth_v517 import SCHEMA, evaluate_ground_truth, main
from test_xlsx_tables_v514 import fixture


class GroundTruthV517Tests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.directory = Path(temp.name)
        self.pdf = self.directory / 'inventory.pdf'
        self.xlsx = self.directory / 'inventory.xlsx'
        self.annotations = self.directory / 'truth.json'
        fixture(self.pdf)
        convert_office(self.pdf, self.xlsx, table_mode=True)
        # Ground truth is explicitly known from fixture() inputs, not find_tables().
        self.truth = {
            'schema': SCHEMA,
            'source_sha256': hashlib.sha256(self.pdf.read_bytes()).hexdigest(),
            'origin': 'synthetic_known_at_generation',
            'tables': [{'page': 1, 'index': 1,
                        'cells': [['Volume', 'Capacite'], ['SAN-A', '20 To'], ['SAN-B', '50 To']],
                        'merges': []}],
        }
        self.save_truth()

    def save_truth(self):
        self.annotations.write_text(json.dumps(self.truth, ensure_ascii=False), encoding='utf-8')

    def edit(self, callback):
        book = load_workbook(self.xlsx)
        callback(book)
        book.save(self.xlsx)
        book.close()

    def test_independent_exact_real_pdf_conversion(self):
        report = evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)
        self.assertEqual(report['status'], 'PASS')
        self.assertEqual(report['annotated_cells'], 6)
        self.assertEqual(report['exact_cells'], 6)
        self.assertEqual(report['table_precision'], 1)
        self.assertEqual(report['table_recall'], 1)
        self.assertEqual(report['semantic_fidelity'], 'annotated_cell_content_only')
        self.assertNotIn('SAN-A', json.dumps(report))
        self.assertEqual(len(report['annotation_sha256']), 64)

    def test_wrong_value_measured_not_hidden_by_extractor(self):
        self.edit(lambda book: setattr(book.worksheets[1]['B4'], 'value', '99 To'))
        report = evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)
        self.assertEqual(report['status'], 'DEVIATION')
        self.assertEqual(report['exact_cells'], 5)
        self.assertEqual(report['cell_exact_accuracy'], 0.833333)

    def test_whitespace_normalized_score_is_separate(self):
        self.edit(lambda book: setattr(book.worksheets[1]['A4'], 'value', 'SAN-A  '))
        report = evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)
        self.assertEqual(report['status'], 'DEVIATION')
        self.assertEqual(report['exact_cells'], 5)
        self.assertEqual(report['cell_normalized_accuracy'], 1)

    def test_missing_table_penalized(self):
        self.edit(lambda book: book.remove(book.worksheets[1]))
        report = evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)
        self.assertEqual(report['matched_tables'], 0)
        self.assertEqual(report['table_recall'], 0)
        self.assertEqual(report['exact_cells'], 0)

    def test_unexpected_table_penalized(self):
        def mutate(book):
            sheet = book.create_sheet('P0001-T002')
            sheet['A1'], sheet['B1'] = 'Page PDF', 1
            sheet['C1'], sheet['D1'] = 'Tableau', 2
            sheet['A3'] = 'EXTRA'
        self.edit(mutate)
        report = evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)
        self.assertEqual(report['observed_tables'], 2)
        self.assertEqual(report['table_precision'], 0.5)
        self.assertEqual(report['extra_nonempty_cells'], 1)

    def test_extra_cell_penalized(self):
        self.edit(lambda book: setattr(book.worksheets[1]['A6'], 'value', 'EXTRA'))
        report = evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)
        self.assertEqual(report['status'], 'DEVIATION')
        self.assertEqual(report['extra_nonempty_cells'], 1)

    def test_merged_cell_ground_truth_match_and_mismatch(self):
        self.truth['tables'][0]['cells'][0] = ['Volume', None]
        self.truth['tables'][0]['merges'] = [[1, 1, 1, 2]]
        self.save_truth()
        self.edit(lambda book: book.worksheets[1].merge_cells('A3:B3'))
        report = evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)
        self.assertEqual(report['status'], 'PASS')
        self.assertEqual(report['matched_merges'], 1)
        self.edit(lambda book: book.worksheets[1].unmerge_cells('A3:B3'))
        report = evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)
        self.assertEqual(report['merge_recall'], 0)
        self.assertEqual(report['status'], 'DEVIATION')

    def test_wrong_pdf_hash_fails_closed(self):
        self.truth['source_sha256'] = '0' * 64
        self.save_truth()
        with self.assertRaisesRegex(ValueError, 'non liee'):
            evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)

    def test_duplicate_json_keys_rejected(self):
        self.annotations.write_text('{"schema": "x", "schema": "y"}', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, 'dupliquee'):
            evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)

    def test_invalid_merge_overlap_rejected(self):
        self.truth['tables'][0]['cells'][0] = ['Volume', None]
        self.truth['tables'][0]['merges'] = [[1, 1, 1, 2], [1, 1, 2, 1]]
        self.save_truth()
        with self.assertRaisesRegex(ValueError, 'chevauchantes'):
            evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)

    def test_nonrectangular_grid_rejected(self):
        self.truth['tables'][0]['cells'][1].pop()
        self.save_truth()
        with self.assertRaisesRegex(ValueError, 'rectangulaire'):
            evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)

    def test_active_formula_in_legacy_sheet_rejected(self):
        self.edit(lambda book: setattr(book.worksheets[0]['A30'], 'value', '=1+1'))
        with self.assertRaisesRegex(ValueError, 'Formule'):
            evaluate_ground_truth(self.pdf, self.xlsx, self.annotations)

    def test_symlink_annotation_rejected(self):
        link = self.directory / 'link.json'
        try:
            link.symlink_to(self.annotations)
        except OSError:
            self.skipTest('Symlinks unavailable on this Windows runner')
        with self.assertRaisesRegex(ValueError, 'symbolique'):
            evaluate_ground_truth(self.pdf, self.xlsx, link)

    def test_cli_deviation_exit_three_and_no_text_leak(self):
        self.edit(lambda book: setattr(book.worksheets[1]['B4'], 'value', 'bad'))
        self.assertEqual(main(['--pdf', str(self.pdf), '--xlsx', str(self.xlsx),
                               '--annotations', str(self.annotations)]), 3)

    def test_cli_success_report_no_overwrite(self):
        report = self.directory / 'metrics.json'
        args = ['--pdf', str(self.pdf), '--xlsx', str(self.xlsx),
                '--annotations', str(self.annotations), '--report', str(report)]
        self.assertEqual(main(args), 0)
        self.assertNotIn('SAN-A', report.read_text(encoding='utf-8'))
        self.assertEqual(main(args), 2)


if __name__ == '__main__':
    unittest.main()
