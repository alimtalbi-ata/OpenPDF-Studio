"""V51.4 opt-in editable PDF table worksheets with fail-closed source audit.

PyMuPDF detects geometric tables heuristically; exact extraction consistency
is verified, not semantic correctness or original Excel formatting.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import fitz
from openpyxl import load_workbook

MAX_TABLES = 256
MAX_CELLS = 200_000
MAX_ROWS_PER_TABLE = 4096
MAX_COLUMNS = 128
MAX_CELL_UTF16 = 32767


@dataclass(frozen=True)
class TableEvidence:
    tables: int
    cells: int
    pages_with_tables: tuple[int, ...]


def _tables(source):
    count = 0
    cells = 0
    pages = []
    with fitz.open(source) as pdf:
        if pdf.needs_pass:
            raise ValueError('PDF chiffre : extraction des tableaux refusee')
        for page_number, page in enumerate(pdf, 1):
            if page_number > 9999:
                raise ValueError('Nombre de pages hors limite du mode tableaux')
            found = page.find_tables()
            for table_number, table in enumerate(found.tables, 1):
                count += 1
                if count > MAX_TABLES or table_number > 999:
                    raise ValueError('Trop de tableaux detectes')
                matrix = table.extract()
                if not matrix or len(matrix) > MAX_ROWS_PER_TABLE:
                    raise ValueError('Dimensions de tableau invalides')
                columns = max((len(row) for row in matrix), default=0)
                if not 1 <= columns <= MAX_COLUMNS:
                    raise ValueError('Nombre de colonnes hors limite')
                for row in matrix:
                    if len(row) > columns:
                        raise ValueError('Tableau irregulier')
                    for value in row:
                        if value is not None and not isinstance(value, str):
                            raise ValueError('Cellule PDF non textuelle')
                        if value is not None and len(value.encode('utf-16-le')) // 2 > MAX_CELL_UTF16:
                            raise ValueError('Cellule Excel trop longue : pas de troncature')
                cells += len(matrix) * columns
                if cells > MAX_CELLS:
                    raise ValueError('Trop de cellules detectees')
                if page_number not in pages:
                    pages.append(page_number)
                yield page_number, table_number, matrix, columns
    if count == 0:
        raise ValueError('Aucun tableau detecte : aucune sortie enrichie')


def _expected(source):
    result = list(_tables(source))
    return result, TableEvidence(len(result), sum(len(rows) * width for _, _, rows, width in result),
                                 tuple(dict.fromkeys(page for page, _, _, _ in result)))


def enrich_xlsx_with_tables(source, destination):
    """Add editable table worksheets to the legacy text-preserving XLSX stage.

    Must be called only on a private temporary workbook, before atomic publish.
    """
    source, destination = Path(source), Path(destination)
    expected, evidence = _expected(source)
    wb = load_workbook(destination)
    try:
        if len(wb.worksheets) != 1 or wb.worksheets[0].title != 'Texte PDF':
            raise ValueError('Classeur Excel source inattendu')
        for page, index, rows, columns in expected:
            sheet = wb.create_sheet(f'P{page:04d}-T{index:03d}')
            sheet['A1'] = 'Page PDF'
            sheet['B1'] = page
            sheet['C1'] = 'Tableau'
            sheet['D1'] = index
            for row_number, row in enumerate(rows, 3):
                for column_number in range(1, columns + 1):
                    value = row[column_number - 1] if column_number <= len(row) else None
                    if value not in (None, ""):
                        cell = sheet.cell(row_number, column_number, value)
                        cell.data_type = 's'  # Never evaluate a PDF-origin Excel formula.
        wb.save(destination)
    finally:
        wb.close()
    verify_xlsx_tables(source, destination)
    return evidence


def verify_xlsx_tables(source, destination):
    """Re-extract source tables and compare every cell to the saved workbook."""
    expected, evidence = _expected(source)
    wb = load_workbook(destination, read_only=True, data_only=False)
    try:
        if len(wb.worksheets) != len(expected) + 1 or wb.worksheets[0].title != 'Texte PDF':
            raise ValueError('Feuilles Excel manquantes ou supplementaires')
        for position, (page, index, rows, columns) in enumerate(expected, 1):
            sheet = wb.worksheets[position]
            if sheet.title != f'P{page:04d}-T{index:03d}':
                raise ValueError('Ordre ou nom des tableaux Excel invalide')
            if (sheet['A1'].value, sheet['B1'].value, sheet['C1'].value, sheet['D1'].value) != ('Page PDF', page, 'Tableau', index):
                raise ValueError('Provenance du tableau Excel invalide')
            if sheet.max_row != len(rows) + 2 or sheet.max_column != max(columns, 4):
                raise ValueError('Dimensions du tableau Excel modifiees')
            for row_number, row in enumerate(rows, 3):
                for column_number in range(1, columns + 1):
                    cell = sheet.cell(row_number, column_number)
                    value = row[column_number - 1] if column_number <= len(row) else None
                    if cell.value != (None if value == '' else value) or cell.data_type == 'f':
                        raise ValueError('Cellule de tableau Excel modifiee ou formule')
        return evidence
    finally:
        wb.close()
