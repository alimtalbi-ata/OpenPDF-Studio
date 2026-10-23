"""Generate V50.7-specific Word/PDF from the actual source contract.

Optional real UI screenshot is used only when a capture exists and is explicitly
identified as Linux Xvfb, never as a Windows capture.
Author: Abdelhalim Talbi.
"""
from pathlib import Path
import re

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / '05-Documentation'
PROOF = ROOT / '04-Tests-et-Qualification' / 'Preuves-Reelles-V50.7'
TITLE = 'OpenPDF Studio V50.7 - Architecture et qualification'
AUTHOR = 'Abdelhalim Talbi'
SECTIONS = [
    ('Objectif et apport produit',
     'V50.7 ferme une lacune de la compression automatique : un recu RGB dont le nombre de pages ou de pixels est seulement plausible, mais incomplet, ne peut plus autoriser la publication. Le controle relit la geometrie reelle de toutes les pages, rotations incluses, sans refaire un rendu RGB.'),
    ('Pipeline et decision',
     'Source PDF -> candidats temporaires par profil -> moteur de compression -> audit RGB a echelle bornee -> validation des SHA-256 et tailles V50.6 -> verification independante de la couverture V50.7 -> selection du plus petit candidat conforme -> publication exclusive. Toute incoherence rejette le candidat avant publication.'),
    ('Controles V50.7',
     'Nombre de pages source/candidat identique ; pages_compared exactement egal au nombre de pages ; geometries, cropbox et rotations identiques ; pixels_compared egal a la somme des dimensions raster prevues ; saved_bytes egal a source_bytes moins candidate_bytes ; refus des PDFs chiffres et des donnees invalides.'),
    ('Tests et reproductibilite',
     'Douze nouveaux tests couvrent un vrai PDF synthetique a deux pages dont une tournee, des recus falsifies (pages, pixels, octets), un PDF chiffre, un PDF invalide, la collision de destination et la source inchangee. Les tests V50.6, V50.5 et V50.4 sont rejoues en CI Ubuntu/Windows. Les resultats exacts de CI sont a consulter dans la PR, pas a deduire de ce document.'),
    ('Qualification Windows et preuves',
     'La compilation, les auto-tests du portable, la construction de l installateur, son installation et sa desinstallation sont des controles distincts. Les preuves reelles de CI doivent etre consultees avant toute diffusion. Une capture Linux Xvfb, si jointe, n est pas une capture Windows.'),
    ('Limites et risques',
     'La seconde verification controle l exhaustivite des mesures et non la valeur de chaque pixel. L audit RGB reste a echelle bornee ; il ne certifie pas PDF/A, OCR, accessibilite, impression ou fidelite des conversions Office. Les PDFs synthetiques ne remplacent pas un corpus metier.'),
    ('Prochaine priorite',
     'Etendre le corpus metier aux scans, tableaux, images, formulaires et liens ; qualifier humainement l interface Windows ; traiter les tests historiques et approfondir la fidelite des conversions Word/Excel/PowerPoint editables.'),
]


def generate():
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    from docx.shared import Inches, Pt
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PDFImage, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors

    DOCS.mkdir(exist_ok=True)
    diagram = DOCS / 'Architecture-Couverture-RGB-V50.7.png'
    im = Image.new('RGB', (1320, 280), 'white')
    d = ImageDraw.Draw(im)
    font = ImageFont.load_default(size=24)
    boxes = [('PDF source', 15), ('Compression\nprivee', 230), ('Audit RGB\nV50.6', 445),
             ('Couverture\nV50.7', 660), ('SHA-256\nfinal', 875), ('Publication\nexclusive', 1090)]
    for label, x in boxes:
        d.rounded_rectangle((x, 65, x+190, 205), radius=16, outline='#1a3850', width=4, fill='#eef5fa')
        d.multiline_text((x+15, 111), label, fill='#102a43', spacing=9, font=font)
        if x < 1090:
            d.line((x+191, 135, x+212, 135), fill='#1a3850', width=4)
    im.save(diagram)

    doc = Document()
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = 'OPENPDF STUDIO  /  V50.7'
    doc.add_heading(TITLE, 0)
    doc.add_paragraph('Auteur : ' + AUTHOR)
    doc.add_paragraph('Documentation specifique V50.7 - 23 octobre 2026')
    doc.add_picture(str(diagram), width=Inches(6.3))
    for heading, body in SECTIONS:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)
    doc.add_heading('Statuts de qualification (ne pas confondre)', level=1)
    rows = [('Controle', 'Statut a documenter'),
            ('Tests fonctionnels Linux', 'Rapports locaux / CI'),
            ('CI Ubuntu/Windows', 'Journaux GitHub Actions'),
            ('Portable Windows', 'Recu de smoke natif'),
            ('Installateur Windows', 'Recu installation/desinstallation'),
            ('Validation humaine GUI', 'Non realisee')]
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Light Shading Accent 1'
    for a, b in rows:
        c = t.add_row().cells
        c[0].text, c[1].text = a, b
    screenshot = PROOF / 'UI-Linux-Xvfb-V50.7.png'
    if screenshot.is_file():
        doc.add_heading('Capture reelle de l interface (Linux Xvfb)', level=1)
        doc.add_paragraph('Capture de l application lancee sous Linux/Xvfb. Ce n est pas une preuve de qualification graphique Windows.')
        doc.add_picture(str(screenshot), width=Inches(6.2))
    doc.core_properties.author = AUTHOR
    doc.save(DOCS / 'OpenPDF-Studio-V50.7-Guide.docx')

    styles = getSampleStyleSheet()
    story = [Paragraph(TITLE, styles['Title']), Spacer(1, 12),
             Paragraph('Auteur : ' + AUTHOR + ' | 23 octobre 2026', styles['Normal']),
             Spacer(1, 16), PDFImage(str(diagram), width=495, height=105), Spacer(1, 16)]
    for heading, body in SECTIONS:
        story += [Paragraph(heading, styles['Heading2']), Paragraph(body, styles['BodyText']), Spacer(1, 10)]
    story.append(Paragraph('Statuts de qualification (preuves distinctes)', styles['Heading2']))
    table = Table(rows, colWidths=[165, 310])
    table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
        ('GRID', (0,0), (-1,-1), 0.4, colors.grey), ('VALIGN',(0,0),(-1,-1),'TOP'),
        ('LEFTPADDING',(0,0),(-1,-1),8), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story += [table, Spacer(1, 15)]
    if screenshot.is_file():
        story += [Paragraph('Capture reelle de l interface - Linux Xvfb (pas Windows)', styles['Heading2']),
                  PDFImage(str(screenshot), width=475, height=309)]
    SimpleDocTemplate(str(DOCS / 'OpenPDF-Studio-V50.7-Guide.pdf'), pagesize=A4,
                      leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48).build(story)
    return diagram


if __name__ == '__main__':
    print(generate())
