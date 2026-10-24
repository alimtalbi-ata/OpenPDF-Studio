"""Independent, bounded PDF-table -> XLSX fidelity and security audit.

Checks exact *extractor parity*, not semantic table recognition or visual fidelity.
Does not copy source text into the report. Author: Abdelhalim Talbi.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import stat
import tempfile
import zipfile

import fitz
from openpyxl import load_workbook

from openpdf_xlsx_tables_v514 import _expected, verify_xlsx_tables

MAX_SOURCE_BYTES = 256 * 1024 * 1024
MAX_XLSX_BYTES = 128 * 1024 * 1024
MAX_ZIP_ENTRIES = 4096
MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_XML_BYTES = 64 * 1024 * 1024
MAX_RATIO = 200
FORBIDDEN_PARTS = ('xl/externallinks/', 'xl/embeddings/', 'xl/activex/',
                   'xl/connections.xml', 'xl/vbaproject.bin')


def _file_hash(path: Path, maximum: int) -> tuple[str, int]:
    if path.is_symlink() or not path.is_file():
        raise ValueError('Fichier absent ou lien symbolique interdit')
    size = path.stat().st_size
    if size > maximum:
        raise ValueError('Fichier depassant la limite de taille')
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        while block := stream.read(1024 * 1024):
            digest.update(block)
    if path.is_symlink() or path.stat().st_size != size:
        raise ValueError('Source modifiee pendant le controle')
    return digest.hexdigest(), size


def _xlsx_preflight(path: Path) -> None:
    """Reject dangerous/oversized ZIP members before openpyxl decompresses XML."""
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if not infos or len(infos) > MAX_ZIP_ENTRIES:
                raise ValueError('Nombre de composants XLSX hors limite')
            names = set()
            total = 0
            for item in infos:
                name = item.filename.replace('\\', '/')
                lower = name.lower()
                if (lower in names or name.startswith('/') or '..' in name.split('/')
                        or name.startswith('\\') or ':' in name or '\x00' in name):
                    raise ValueError('Chemin XLSX dangereux ou duplique')
                names.add(lower)
                if any(lower.startswith(p) for p in FORBIDDEN_PARTS):
                    raise ValueError('Contenu XLSX actif ou externe interdit')
                if stat.S_IFMT(item.external_attr >> 16) == stat.S_IFLNK:
                    raise ValueError('Lien symbolique dans XLSX interdit')
                total += item.file_size
                if (total > MAX_UNCOMPRESSED_BYTES
                        or (lower.endswith(('.xml', '.rels')) and item.file_size > MAX_XML_BYTES)
                        or (item.file_size > 1024 * 1024
                            and item.file_size > MAX_RATIO * max(item.compress_size, 1))):
                    raise ValueError('Decompression XLSX excessive')
            if '[content_types].xml' not in names or 'xl/workbook.xml' not in names:
                raise ValueError('Structure XLSX incomplete')
            if archive.testzip() is not None:
                raise ValueError('CRC XLSX invalide')
    except (zipfile.BadZipFile, OSError, RuntimeError) as exc:
        raise ValueError('Archive XLSX invalide') from exc


def audit_xlsx_fidelity(source, workbook) -> dict:
    """Audit a generated workbook, raising on any integrity/security mismatch."""
    source, workbook = Path(source), Path(workbook)
    if source.resolve() == workbook.resolve():
        raise ValueError('Source et resultat identiques')
    pdf_hash, pdf_bytes = _file_hash(source, MAX_SOURCE_BYTES)
    xlsx_hash, xlsx_bytes = _file_hash(workbook, MAX_XLSX_BYTES)
    _xlsx_preflight(workbook)
    expected, _ = _expected(source)
    evidence = verify_xlsx_tables(source, workbook)
    wb = load_workbook(workbook, read_only=False, data_only=False, keep_links=False)
    try:
        for sheet in wb.worksheets:
            # Sparse XLSX rows may claim 1,048,576 rows: inspect only real cells.
            for cell in sheet._cells.values():
                if cell.data_type == 'f' or cell.hyperlink is not None:
                    raise ValueError('Formule ou hyperlien actif dans XLSX')
        merges = sum(len(entry[4]) if len(entry) == 5 else 0 for entry in expected)
        nonempty = sum(sum(value not in (None, '') for value in row)
                       for entry in expected for row in entry[2])
    finally:
        wb.close()
    with fitz.open(source) as pdf:
        page_count = pdf.page_count
    if _file_hash(source, MAX_SOURCE_BYTES) != (pdf_hash, pdf_bytes):
        raise ValueError('PDF modifie pendant le controle')
    if _file_hash(workbook, MAX_XLSX_BYTES) != (xlsx_hash, xlsx_bytes):
        raise ValueError('XLSX modifie pendant le controle')
    return {
        'schema': 'openpdf.xlsx.fidelity.v1',
        'status': 'PASS',
        'scope': 'exact_extractor_parity_only',
        'semantic_fidelity': 'not_assessed',
        'visual_fidelity': 'not_assessed',
        'source_sha256': pdf_hash,
        'source_bytes': pdf_bytes,
        'workbook_sha256': xlsx_hash,
        'workbook_bytes': xlsx_bytes,
        'pages': page_count,
        'pages_with_tables': list(evidence.pages_with_tables),
        'tables': evidence.tables,
        'grid_cells': evidence.cells,
        'nonempty_cells': nonempty,
        'merged_ranges': merges,
        'page_table_coverage': round(len(evidence.pages_with_tables) / page_count, 6) if page_count else 0,
    }


def write_report_new(path, evidence: dict) -> None:
    """Publish JSON atomically without overwriting an existing report."""
    target = Path(path)
    if target.is_symlink() or target.exists():
        raise FileExistsError('Rapport deja present ou lien symbolique')
    if not target.parent.is_dir() or target.parent.is_symlink():
        raise ValueError('Repertoire de rapport non fiable')
    payload = (json.dumps(evidence, ensure_ascii=False, sort_keys=True, indent=2) + '\n').encode()
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode='wb', prefix='.openpdf-audit-',
                                         suffix='.tmp', dir=target.parent, delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
        os.link(temporary, target, follow_symlinks=False)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description='Audit PDF -> Excel, sans texte source dans le rapport')
    parser.add_argument('--pdf', required=True)
    parser.add_argument('--xlsx', required=True)
    parser.add_argument('--report')
    args = parser.parse_args(argv)
    try:
        evidence = audit_xlsx_fidelity(args.pdf, args.xlsx)
        if args.report:
            write_report_new(args.report, evidence)
    except (ValueError, FileExistsError, OSError) as exc:
        print(json.dumps({'status': 'FAIL', 'reason': str(exc)}, ensure_ascii=False))
        return 2
    print(json.dumps(evidence, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
