"""Generate V554-specific Word/PDF release notes and an architecture diagram.
Author: Abdelhalim Talbi. The diagram is explicitly not a GUI screenshot.
The generator embeds the source SHA/PR/workflow identifiers supplied by CI.
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
AUTHOR = "Abdelhalim Talbi"
TITLE = "OpenPDF Studio V51.4 / V554 - ZIP expansion-ratio guard"
OUTPUT_DOCX = "OpenPDF-Studio-V51.4-V554-Guide.docx"
OUTPUT_PDF = "OpenPDF-Studio-V51.4-V554-Guide.pdf"
OUTPUT_SVG = "Architecture-V554-ZIP-Guard.svg"
OUTPUT_PNG = "Architecture-V554-ZIP-Guard.png"


def build_context(*, commit_sha=None, pr_number=None, workflow_run_id=None):
    sha = commit_sha or os.environ.get("OPENPDF_HEAD_SHA") or os.environ.get("GITHUB_SHA", "")
    if not re.fullmatch(r"[0-9a-f]{40}", str(sha)):
        sha = "NOT-VERIFIED"
    pr = str(pr_number if pr_number is not None else os.environ.get("OPENPDF_PULL_REQUEST_NUMBER", "")).strip()
    run_id = str(workflow_run_id if workflow_run_id is not None else os.environ.get("OPENPDF_RUN_ID", "")).strip()
    return {"commit_sha": sha, "pr_number": pr or "not supplied", "workflow_run_id": run_id or "not supplied"}


def build_sections(context):
    run_url = (
        "https://github.com/alimtalbi-ata/OpenPDF-Studio/actions/runs/"
        + context["workflow_run_id"]
        if context["workflow_run_id"].isdigit()
        else "not available"
    )
    pr_url = (
        "https://github.com/alimtalbi-ata/OpenPDF-Studio/pull/" + context["pr_number"]
        if context["pr_number"].isdigit()
        else "not available"
    )
    exact = context["commit_sha"] != "NOT-VERIFIED"
    return [
        ("1. Risque traité",
         "Le vérificateur du ZIP cumulatif conservait déjà des plafonds absolus du nombre de membres "
         "et de la taille décompressée. Une archive très compressée peut toutefois amplifier les ressources "
         "CPU et les E/S consommées avant qu'une limite absolue soit atteinte."),
        ("2. Contrôles V554",
         "Avant de lire le contenu de n'importe quel membre, le packaging inspecte les tailles compressées "
         "et décompressées du répertoire central ZIP. Un membre d'au moins 1 MiB est refusé au-delà de 200:1. "
         "L'ensemble de l'archive est refusé au-delà de 100:1 lorsque la taille décompressée totale atteint "
         "1 MiB. Les plafonds existants restent 200 000 membres, 8 GiB décompressés au total et 4 GiB par membre."),
        ("3. Tests fonctionnels",
         "Quatre tests V554 couvrent une amplification d'un membre, une amplification répartie entre de petits "
         "membres, un ratio borné accepté et des métadonnées de taille invalides. Le workflow de fermeture des "
         "sources exécute cette suite sur Ubuntu et Windows. Les tests passent avant que la phase de packaging "
         "ne démarre dans la même exécution CI."),
        ("4. Qualification Windows et archive",
         "La CI en amont construit le portable Windows, lance ses tests de démarrage/fonctionnement/OCR, "
         "construit l'installateur et vérifie l'installation, le lancement, le runtime OCR et la désinstallation. "
         "La dernière étape du workflow doit ensuite reconstruire le ZIP cumulatif et vérifier CRC, inventaire "
         "SHA-256 et reçus des binaires. Ce guide seul n'est pas un reçu de vérification de l'archive finale."),
        ("5. Limites de preuve",
         "Les ratios sont des seuils de défense conservateurs et peuvent rejeter une archive légitime extrêmement "
         "compressible. Ils complètent les contrôles de noms Unicode et de SHA-256; ils n'authentifient pas "
         "l'éditeur. Aucune signature Authenticode, revue graphique humaine Windows, précision OCR certifiée "
         "ou fidélité visuelle Office garantie n'est revendiquée par V554."),
        ("6. Prochaine priorité",
         "Consolider la chaîne de publication V549-V554 sans fusion automatique, puis mesurer la fidélité "
         "PDF vers DOCX/XLSX/PPTX sur un corpus métier annoté indépendamment, avec résultats par document et page."),
        ("Traçabilité CI",
         f"Commit source: {context['commit_sha']}\nPR: {pr_url}\nRun CI: {run_url}\n"
         + ("La génération est liée au SHA fourni par le workflow." if exact else "Le SHA n'a pas été fourni: qualification non vérifiée.")),
    ]


def svg_diagram():
    # This is a generated architecture diagram, not a screenshot of the application.
    return """<svg xmlns="http://www.w3.org/2000/svg" width="1120" height="330" viewBox="0 0 1120 330" role="img" aria-labelledby="title desc">
<title id="title">OpenPDF Studio V554 ZIP resource preflight</title>
<desc id="desc">Generated architecture diagram showing central-directory metadata checked before decompression.</desc>
<defs><marker id="arrow" markerWidth="9" markerHeight="9" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#24566b"/></marker></defs>
<rect x="18" y="78" width="205" height="96" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="120" y="110" text-anchor="middle" font-family="Segoe UI,Arial" font-size="15" fill="#153747">ZIP central directory</text>
<text x="120" y="134" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#153747">compressed / raw sizes</text>
<path d="M223 126 L260 126" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="267" y="64" width="220" height="124" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="377" y="94" text-anchor="middle" font-family="Segoe UI,Arial" font-size="15" fill="#153747">Resource preflight</text>
<text x="377" y="120" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#153747">Per member: max 200:1</text>
<text x="377" y="144" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#153747">Aggregate: max 100:1</text>
<text x="377" y="166" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#153747">threshold: 1 MiB</text>
<path d="M487 126 L523 126" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="530" y="78" width="175" height="96" rx="12" fill="#fff3df" stroke="#a75c19" stroke-width="2"/>
<text x="617" y="109" text-anchor="middle" font-family="Segoe UI,Arial" font-size="15" fill="#75410d">Limit exceeded?</text>
<text x="617" y="134" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#75410d">Reject before read</text>
<path d="M705 126 L741 126" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="748" y="78" width="165" height="96" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="830" y="109" text-anchor="middle" font-family="Segoe UI,Arial" font-size="15" fill="#153747">Allowed archive</text>
<text x="830" y="134" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#153747">Read / hash members</text>
<path d="M913 126 L949 126" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="956" y="78" width="145" height="96" rx="12" fill="#e7f5e9" stroke="#2b7545" stroke-width="2"/>
<text x="1028" y="109" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#1d5833">CRC + SHA-256</text>
<text x="1028" y="134" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#1d5833">inventory</text>
<rect x="267" y="225" width="834" height="64" rx="8" fill="#fbe7e7" stroke="#a33333" stroke-width="1.5"/>
<text x="684" y="249" text-anchor="middle" font-family="Segoe UI,Arial" font-size="14" fill="#8b2525">Existing hard limits: 200,000 members / 8 GiB total / 4 GiB per member</text>
<text x="684" y="272" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#8b2525">Architecture diagram only - not a real application screenshot or a publisher signature.</text>
</svg>"""


def draw_diagram_png(path):
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1120, 330
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = next((p for p in (
        "C:/Windows/Fonts/segoeui.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ) if Path(p).is_file()), None)
    font = ImageFont.truetype(font_path, 17) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(font_path, 14) if font_path else ImageFont.load_default()
    boxes = [
        (18, 78, 223, 174, "ZIP central directory", "compressed / raw sizes", "#e8f3f5", "#24566b"),
        (267, 64, 487, 188, "Resource preflight", "member 200:1 | all 100:1", "#e8f3f5", "#24566b"),
        (530, 78, 705, 174, "Limit exceeded?", "reject before read", "#fff3df", "#a75c19"),
        (748, 78, 913, 174, "Allowed archive", "read / hash members", "#e8f3f5", "#24566b"),
        (956, 78, 1101, 174, "CRC + SHA-256", "inventory", "#e7f5e9", "#2b7545"),
    ]
    for x0, y0, x1, y1, label, sub, fill, outline in boxes:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=12, fill=fill, outline=outline, width=3)
        draw.text((x0 + 9, y0 + 25), label, font=small, fill="#153747")
        draw.text((x0 + 9, y0 + 57), sub, font=small, fill="#153747")
    for x0, x1 in ((223, 267), (487, 530), (705, 748), (913, 956)):
        cy = 126
        draw.line((x0 + 2, cy, x1 - 7, cy), fill="#24566b", width=3)
        draw.polygon([(x1 - 7, cy), (x1 - 17, cy - 6), (x1 - 17, cy + 6)], fill="#24566b")
    draw.rounded_rectangle((267, 225, 1101, 289), radius=8, fill="#fbe7e7", outline="#a33333", width=2)
    draw.text((285, 240), "Hard limits: 200,000 members / 8 GiB total / 4 GiB per member", font=small, fill="#8b2525")
    draw.text((285, 263), "Architecture diagram only - not a screenshot or a publisher signature.", font=small, fill="#8b2525")
    image.save(path, format="PNG", optimize=True)


def generate(root=ROOT, *, commit_sha=None, pr_number=None, workflow_run_id=None):
    from docx import Document
    from docx.shared import Inches as DocxInches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PdfImage, PageBreak, Table, TableStyle, KeepTogether

    root = Path(root)
    docs = root / "05-Documentation"
    docs.mkdir(parents=True, exist_ok=True)
    context = build_context(commit_sha=commit_sha, pr_number=pr_number, workflow_run_id=workflow_run_id)
    sections = build_sections(context)
    (docs / OUTPUT_SVG).write_text(svg_diagram(), encoding="utf-8")
    draw_diagram_png(docs / OUTPUT_PNG)

    doc = Document()
    doc.core_properties.author = AUTHOR
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "V554 cumulative ZIP security and exact-SHA qualification"
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(TITLE, level=0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("Auteur : ").bold = True
    p.add_run(AUTHOR + "    ")
    p.add_run("Version applicative : ").bold = True
    p.add_run("51.4    ")
    p.add_run("Jalon : ").bold = True
    p.add_run("V554")
    doc.add_paragraph(f"Commit source : {context['commit_sha']}")
    doc.add_paragraph(f"PR : {context['pr_number']}    |    Workflow : {context['workflow_run_id']}")
    if context["commit_sha"] == "NOT-VERIFIED":
        note = doc.add_paragraph("Qualification du SHA non vérifiée : ce document ne constitue pas un reçu de validation.")
        for run in note.runs:
            run.bold = True
    doc.add_picture(str(docs / OUTPUT_PNG), width=DocxInches(6.35))
    cap = doc.add_paragraph("Figure 1. Contrôle des ressources ZIP avant lecture des contenus. Schéma généré, pas une capture de l'application.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, body in sections:
        if heading == "4. Qualification Windows et archive":
            doc.add_page_break()
        doc.add_heading(heading, level=1)
        for line in body.split("\n"):
            if line.strip():
                doc.add_paragraph(line)
    doc.add_heading("Résumé des niveaux de preuve", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, ("Niveau", "V554", "Interprétation")):
        cell.text = value
    rows = [
        ("Tests fonctionnels", "4 tests, Ubuntu + Windows", "Tests de garde de l'archive"),
        ("Windows natif", "Portable + installateur", "Smoke, install, lancement et désinstallation"),
        ("Archive cumulée", "Rejouée dans le workflow", "CRC, inventaire et SHA-256 requis"),
        ("Signature éditeur", "Non revendiquée", "SHA-256 seul n'authentifie pas l'origine"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    doc.save(docs / OUTPUT_DOCX)

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="V554Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12, spaceAfter=6))
    styles.add(ParagraphStyle(name="V554Small", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, leading=9, spaceAfter=4))
    story = [
        Paragraph(escape(TITLE), styles["Title"]),
        Spacer(1, 4 * mm),
        Paragraph(escape(f"Auteur : {AUTHOR} | Version applicative : 51.4 | Jalon : V554"), styles["V554Body"]),
        Paragraph(escape(f"Commit source : {context['commit_sha']} | PR : {context['pr_number']} | Workflow : {context['workflow_run_id']}"), styles["V554Small"]),
        Spacer(1, 3 * mm),
        PdfImage(str(docs / OUTPUT_PNG), width=174 * mm, height=174 * mm * 330 / 1120),
        Paragraph("Figure 1. Schéma d'architecture généré; ce n'est pas une capture de l'application.", styles["Italic"]),
        Spacer(1, 4 * mm),
    ]
    for heading, body in sections:
        story.append(Paragraph(escape(heading), styles["Heading2"]))
        for line in body.split("\n"):
            if line.strip():
                story.append(Paragraph(escape(line), styles["V554Body"]))
    data = [
        ["Niveau", "V554", "Interprétation"],
        ["Tests fonctionnels", "4 tests, Ubuntu + Windows", "Garde archive"],
        ["Windows natif", "Portable + installateur", "Smoke et cycle install/uninstall"],
        ["Archive cumulée", "CRC, inventaire, SHA-256", "Validation finale requise"],
        ["Signature éditeur", "Non revendiquée", "SHA-256 n'authentifie pas l'origine"],
    ]
    table = Table(data, colWidths=[38 * mm, 57 * mm, 79 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24566b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9db3bd")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(Spacer(1, 3 * mm))
    story.append(KeepTogether([Paragraph("Résumé des niveaux de preuve", styles["Heading2"]), table]))
    output_pdf = docs / OUTPUT_PDF
    SimpleDocTemplate(str(output_pdf), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title=TITLE, author=AUTHOR).build(story)
    for path in (docs / OUTPUT_DOCX, output_pdf, docs / OUTPUT_PNG, docs / OUTPUT_SVG):
        if not path.is_file() or path.stat().st_size <= 0:
            raise RuntimeError("Generated document is missing or empty: " + str(path))
    return {"context": context, "docx": str(docs / OUTPUT_DOCX), "pdf": str(output_pdf), "png": str(docs / OUTPUT_PNG), "svg": str(docs / OUTPUT_SVG)}


if __name__ == "__main__":
    result = generate()
    print("V554 versioned documentation generated:")
    for key, value in result.items():
        print(f"{key}: {value}")
