import tempfile,unittest
from pathlib import Path
import fitz
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from openpdf_conversion_audit_v346 import audit_conversion
class TestAudit(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.base=Path(self.tmp.name)
  d=fitz.open();d.new_page();d.save(self.base/'input.pdf');d.close()
 def test_docx(self):
  d=Document();d.add_paragraph('Example');d.save(self.base/'out.docx');r=audit_conversion(self.base/'input.pdf',self.base/'out.docx');self.assertTrue(r.valid);self.assertEqual(len(r.sha256),64)
 def test_xlsx(self):
  w=Workbook();w.active['A1']='Example';w.save(self.base/'out.xlsx');self.assertEqual(audit_conversion(self.base/'input.pdf',self.base/'out.xlsx').output_units,1)
 def test_pptx(self):
  p=Presentation();p.slides.add_slide(p.slide_layouts[6]);p.save(self.base/'out.pptx');self.assertEqual(audit_conversion(self.base/'input.pdf',self.base/'out.pptx').output_units,1)
 def test_missing(self):
  with self.assertRaises(ValueError):audit_conversion(self.base/'input.pdf',self.base/'missing.docx')
 def test_corrupt(self):
  (self.base/'out.docx').write_bytes(b'not a docx')
  with self.assertRaises(ValueError):audit_conversion(self.base/'input.pdf',self.base/'out.docx')
 def test_same_file(self):
  with self.assertRaises(ValueError):audit_conversion(self.base/'input.pdf',self.base/'input.pdf')
if __name__=='__main__':unittest.main()
