"""Generate version-specific V557 Word/PDF documentation and architecture diagrams.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations

import os
import re
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
AUTHOR = "Abdelhalim Talbi"
TITLE = "OpenPDF Studio V51.4 / V557 - Strict JSON release integrity"
DOCX_NAME = "OpenPDF-Studio-V51.4-V557-Strict-JSON-Guide.docx"
PDF_NAME = "OpenPDF-Studio-V51.4-V557-Strict-JSON-Guide.pdf"
PNG_NAME = "Architecture-V557-Strict-JSON.png"
SVG_NAME = "Architecture-V557-Strict-JSON.svg"


def build_context(*, commit_sha=None, pr_number=None, workflow_run_id=None):
    sha = commit_sha
    if sha is None:
        sha = os.environ.get("OPENPDF_HEAD_SHA") or os.environ.get("GITHUB_SHA", "")
    if not isinstance(sha, str) or not re.fullmatch(r"[0-9a-f]{40}", sha):
        sha = "NOT-VERIFIED"
    pr = str(pr_number if pr_number is not None else os.environ.get("OPENPDF_PULL_REQUEST_NUMBER", "")).strip()
    run_id = str(workflow_run_id if workflow_run_id is not None else os.environ.get("OPENPDF_RUN_ID", "")).strip()
    return {"commit_sha": sha, "pr_number": pr or "not supplied", "workflow_run_id": run_id or "not supplied"}


def build_sections(context):
    pr_url = ("https://github.com/alimtalbi-ata/OpenPDF-Studio/pull/" + context["pr_number"]
              if context["pr_number"].isdigit() else "not available")
    run_url = ("https://github.com/alimtalbi-ata/OpenPDF-Studio/actions/runs/" + context["workflow_run_id"]
               if context["workflow_run_id"].isdigit() else "not available")
    sha_verified_format = context["commit_sha"] != "NOT-VERIFIED"
    return [
        ("1. Ambiguïté corrigée",
         "Le parseur JSON standard de Python accepte par défaut des clés répétées et conserve la dernière valeur. Il accepte aussi les constantes numériques NaN et Infinity, qui ne font pas partie du JSON strict interopérable. Des outils différents peuvent donc interpréter un même manifeste ou reçu de façon différente."),
        ("2. Implémentation",
         "V557 introduit un parseur unique utilisant object_pairs_hook pour détecter les clés dupliquées à tous les niveaux, ainsi qu'un rejet explicite des constantes non finies. Il traite l'encodage UTF-8, accepte un BOM UTF-8, et transforme les erreurs en refus explicites. Le contrôle est appliqué à l'inventaire SHA-256 du ZIP, au manifeste binaire, aux reçus de smoke Windows et au reçu d'installation."),
        ("3. Défense en profondeur",
         "La validation stricte du JSON s'ajoute aux contrôles existants : normalisation NFKC/casefold des chemins V553, limites d'amplification V554, cohérence des enregistrements ZIP V555 et liaison du contenu vérifié au même descripteur V556. Les validations de schéma, tailles, empreintes et reçus restent nécessaires après le décodage."),
        ("4. Tests fonctionnels",
         "Les nouveaux tests couvrent les clés dupliquées au niveau racine et imbriqué, les tableaux d'objets, NaN/Infinity/-Infinity, un BOM UTF-8 valide, un JSON tronqué, un UTF-8 invalide et une archive ZIP dont l'inventaire contient réellement une clé répétée. Une archive cohérente reste acceptée."),
        ("5. Limites et qualification",
         "Ces tests sont des tests fonctionnels du parseur et du vérificateur d'archive. Ils ne remplacent pas les builds portables Windows, l'installateur, l'installation/désinstallation, le test GUI natif ni la validation CRC/inventaire/SHA-256 du ZIP cumulatif sur le SHA exact. Un manifeste à clés uniques ne prouve pas l'authenticité de son éditeur."),
        ("6. Prochaine priorité produit",
         "Poursuivre la consolidation de la chaîne de publication V554-V557, puis renforcer le corpus métier indépendant PDF vers DOCX/XLSX/PPTX avec des mesures distinctes de précision, rappel, ordre, géométrie et fidélité visuelle."),
        ("Traçabilité",
         "Commit: " + context["commit_sha"] + "\nPR: " + pr_url + "\nWorkflow: " + run_url + "\n" +
         ("Le SHA a le format attendu; les reçus CI doivent établir la qualification de ce commit."
          if sha_verified_format else
          "Le SHA source est absent ou invalide; ce document est explicitement non qualifié.")),
    ]


def svg_diagram():
    return """<svg xmlns="http://www.w3.org/2000/svg" width="1120" height="360" viewBox="0 0 1120 360" role="img" aria-labelledby="title desc">
<title id="title">OpenPDF Studio V557 strict JSON release integrity</title>
<desc id="desc">Generated architecture diagram, not an application screenshot.</desc>
<defs><marker id="arrow" markerWidth="9" markerHeight="9" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#24566b"/></marker></defs>
<rect x="48" y="92" width="270" height="148" rx="12" fill="#e8f3f5" stroke="#24566b" stroke-width="2"/>
<text x="183" y="124" text-anchor="middle" font-family="Segoe UI,Arial" font-size="16" fill="#153747">JSON inputs</text>
<text x="183" y="154" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#153747">UTF-8 / optional BOM</text>
<text x="183" y="180" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#153747">bounded archive members</text>
<path d="M318 165 L402 165" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="410" y="72" width="300" height="188" rx="12" fill="#fff3df" stroke="#a75c19" stroke-width="2"/>
<text x="560" y="108" text-anchor="middle" font-family="Segoe UI,Arial" font-size="16" fill="#75410d">Strict JSON parser</text>
<text x="560" y="140" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#75410d">unique object keys</text>
<text x="560" y="166" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#75410d">NaN / Infinity rejected</text>
<text x="560" y="192" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#75410d">schema validation follows</text>
<path d="M710 165 L794 165" stroke="#24566b" stroke-width="3" marker-end="url(#arrow)"/>
<rect x="802" y="92" width="270" height="148" rx="12" fill="#e7f5e9" stroke="#2b7545" stroke-width="2"/>
<text x="937" y="124" text-anchor="middle" font-family="Segoe UI,Arial" font-size="16" fill="#1d5833">Fail-closed receipt</text>
<text x="937" y="154" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#1d5833">inventory + SHA-256</text>
<text x="937" y="180" text-anchor="middle" font-family="Segoe UI,Arial" font-size="13" fill="#1d5833">binary / smoke / install</text>
<rect x="48" y="288" width="1024" height="48" rx="8" fill="#f4f0ff" stroke="#6750a4" stroke-width="1.5"/>
<text x="560" y="308" text-anchor="middle" font-family="Segoe UI,Arial" font-size="12" fill="#3e2d73">V553 path normalization | V554 expansion limits | V555 ZIP records | V556 same descriptor</text>
<text x="560" y="326" text-anchor="middle" font-family="Segoe UI,Arial" font-size="11" fill="#3e2d73">Architecture diagram only — not a screenshot or publisher signature.</text>
</svg>"""


def draw_diagram_png(path):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1120, 360), "white")
    draw = ImageDraw.Draw(image)
    font_path = next((candidate for candidate in (
        "C:/Windows/Fonts/segoeui.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ) if Path(candidate).is_file()), None)
    title_font = ImageFont.truetype(font_path, 16) if font_path else ImageFont.load_default()
    body_font = ImageFont.truetype(font_path, 12) if font_path else ImageFont.load_default()
    boxes = [
        (48, 92, 318, 240, "JSON inputs", "UTF-8 / optional BOM", "bounded archive members", "#e8f3f5", "#24566b"),
        (410, 72, 710, 260, "Strict JSON parser", "unique object keys", "NaN / Infinity rejected", "#fff3df", "#a75c19"),
        (802, 92, 1072, 240, "Fail-closed receipt", "inventory + SHA-256", "binary / smoke / install", "#e7f5e9", "#2b7545"),
    ]
    for x0, y0, x1, y1, title, line1, line2, fill, outline in boxes:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=12, fill=fill, outline=outline, width=2)
        draw.text((x0 + 14, y0 + 28), title, font=title_font, fill="#153747")
        draw.text((x0 + 14, y0 + 68), line1, font=body_font, fill="#153747")
        draw.text((x0 + 14, y0 + 94), line2, font=body_font, fill="#153747")
    for start, end in ((318, 410), (710, 802)):
        y = 165
        draw.line((start + 3, y, end - 9, y), fill="#24566b", width=3)
        draw.polygon([(end - 8, y), (end - 19, y - 6), (end - 19, y + 6)], fill="#24566b")
    draw.rounded_rectangle((48, 288, 1072, 336), radius=8, fill="#f4f0ff", outline="#6750a4", width=2)
    draw.text((64, 298), "V553 path normalization | V554 expansion limits | V555 ZIP records | V556 same descriptor",
              font=body_font, fill="#3e2d73")
    draw.text((64, 316), "Architecture diagram only — not a screenshot or publisher signature.",
              font=body_font, fill="#3e2d73")
    image.save(path, format="PNG", optimize=True)


def generate(root=ROOT, *, commit_sha=None, pr_number=None, workflow_run_id=None):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PdfImage

    root = Path(root)
    docs = root / "05-Documentation"
    docs.mkdir(parents=True, exist_ok=True)
    context = build_context(commit_sha=commit_sha, pr_number=pr_number, workflow_run_id=workflow_run_id)
    sections = build_sections(context)
    svg_path, png_path = docs / SVG_NAME, docs / PNG_NAME
    docx_path, pdf_path = docs / DOCX_NAME, docs / PDF_NAME
    svg_path.write_text(svg_diagram(), encoding="utf-8")
    draw_diagram_png(png_path)

    document = Document()
    document.core_properties.author = AUTHOR
    document.core_properties.title = TITLE
    document.core_properties.subject = "Strict JSON parsing for release integrity evidence"
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading(TITLE, level=0)
    intro = document.add_paragraph()
    intro.add_run("Auteur : ").bold = True
    intro.add_run(AUTHOR + "  ")
    intro.add_run("Version applicative : ").bold = True
    intro.add_run("51.4  ")
    intro.add_run("Jalon : ").bold = True
    intro.add_run("V557")
    document.add_paragraph("Commit source : " + context["commit_sha"])
    document.add_paragraph("PR : " + context["pr_number"] + " | Workflow : " + context["workflow_run_id"])
    if context["commit_sha"] == "NOT-VERIFIED":
        note = document.add_paragraph("SHA absent ou invalide : guide non qualifié.")
        for run in note.runs:
            run.bold = True
    document.add_picture(str(png_path), width=Inches(6.25))
    caption = document.add_paragraph("Figure 1. Schéma généré, pas une capture réelle de l'interface.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for heading, body in sections:
        document.add_heading(heading, level=1)
        for line in body.split("\n"):
            if line.strip():
                document.add_paragraph(line)
    document.save(docx_path)

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="V557Body", parent=styles["BodyText"],
                              fontName="Helvetica", fontSize=9, leading=12, spaceAfter=6))
    story = [
        Paragraph(escape(TITLE), styles["Title"]),
        Paragraph(escape("Auteur : " + AUTHOR + " | Commit : " + context["commit_sha"]), styles["V557Body"]),
        Paragraph(escape("PR : " + context["pr_number"] + " | Workflow : " + context["workflow_run_id"]), styles["V557Body"]),
        Spacer(1, 4 * mm),
        PdfImage(str(png_path), width=180 * mm, height=180 * mm * 360 / 1120),
        Spacer(1, 4 * mm),
    ]
    for heading, body in sections:
        story.append(Paragraph(escape(heading), styles["Heading2"]))
        story.append(Paragraph(escape(body).replace("\n", "<br/>"), styles["V557Body"]))
    pdf = SimpleDocTemplate(str(pdf_path), pagesize=A4, rightMargin=14 * mm,
                            leftMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=TITLE, author=AUTHOR)
    pdf.build(story)
    return {"docx": docx_path, "pdf": pdf_path, "png": png_path, "svg": svg_path}


if __name__ == "__main__":
    result = generate()
    for key, path in result.items():
        print(key + ": " + str(path))
