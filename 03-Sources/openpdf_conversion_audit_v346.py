"""Conservative, offline conversion audit. Structural checks are not layout fidelity claims."""
from pathlib import Path
from dataclasses import dataclass,asdict
import hashlib,zipfile
import fitz

@dataclass(frozen=True)
class ConversionAudit:
    source: str
    destination: str
    output_format: str
    source_pages: int
    output_units: int
    output_bytes: int
    sha256: str
    valid: bool
    warnings: tuple[str,...]
    def as_dict(self):
        result=asdict(self)
        result['warnings']=list(self.warnings)
        return result

def audit_conversion(source,destination):
    source,destination=Path(source),Path(destination)
    if not source.is_file() or source.suffix.lower()!='.pdf':
        raise ValueError('Source PDF introuvable ou extension invalide')
    if not destination.is_file() or destination.stat().st_size==0:
        raise ValueError('Sortie absente ou vide')
    if source.resolve()==destination.resolve():
        raise ValueError('La sortie ne peut pas écraser la source')
    with fitz.open(source) as doc:
        if doc.needs_pass: raise ValueError('PDF protégé : audit indisponible')
        pages=len(doc)
    ext=destination.suffix.lower()
    if ext not in ('.docx','.xlsx','.pptx'):
        raise ValueError('Format Office non pris en charge')
    if not zipfile.is_zipfile(destination):
        raise ValueError('Fichier Office non valide (conteneur ZIP absent)')
    marker={'.docx':'word/document.xml','.xlsx':'xl/workbook.xml','.pptx':'ppt/presentation.xml'}[ext]
    with zipfile.ZipFile(destination) as z:
        members=z.infolist()
        # Refuse pathological Office archives before decompressing any entry.
        if len(members)>10000 or sum(m.file_size for m in members)>512*1024*1024:
            raise ValueError('Archive Office anormalement volumineuse')
        if any(m.file_size>128*1024*1024 for m in members):
            raise ValueError('Entrée Office anormalement volumineuse')
        if z.testzip() is not None or marker not in z.namelist():
            raise ValueError('Structure Office corrompue ou incomplète')
    warnings=['Structure valide, mais fidélité de mise en page non mesurée']
    from openpdf_fidelity_v352 import measure_text_coverage
    coverage=measure_text_coverage(source,destination)
    if coverage.coverage is None:
        warnings.append(coverage.note)
    else:
        warnings.append(f'Couverture lexicale indicative : {coverage.coverage:.1%} ({coverage.matched_words}/{coverage.source_words} mots)')
        if coverage.coverage < 0.8:
            warnings.append('Couverture lexicale faible : contrôler le document Office avant utilisation')
    if ext=='.docx':
        from docx import Document
        d=Document(destination)
        units=len(d.paragraphs)+len(d.tables)
        if not any(p.text.strip() for p in d.paragraphs) and not d.tables:
            warnings.append('Aucun texte ni tableau détecté dans le corps du document')
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(destination,read_only=True,data_only=False)
        try: units=len(wb.worksheets)
        finally: wb.close()
    else:
        from pptx import Presentation
        units=len(Presentation(destination).slides)
        if units!=pages: warnings.append('Nombre de diapositives différent du nombre de pages PDF (indicatif seulement)')
    with destination.open('rb') as stream:
        digest=hashlib.file_digest(stream,'sha256').hexdigest()
    return ConversionAudit(str(source),str(destination),ext[1:],pages,units,destination.stat().st_size,digest,True,tuple(warnings))
