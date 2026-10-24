"""V536 real tests: preserve raster edge content during OCR deskew.

Synthetic fixtures are not independent business-document OCR ground truth.
Author: Abdelhalim Talbi.
"""
from __future__ import annotations
import hashlib
import io
from pathlib import Path
import sys
import tempfile
import unittest

import fitz
from PIL import Image, ImageDraw
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '03-Sources'))
from openpdf_ocr_canvas_v536 import rotate_ocr_canvas
from openpdf_ocr_language_preflight_v486 import preflight_ocr_languages
from openpdf_ocr_office_pipeline_v511 import convert_scanned_office

TABLE = (('ID', 'STOCK', 'PRICE'), ('A01', '102', '55'),
         ('A02', '103', '80'), ('A03', '104', '74'))


def has_english_ocr():
    try:
        preflight_ocr_languages('eng')
        return True
    except (RuntimeError, ValueError, OSError):
        return False


def make_scan(path: Path, angle=8):
    with fitz.open() as source:
        page = source.new_page(width=600, height=400)
        for row, values in enumerate(TABLE):
            for col, value in enumerate(values):
                page.insert_text((55 + col * 160, 100 + row * 50), value, fontsize=17)
        raster = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        tilted = Image.frombytes('RGB', (raster.width, raster.height), raster.samples).rotate(
            angle, resample=Image.Resampling.BICUBIC, expand=False, fillcolor='white')
        output = io.BytesIO()
        tilted.save(output, format='PNG')
    with fitz.open() as document:
        document.new_page(width=600, height=400).insert_image(
            fitz.Rect(0, 0, 600, 400), stream=output.getvalue())
        document.save(path)


class OCRCanvasV536(unittest.TestCase):
    def test_rotation_keeps_all_four_page_edge_markers(self):
        image = Image.new('RGB', (800, 500), 'white')
        draw = ImageDraw.Draw(image)
        for x in (2, 740):
            for y in (2, 440):
                draw.rectangle((x, y, x + 55, y + 55), fill='black')
        baseline = image.rotate(-12, resample=Image.Resampling.BICUBIC,
                                expand=False, fillcolor='white')
        corrected = rotate_ocr_canvas(image, -12)
        dark = lambda bitmap: bitmap.convert('L').point(
            lambda pixel: 1 if pixel < 160 else 0).tobytes().count(1)
        self.assertEqual(corrected.size, image.size)
        # The old implementation discards most corner content. The corrected
        # image must retain more than twice the dark foreground area.
        self.assertGreater(dark(corrected), 2 * dark(baseline))
        self.assertGreater(dark(corrected), 6500)

    def test_interior_table_avoids_unnecessary_downsampling(self):
        image = Image.new('RGB', (800, 500), 'white')
        ImageDraw.Draw(image).rectangle((200, 150, 600, 350), fill='black')
        expected = image.rotate(-8, resample=Image.Resampling.BICUBIC,
                                expand=False, fillcolor='white')
        self.assertEqual(rotate_ocr_canvas(image, -8).tobytes(), expected.tobytes())

    def test_zero_rotation_is_lossless_and_preserves_size(self):
        image = Image.new('RGB', (320, 240), 'navy')
        output = rotate_ocr_canvas(image, 0)
        self.assertEqual(output.tobytes(), image.tobytes())
        self.assertEqual(output.size, image.size)
        self.assertIsNot(output, image)

    def test_invalid_angles_fail_before_processing(self):
        image = Image.new('RGB', (320, 240), 'white')
        for angle in (True, None, '8', float('nan'), float('inf'), 12.1, -12.1):
            with self.subTest(angle=angle), self.assertRaisesRegex(ValueError, 'Angle'):
                rotate_ocr_canvas(image, angle)

    def test_oversized_raster_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'Budget'):
            rotate_ocr_canvas(Image.new('RGB', (200, 100)), 8, max_pixels=10000)

    def test_invalid_pixel_budget_is_rejected(self):
        image = Image.new('RGB', (200, 100))
        for limit in (0, True, -1):
            with self.subTest(limit=limit), self.assertRaisesRegex(ValueError, 'Budget'):
                rotate_ocr_canvas(image, 8, max_pixels=limit)

    def test_service_uses_preserving_rotation(self):
        source = (ROOT / '03-Sources/openpdf_ocr_service_v347.py').read_text(encoding='utf-8')
        self.assertIn('corrected = rotate_ocr_canvas(bitmap, applied_angle)', source)
        self.assertNotIn('corrected = bitmap.rotate(', source)

    @unittest.skipUnless(has_english_ocr(), 'English OCR model unavailable')
    def test_manual_deskew_preserves_editable_xlsx_table(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'tilted.pdf'
            destination = Path(folder) / 'result.xlsx'
            make_scan(source, angle=8)
            original = hashlib.sha256(source.read_bytes()).hexdigest()
            result = convert_scanned_office(source, destination, language='eng',
                                            table_mode=True, deskew_angle_degrees=-8)
            self.assertEqual(result.table_cells_by_page, ((1, 12, 12),))
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original)
            self.assertEqual(hashlib.sha256(destination.read_bytes()).hexdigest(), result.sha256)
            workbook = load_workbook(destination)
            try:
                sheet = workbook['P0001-T001']
                self.assertEqual(tuple(tuple(sheet.cell(r, c).value for c in range(1, 4))
                                       for r in range(3, 7)), TABLE)
            finally:
                workbook.close()

    @unittest.skipUnless(has_english_ocr(), 'English OCR model unavailable')
    def test_automatic_deskew_preserves_editable_xlsx_table(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'tilted.pdf'
            destination = Path(folder) / 'result.xlsx'
            make_scan(source, angle=8)
            result = convert_scanned_office(source, destination, language='eng',
                                            table_mode=True, auto_deskew=True)
            self.assertEqual(result.table_cells_by_page, ((1, 12, 12),))
            self.assertAlmostEqual(result.deskew_pages[0][1], -8, delta=0.75)


if __name__ == '__main__':
    unittest.main()
