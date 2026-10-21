"""V46.1 regression: exact native DOCX text and hidden-text fail-closed gate.

Tests use real synthetic PDF/DOCX files on Linux, not Windows qualification.
Author: Abdelhalim Talbi.
"""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '03-Sources'))
import fitz
from docx import Document
from docx.oxml import OxmlElement
from openpdf_core_v260 import pdf_to_docx
from openpdf_docx_roundtrip_v460 import DocxRoundtripError, verify_docx_text


class HiddenDocxV461(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        root = Path(temp.name)
        self.source, self.target = root / 'input.pdf', root / 'output.docx'
        with fitz.open() as pdf:
            pdf.new_page().insert_text((40, 70), 'Texte éditable café', fontsize=12)
            pdf.new_page().insert_text((40, 70), 'Deuxième page', fontsize=12)
            pdf.save(self.source)
        pdf_to_docx(self.source, self.target)

    def test_native_export_passes(self):
        self.assertGreater(verify_docx_text(self.source, self.target), 2)

    def test_hidden_run_fails(self):
        doc = Document(self.target)
        doc.paragraphs[0].runs[0].font.hidden = True
        doc.save(self.target)
        with self.assertRaisesRegex(DocxRoundtripError, 'masqué'):
            verify_docx_text(self.source, self.target)

    def test_hidden_style_fails(self):
        doc = Document(self.target)
        doc.styles['Normal'].font.hidden = True
        doc.save(self.target)
        with self.assertRaisesRegex(DocxRoundtripError, 'masqué'):
            verify_docx_text(self.source, self.target)

    def test_spec_vanish_fails(self):
        doc = Document(self.target)
        doc.paragraphs[0].runs[0]._r.get_or_add_rPr().append(OxmlElement('w:specVanish'))
        doc.save(self.target)
        with self.assertRaisesRegex(DocxRoundtripError, 'masqué'):
            verify_docx_text(self.source, self.target)

    def test_page_break_removed_fails(self):
        doc = Document(self.target)
        node = next(p._element for p in doc.paragraphs if p._p.xpath('.//w:br[@w:type="page"]'))
        node.getparent().remove(node)
        doc.save(self.target)
        with self.assertRaises(DocxRoundtripError):
            verify_docx_text(self.source, self.target)


if __name__ == '__main__':
    unittest.main()
