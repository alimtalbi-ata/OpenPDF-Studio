"""End-to-end smoke: actual PDF -> editable Office, through public dispatcher."""
import sys,tempfile,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'03-Sources'))
import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from openpdf_operations import run_operation

class ActualOfficeConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.source=self.root/'sample.pdf'
        with fitz.open() as pdf:
            page=pdf.new_page()
            page.insert_text((72,72),'ALPHA BETA 123')
            pdf.save(self.source)
    def test_docx_contains_editable_text(self):
        target=self.root/'sample.docx'
        run_operation('docx',[self.source],target)
        doc=Document(target)
        self.assertIn('ALPHA BETA 123',' '.join(p.text for p in doc.paragraphs))
    def test_xlsx_contains_editable_cells(self):
        target=self.root/'sample.xlsx'
        run_operation('xlsx',[self.source],target)
        wb=load_workbook(target)
        try:
            values=[str(cell.value) for row in wb.active for cell in row if cell.value is not None]
            self.assertIn('ALPHA BETA 123',values)
        finally:wb.close()
    def test_pptx_contains_editable_text(self):
        target=self.root/'sample.pptx'
        run_operation('pptx',[self.source],target)
        presentation=Presentation(target)
        self.assertEqual(len(presentation.slides),1)
        content=' '.join(shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame)
        self.assertIn('ALPHA BETA 123',content)
    def test_merge_and_extract_actual_pdf(self):
        merged=self.root/'merged.pdf'
        run_operation('fusion',[self.source,self.source],merged)
        with fitz.open(merged) as pdf:self.assertEqual(pdf.page_count,2)
        extracted=self.root/'extracted.pdf'
        run_operation('extraction',[merged],extracted,first=2,last=2)
        with fitz.open(extracted) as pdf:
            self.assertEqual(pdf.page_count,1)
            self.assertIn('ALPHA BETA 123',pdf[0].get_text())
if __name__=='__main__':unittest.main()
