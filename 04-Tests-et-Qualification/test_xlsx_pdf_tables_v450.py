"""Verify real PDF table geometry is exported into separate Excel cells."""
import sys,tempfile,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'03-Sources'))
import fitz
from openpyxl import load_workbook
from openpdf_operations import run_operation

class PDFTableToExcel(unittest.TestCase):
    def test_vector_table_becomes_cells(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);src=root/'table.pdf';dst=root/'table.xlsx'
            with fitz.open() as pdf:
                page=pdf.new_page(width=400,height=300)
                shape=page.new_shape()
                for x in (40,160,280):
                    shape.draw_line((x,60),(x,160))
                for y in (60,110,160):
                    shape.draw_line((40,y),(280,y))
                shape.finish(color=(0,0,0),width=1);shape.commit()
                for x,y,label in [(55,90,'Name'),(175,90,'Amount'),(55,140,'Alpha'),(175,140,'123')]:
                    page.insert_text((x,y),label,fontsize=12)
                pdf.save(src)
            run_operation('xlsx',[src],dst)
            wb=load_workbook(dst,read_only=True,data_only=False)
            self.assertIn('Texte PDF',wb.sheetnames)
            self.assertIn('Tableau 1',wb.sheetnames)
            rows=list(wb['Tableau 1'].values)
            self.assertTrue(any('Alpha' in row and '123' in row for row in rows))
if __name__=='__main__':unittest.main()
